import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl

from lib2.PARTS.TECH_DZT import part_TECH

from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartOfTECH_DZT_GUI:

    OUTPUT_PARAMS = [
            "APTTECHLGC_1_OILPTRC1_FuncEnabled", "APTTECHLGC_1_OILPTRC1_FuncOperDisabled", "APTTECHLGC_1_OILPTRC1_OpOIL", "APTTECHLGC_1_OILPTRC1_OpOnSignal", "APTTECHLGC_1_OILPTRC1_BlockOIL", "APTTECHLGC_1_WINPTRC1_FuncEnabled", "APTTECHLGC_1_WINPTRC1_FuncOperDisabled", "APTTECHLGC_1_WINPTRC1_OpWIN", "APTTECHLGC_1_WINPTRC1_OpOnSignal", "APTTECHLGC_1_WINPTRC1_BlockWIN", "APTTECHLGC_1_VLVPTRC1_FuncEnabled", "APTTECHLGC_1_VLVPTRC1_FuncOperDisabled", "APTTECHLGC_1_VLVPTRC1_OpVLV", "APTTECHLGC_1_VLVPTRC1_OpOnSignal", "APTTECHLGC_1_VLVPTRC1_BlockVLV", "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_FuncEnabled", "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_FuncOperDisabled", "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_Op", "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_OpOnSignal", "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_FuncEnabled", "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_FuncOperDisabled", "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_Op", "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_OpOnSignal", "ALMTECHLGC_UIRZ_1_LEVPTRC1_FuncEnabled", "ALMTECHLGC_UIRZ_1_LEVPTRC1_FuncOperDisabled", "ALMTECHLGC_UIRZ_1_LEVPTRC1_Op", "ALMTECHLGC_UIRZ_1_LEVPTRC1_OpOnSignal", "TALMGASLGC_1_PTRC1_FuncEnabled", "TALMGASLGC_1_PTRC1_FuncOperDisabled", "TALMGASLGC_1_PTRC1_Op", "TALMGASLGC_1_PTRC1_OpOnSignal", "TALMGASLGC_1_PTRC1_BlockGASProtSign", "TTRGASLGC_1_PTRC1_FuncEnabled", "TTRGASLGC_1_PTRC1_FuncOperDisabled", "TTRGASLGC_1_PTRC1_Op", "TTRGASLGC_1_PTRC1_OpOnSignal", "TTRGASLGC_1_PTRC1_BlockGASProtTrip", "TLTCGASLGC_1_PTRC1_FuncEnabled", "TLTCGASLGC_1_PTRC1_FuncOperDisabled", "TLTCGASLGC_1_PTRC1_Op", "TLTCGASLGC_1_PTRC1_OpOnSignal", "TLTCGASLGC_1_PTRC1_BlockGASLTCProt", 
            "TPRMOFFLVLGC_1_PTRC1_FuncEnabled", "TPRMOFFLVLGC_1_PTRC1_FuncOperDisabled", "pusk_ptrc1_tprmofflvlgc", "TPRMOFFLVLGC_1_PTRC1_Op", "TPRMOFFLVLGC_1_RBRE1_FuncEnabled", "TPRMOFFLVLGC_1_RBRE1_FuncOperDisabled", "TPRMOFFLVLGC_1_RBRE1_BlkOp",
            "DZT2_SignAssembly_1_GASSign", "DZT2_SignAssembly_1_LowIsolGAS", "DZT2_SignAssembly_1_GASBlock", "DZT2_SignAssembly_1_TECHSign", "DZT2_SignAssembly_1_LowIsolTECH", "DZT2_SignAssembly_1_TECHBlock", "DZT2_SignAssembly_1_ALMSign", "DZT2_SignAssembly_1_OpExt", "DZT2_SignAssembly_1_CtlCir", "DZT2_SignAssembly_1_TestBlock", "DZT2_SignAssembly_1_OCSign", "DZT2_SignAssembly_1_GAS_OCControlSignAssem", "DZT2_SignAssembly_1_TECH_OCControlSignAssem", "DZT2_SignAssembly_1_OCnnSign", "DZT2_SignAssembly_1_ExtSignGen", 
            "DZT2_LVALH_1_CALH1_Alarm"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование Газовых и Технологических Защит ДЗТ2. вер.0 от 15.05.25, вер.1 от 30.05.26")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация SGF-параметров
        self.sgf_params = {
            "APTTECHLGC_1_OILPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "APTTECHLGC_1_WINPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "APTTECHLGC_1_VLVPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis": tk.IntVar(value=0),
            "TALMGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TALMGASLGC_1_PTRC1_LowIsolSignCtrl": tk.IntVar(value=0),
            "TTRGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TTRGASLGC_1_PTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "TLTCGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TLTCGASLGC_1_PTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA1": tk.IntVar(value=0),                       
            "DZT2_SignAssembly_1_Ctl_SA2": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA3": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA4": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG2": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG3": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_GAS_OCControl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_TECH_OCControl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_ARCnn1_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_ARCnn2_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_CBFPnn1_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_CBFPnn2_OCControl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }

        # Настройки (T-параметры)
        self.settings = {
            #"T1_oilptrc1_apttechlgc": tk.DoubleVar(value=1),
            #"T1_winptrc1_apttechlgc": tk.DoubleVar(value=1),
            #"T1_vlvptrc1_apttechlgc": tk.DoubleVar(value=1),
            "APTTECHLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),            
            "TALMGASLGC_1_PTRC1_TopOnBlk": tk.DoubleVar(value=1),
            "TTRGASLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),
            "TLTCGASLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),
        }

        # Входные параметры для Step()
        self.input_vars = {
            "VYVOD": tk.IntVar(value=0),
            "OV_tz": tk.IntVar(value=0),
            "OV_dtm": tk.IntVar(value=0),
            "OV_dto": tk.IntVar(value=0),
            "OV_rd": tk.IntVar(value=0),
            "NaSign_dtm": tk.IntVar(value=0),
            "NaSign_dto": tk.IntVar(value=0),
            "NaSign_rd": tk.IntVar(value=0),
            "srabKontOtkl_m": tk.IntVar(value=0),
            "srabKontSign_m": tk.IntVar(value=0),
            "srabKontOtkl_o": tk.IntVar(value=0),
            "srabKontSign_o": tk.IntVar(value=0),
            "srabKontOtkl_rd": tk.IntVar(value=0),
            "srabKI_m": tk.IntVar(value=0),
            "srabKI_o": tk.IntVar(value=0),
            "srabKI_rd": tk.IntVar(value=0),
            "Sbros": tk.IntVar(value=0),
            "OV_ts": tk.IntVar(value=0),
            "OV_pk": tk.IntVar(value=0),
            "OV_ok": tk.IntVar(value=0),
            "OV_lev": tk.IntVar(value=0),
            "NaSign_pk": tk.IntVar(value=0),
            "NaSign_ok": tk.IntVar(value=0),
            "NaSign_lev": tk.IntVar(value=0),
            "srabKontOtkl_pk": tk.IntVar(value=0),
            "srabKontOtkl_ok": tk.IntVar(value=0),
            "srabKontOtkl_lev": tk.IntVar(value=0),
            "OV_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "NaOtkl_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "OV_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "NaSign_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "OV_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "NaSign_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "OV_tprmofflvlgc": tk.IntVar(value=0),
            "OV_ptrc1_tprmofflvlgc": tk.IntVar(value=0),
            "OV_rbre1_tprmofflvlgc": tk.IntVar(value=0),
            "oil_t_hi_level": tk.IntVar(value=0),
            "oil_ltc_hi_level": tk.IntVar(value=0),
            "oil_ltc_lo_level": tk.IntVar(value=0),
            "oil_ltc_lo_temp": tk.IntVar(value=0),
            "otkaz_so": tk.IntVar(value=0),
            "neisp_so": tk.IntVar(value=0),
            "vnesh_otk_zdz1": tk.IntVar(value=0),
            "vnesh_otk_zdz2": tk.IntVar(value=0),
            "vnesh_otk_urov1": tk.IntVar(value=0),
            "vnesh_otk_urov2": tk.IntVar(value=0),
            "vnesh_otkl_so": tk.IntVar(value=0),            
        }

        self.output_labels = {}  # Для вывода результатов

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        self.create_widgets()

    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def _get_output_names(self):
        return self.OUTPUT_PARAMS  

    def create_widgets(self):
        # Фрейм для SGF-параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.sgf_params.items():

            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)

            row += 1
            if row >= 11:
                row = 0
                col += 2

        # Фрейм для настроек (T-параметры)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.settings.items():

            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)

            row += 1
            if row >= 2:
                row = 0
                col += 2

        # Фрейм для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Фрейм для входных параметров
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")

                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):

                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)

                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)

            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Фрейм для выходных параметров
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row, col = 0, 0
        for output in self.OUTPUT_PARAMS:

            label = ttk.Label(output_frame, text=output, width=45, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)

            row += 1
            if row >= 22:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_TECH(
            # Передаем SGF-параметры из self.sgf_params
            SGF1_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_EnaDis"].get(),
            SGF2_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl"].get(),
            #T1_oilptrc1_apttechlgc=self.settings["T1_oilptrc1_apttechlgc"].get(),
            SGF1_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_EnaDis"].get(),
            SGF2_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl"].get(),
            #T1_winptrc1_apttechlgc=self.settings["T1_winptrc1_apttechlgc"].get(),
            SGF1_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_EnaDis"].get(),
            SGF2_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl"].get(),
            #T1_vlvptrc1_apttechlgc=self.settings["T1_vlvptrc1_apttechlgc"].get(),
            T1_apttechlgc=self.settings["APTTECHLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_prvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis"].get(),
            SGF1_shvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis"].get(),
            SGF1_levptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis"].get(),
            SGF1_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_LowIsolSignCtrl"].get(),
            T1_ptrc1_talmgaslgc=self.settings["TALMGASLGC_1_PTRC1_TopOnBlk"].get(),
            SGF1_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_ttrgaslgc=self.settings["TTRGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_tltcgaslgc=self.settings["TLTCGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF1_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG1"].get(),
            SGF6_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG2"].get(),
            SGF7_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG3"].get(),
            SGF8_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_GAS_OCControl_SG1"].get(),
            SGF9_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_TECH_OCControl_SG1"].get(),
            SGF10_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_ARCnn1_OCControl"].get(),
            SGF11_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_ARCnn2_OCControl"].get(),
            SGF12_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_CBFPnn1_OCControl"].get(),
            SGF13_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_CBFPnn2_OCControl"].get(),
        )
        print("part_TECH initialized")

    def start_polling(self):
        if not self.part:
            print("part_TECH not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")


    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfTECH_DZT_GUI(root)
    root.mainloop()