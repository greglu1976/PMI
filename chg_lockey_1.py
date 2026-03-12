import os
import openpyxl

def set_lockey_value():
    """
    Устанавливает значение столбца 'LocKey' = 1 на листе 'Inputs'
    во всех xlsx файлах текущей папки
    """
    current_dir = os.getcwd()
    xlsx_files = [f for f in os.listdir(current_dir) if f.endswith('.xlsx') and not f.startswith('~$')]
    
    print(f"Найдено {len(xlsx_files)} файлов")
    print("-" * 50)
    
    for file_name in xlsx_files:
        file_path = os.path.join(current_dir, file_name)
        print(f"Обработка: {file_name}")
        
        try:
            wb = openpyxl.load_workbook(file_path)
            
            if "Inputs" in wb.sheetnames:
                sheet = wb["Inputs"]
                lockey_col = None
                
                # Поиск столбца LocKey
                for col in range(1, sheet.max_column + 1):
                    if sheet.cell(row=1, column=col).value == "LocKey":
                        lockey_col = col
                        break
                
                # Заполнение значениями
                if lockey_col:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row=row, column=lockey_col, value=1)
                    wb.save(file_path)
                    print(f"  ✅ LocKey установлен в 1 ({sheet.max_row - 1} строк)")
                else:
                    print(f"  ⚠️ Столбец 'LocKey' не найден")
            else:
                print(f"  ⚠️ Лист 'Inputs' не найден")
            
            wb.close()
            
        except Exception as e:
            print(f"  ❌ Ошибка: {e}")
    
    print("-" * 50)
    print("Готово!")
    input("Нажмите Enter, чтобы выйти...")

if __name__ == "__main__":
    set_lockey_value()