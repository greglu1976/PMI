import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl

from lib._PARTS.LO_T import part_LO 

class PartLO_GUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ЛО, УРОВ, СС, ПС, ЛО ВН, ЛО НН М300-Т. вер.0 от 07.04.25, вер.1 - 2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация SGF-параметров
        self.sgf_params = {
            "LVTTOC_1_KschemeCT": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTUV1_VoltStrCond": tk.IntVar(value=0),
            "LVTTOC_1_PHAR1_RegBlock": tk.IntVar(value=0),
            "LVTTOC_1_RBLC1_StepSel": tk.IntVar(value=0),
            "T_LVTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "T_LVTOC_1_PTOC1_KschemeCT": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_EnaDis": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_RstFrmCLS": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_KnobCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_EnaDis": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CurrentPickUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CBOSTypeCtrl": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_ActUpSwitch": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent": tk.IntVar(value=0),
            "TOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC3_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl": tk.IntVar(value=0),
            "T_HVTCBOFF_1_HVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVCBRECRBRE1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVBTSRBLC1_EnaDis": tk.IntVar(value=0), 
            "T_SignAssembly_1_Ctl_SA1": tk.IntVar(value=0),                       
            "T_SignAssembly_1_Ctl_SA2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA3": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA4": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA5": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG1": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_GAS_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_TECH_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_OCcir_CB": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_ARCnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_CBFPnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_IEDvt_OCControl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }

        # Настройки (T-параметры)
        self.settings = {
            "LVTTOC_1_PTOC1_Top": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTOC1_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC1_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_Top": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC3_Top": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTOC3_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC3_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTUV1_Uop": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTUV1_U2op": tk.DoubleVar(value=1),
            "LVTTOC_1_PHAR1_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PHAR1_PhStr": tk.DoubleVar(value=40),
            "T_LVTOC_1_PTOC1_Top": tk.DoubleVar(value=1),
            "T_LVTOC_1_PTOC1_Iop": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_EnBlk": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_ElmgWorking": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Top": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Iop": tk.DoubleVar(value=1),
            "T_HVTCBOFF_1_HVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
            "T1_lvcbptrc1_lvtcboff": tk.DoubleVar(value=1),
        }

        # Входные параметры для Step()
        self.input_vars = {
            "VYVOD": tk.IntVar(value=0),
            "OV_lvttoc": tk.IntVar(value=0),
            "OV_ptoc1_lvtoc": tk.IntVar(value=0),
            "OV_tofflvlg": tk.IntVar(value=0),
            "OVlo_tofflvlg": tk.IntVar(value=0),
            "OVzapv_tofflvlg": tk.IntVar(value=0),
            "OVzavr_tofflvlg": tk.IntVar(value=0),
            "OV_hvcbptrc1_hvtcboff": tk.IntVar(value=0),
            "vnesh_otkl_zdz": tk.IntVar(value=0),
            "vnesh_otkl_urov": tk.IntVar(value=0),
            "OV_lvtcboff": tk.IntVar(value=0),
            "OV_lvcbptrc1_lvtcboff": tk.IntVar(value=0),
            "OV_lvcbrecrbre1_lvtcboff": tk.IntVar(value=0),
            "OV_lvbtsrblc1_lvtcboff": tk.IntVar(value=0),
            "OV_rcbf1_lvcbsup": tk.IntVar(value=0),
            "vnesh_blok_upr_V": tk.IntVar(value=0),
            "OV_rbrf1_tpbrf": tk.IntVar(value=0),
            "pusk_urov_vnesh": tk.IntVar(value=0),
            "kontr_emo1": tk.IntVar(value=0),
            "kontr_emo2": tk.IntVar(value=0),                        
            "Polozh_SA1": tk.IntVar(value=0),
            "Polozh_SA2": tk.IntVar(value=0),
            "Polozh_SA3": tk.IntVar(value=0),
            "Polozh_SA4": tk.IntVar(value=0),
            "Polozh_SA5": tk.IntVar(value=0),
            "Polozh_SG1": tk.IntVar(value=0),
            "Polozh_SG2": tk.IntVar(value=0),
            "ot_gz": tk.IntVar(value=0),
            "ot_tz": tk.IntVar(value=0),
            "ot_v": tk.IntVar(value=0),
            "ot_zdz_nn": tk.IntVar(value=0),
            "ot_urov_nn": tk.IntVar(value=0),
            "ot_ieu_tn": tk.IntVar(value=0),
            "vnesh_sign1": tk.IntVar(value=0),
            "vnesh_sign2": tk.IntVar(value=0),
            "vnesh_sign3": tk.IntVar(value=0),
            "vnesh_sign4": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),            
        }

        self.output_labels = {}  # Для вывода результатов
        self.create_widgets()

    def create_widgets(self):
        # Фрейм для SGF-параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 16:
                row = 0
                col += 2

        # Фрейм для настроек (T-параметры)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Фрейм для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Фрейм для входных параметров
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(input_frame, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(input_frame, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Фрейм для выходных параметров
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")
        outputs = [
            "mtz_pusk_lvttoc", "mtz_srab_ptoc1_lvttoc", "mtz_srab_ptoc2_lvttoc", "mtz_srab_ptoc3_lvttoc", "pusk_ptoc1_lvtoc", "srab_ptoc1_lvtoc",
        "vvod_ptrc1_tofflvlgc", "oper_vyvod_ptrc1_tofflvlgc", "pusk_ptrc1_tofflvlgc", "srab_ptrc1_tofflvlgc", "vvod_rblc1_tofflvlgc", "oper_vyvod_rblc1_tofflvlgc", "zapret_rblc1_tofflvlgc", "vvod_rbre1_tofflvlgc", "oper_vyvod_rbre1_tofflvlgc", "zapret_rbre1_tofflvlgc", "vvod_hvcbptrc1_hvtcboff", "oper_vyvod_hvcbptrc1_hvtcboff", "otkl_hvcbptrc1_hvtcboff", "otkl_avar_hvcbptrc1_hvtcboff", "vvod_lvcbptrc1_lvtcboff", "oper_vyvod_lvcbptrc1_lvtcboff", "otkl_lvcbptrc1_lvtcboff", "otkl_avar_lvcbptrc1_lvtcboff", "vvod_lvcbrecrbre1_lvtcboff", "oper_vyvod_lvcbrecrbre1_lvtcboff", "zapret_lvcbrecrbre1_lvtcboff", "vvod_lvbtsrblc1_lvtcboff", "oper_vyvod_lvbtsrblc1_lvtcboff", "zapret_lvbtsrblc1_lvtcboff", "blok_otkl_rcbf1_lvcbsup",
       "vvod_rbrf1_tpbrf", "oper_vyvod_rbrf1_tpbrf", "uskorenie_rbrf1_tpbrf", "srab_rbrf1_tpbrf", "pusk_rbrf1_tpbrf", "io_rbrf1_tpbrf", "srab_na_sebya_rbrf1_tpbrf",
        "SS_gz_sign", "SS_gz_zablok", "SS_gz_nizk_isol", "SS_tz_sign", "SS_tz_nizk_isol", "SS_tz_zablok", "SS_ts_sign", "SS_vnesh_otkl", "SS_vyh_zepi_razobr", "SS_bi_vyved", "SS_ot_sign", "SS_neispr_ot_gz",  "SS_neispr_ot_tz", "SS_neispr_ot_v", "SS_ot_nn_sign", "SS_prev_vrem_per_ka", "SS_obsh_vnesh_sign", "pusk_lvalv"
        ]

        row, col = 0, 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=33, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_LO(
            # Передаем SGF-параметры из self.sgf_params
            SGF1=self.sgf_params["LVTTOC_1_KschemeCT"].get(),
            SGF1_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VolMod"].get(),
            SGF3_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_MICMod"].get(),
            SGF4_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_ExtVFlMod"].get(),
            SGF5_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VCMod"].get(),
            SGF6_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_SBMod"].get(),
            SGF1_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_EnaDis"].get(),
            SGF2_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VolMod"].get(),
            SGF3_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_MICMod"].get(),
            SGF4_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_ExtVFlMod"].get(),
            SGF5_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VCMod"].get(),
            SGF6_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_SBMod"].get(),
            SGF1_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_EnaDis"].get(),
            SGF2_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VolMod"].get(),
            SGF3_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_MICMod"].get(),
            SGF4_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_ExtVFlMod"].get(),
            SGF5_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VCMod"].get(),
            SGF6_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_SBMod"].get(),
            SGF1_ptuv1=self.sgf_params["LVTTOC_1_PTUV1_VoltStrCond"].get(),
            SGF1_phar1=self.sgf_params["LVTTOC_1_PHAR1_RegBlock"].get(),
            SGF1_rblc1=self.sgf_params["LVTTOC_1_RBLC1_StepSel"].get(),
            SGF1_ptoc1_lvtoc=self.sgf_params["T_LVTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1_lvtoc=self.sgf_params["T_LVTOC_1_PTOC1_KschemeCT"].get(),
            SGF1_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_EnaDis"].get(),
            SGF2_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol"].get(),
            SGF3_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault"].get(),
            SGF4_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_RstFrmCLS"].get(),
            SGF5_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl"].get(),
            SGF6_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault"].get(),
            SGF7_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_KnobCtrl"].get(),
            SGF8_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm"].get(),
            SGF1_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_EnaDis"].get(),
            SGF2_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp"].get(),
            SGF3_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CurrentPickUp"].get(),
            SGF4_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CBOSTypeCtrl"].get(),
            SGF5_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_ActUpSwitch"].get(),
            SGF6_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF2_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC2_Ctrl"].get(),
            SGF3_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC3_Ctrl"].get(),
            SGF1_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_EnaDis"].get(),
            SGF2_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl"].get(),
            SGF3_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl"].get(),
            SGF1_hvcbptrc1_hvtcboff=self.sgf_params["T_HVTCBOFF_1_HVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbptrc1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbrecrbre1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVCBRECRBRE1_EnaDis"].get(),
            SGF1_lvbtsrblc1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVBTSRBLC1_EnaDis"].get(),
            SGF1_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["T_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["T_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_lvalh=self.sgf_params["T_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_lvalh=self.sgf_params["T_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_lvalh=self.sgf_params["T_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF14_lvalh=self.sgf_params["T_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA5"].get(),
            SGF6_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SG1"].get(),
            SGF7_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SG2"].get(),
            SGF8_tsa=self.sgf_params["T_SignAssembly_1_Ctl_GAS_OCControl"].get(),
            SGF9_tsa=self.sgf_params["T_SignAssembly_1_Ctl_TECH_OCControl"].get(),
            SGF10_tsa=self.sgf_params["T_SignAssembly_1_Ctl_OCcir_CB"].get(),
            SGF11_tsa=self.sgf_params["T_SignAssembly_1_Ctl_ARCnn_OCControl"].get(),
            SGF12_tsa=self.sgf_params["T_SignAssembly_1_Ctl_CBFPnn_OCControl"].get(),
            SGF13_tsa=self.sgf_params["T_SignAssembly_1_Ctl_IEDvt_OCControl"].get(),
            T1_ptoc1=self.settings["LVTTOC_1_PTOC1_Top"].get(),
            Iset_ptoc1=self.settings["LVTTOC_1_PTOC1_Iop"].get(),
            Icoarse_ptoc1=self.settings["LVTTOC_1_PTOC1_IopCoars"].get(),
            T1_ptoc2=self.settings["LVTTOC_1_PTOC2_Top"].get(),
            Iset_ptoc2=self.settings["LVTTOC_1_PTOC2_Iop"].get(),
            Icoarse_ptoc2=self.settings["LVTTOC_1_PTOC2_IopCoars"].get(),
            T1_ptoc3=self.settings["LVTTOC_1_PTOC3_Top"].get(),
            Iset_ptoc3=self.settings["LVTTOC_1_PTOC3_Iop"].get(),
            Icoarse_ptoc3=self.settings["LVTTOC_1_PTOC3_IopCoars"].get(),
            Uop_ptuv1=self.settings["LVTTOC_1_PTUV1_Uop"].get(),
            U2op_ptuv1=self.settings["LVTTOC_1_PTUV1_U2op"].get(),
            Imax_phar1=self.settings["LVTTOC_1_PHAR1_Iop"].get(),
            Ratio_phar1=self.settings["LVTTOC_1_PHAR1_PhStr"].get(),
            T1_ptoc1_lvtoc=self.settings["T_LVTOC_1_PTOC1_Top"].get(),
            Iset_ptoc1_lvtoc=self.settings["T_LVTOC_1_PTOC1_Iop"].get(),
            T1_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),
            T2_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(),
            T3_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),
            T1_rbrf1_tpbrf=self.settings["T_TPBRF_1_GENRBRF1_Top"].get(),
            Iset_rbrf1_tpbrf=self.settings["T_TPBRF_1_GENRBRF1_Iop"].get(),
            T1_hvcbptrc1_hvtcboff=self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(),
            T1_lvcbptrc1_lvtcboff=self.settings["T1_lvcbptrc1_lvtcboff"].get(),
        )
        print("part_LO initialized")

    def start_polling(self):
        if not self.part:
            print("part_LO not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")


    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = PartLO_GUI(root)
    root.mainloop()