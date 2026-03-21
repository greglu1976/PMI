# автоматическое тестирование ФСУ в части МТЗ, КЦН, ПС + ЛЗТ, ТК ЗДЗ для исполнения Т
# с графическим интерфейсом

import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension

import itertools
import openpyxl
import json
from lib2.PARTS.MTZ_T_inout import partOfFsuInTOC
from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartOfFsuInTOC_GUI:

    # === СПИСОК ВЫХОДНЫХ ПАРАМЕТРОВ (единое определение) ===
    OUTPUT_PARAMS = [            
            "LVRBVTR_1_RVTR1_FuncEnabled", "LVRBVTR_1_RVTR1_FuncOperDisabled", "LVRBVTR_1_RVTR1_DE_Upp", "LVRBVTR_1_RVTR1_DE_U2", "LVRBVTR_1_RVTR1_Str", "LVRBVTR_1_RVTR1_Alm",
            "LVTTOC_1_PTOC1_FuncEnabled", "LVTTOC_1_PTOC1_FuncOperDisabled", "LVTTOC_1_PTOC1_IAStr", "LVTTOC_1_PTOC1_IBStr", "LVTTOC_1_PTOC1_ICStr",
            "LVTTOC_1_PTOC1_Str", "LVTTOC_1_PTOC1_OpOnSignal", "LVTTOC_1_PTOC1_Op", "LVTTOC_1_PTOC1_DE_IA", "LVTTOC_1_PTOC1_DE_IB",
            "LVTTOC_1_PTOC1_DE_IC", "LVTTOC_1_PTOC2_FuncEnabled", "LVTTOC_1_PTOC2_FuncOperDisabled", "LVTTOC_1_PTOC2_IAStr", "LVTTOC_1_PTOC2_IBStr",
            "LVTTOC_1_PTOC2_ICStr", "LVTTOC_1_PTOC2_Str", "LVTTOC_1_PTOC2_OpOnSignal", "LVTTOC_1_PTOC2_Op", "LVTTOC_1_PTOC2_DE_IA",
            "LVTTOC_1_PTOC2_DE_IB", "LVTTOC_1_PTOC2_DE_IC", "LVTTOC_1_PTOC3_FuncEnabled", "LVTTOC_1_PTOC3_FuncOperDisabled", "LVTTOC_1_PTOC3_IAStr",
            "LVTTOC_1_PTOC3_IBStr", "LVTTOC_1_PTOC3_ICStr", "LVTTOC_1_PTOC3_Str", "LVTTOC_1_PTOC3_OpOnSignal", "LVTTOC_1_PTOC3_Op",
            "LVTTOC_1_PTOC3_DE_IA", "LVTTOC_1_PTOC3_DE_IB", "LVTTOC_1_PTOC3_DE_IC", "LVTTOC_1_PTUV1_UndervoltStr",
            "LVTTOC_1_PHAR1_PharmonicStrA", "LVTTOC_1_PHAR1_PharmonicStrB", "LVTTOC_1_PHAR1_PharmonicStrC", "LVTTOC_1_PHAR1_PharmonicStr", "LVTTOC_1_RBLC1_LbpBlkOp",
            "LVTTOC_1_LVTTOC_Str", "TOFFLVLGC_1_PTRC1_FuncEnabled", "TOFFLVLGC_1_PTRC1_FuncOperDisabled", "TOFFLVLGC_1_PTRC1_Str", "TOFFLVLGC_1_PTRC1_Op", "TOFFLVLGC_1_LVCBRBLC1_FuncEnabled", "TOFFLVLGC_1_LVCBRBLC1_FuncOperDisabled", "TOFFLVLGC_1_LVCBRBLC1_BlkOp",
            "TOFFLVLGC_1_RBRE1_FuncEnabled", "TOFFLVLGC_1_RBRE1_FuncOperDisabled", "TOFFLVLGC_1_RBRE1_BlkOp", "T_LVALH_1_CALH1_Alarm", 
            "T_LVARCTOC_1_PTOC1_Str", 
            "TTOCLGC_UIRZ_1_PTRC1_Str", 
            "IAB", "dIAB", "IBC", "dIBC", "ICA", "dICA", "I2", "I0", "I1", "UAB_ptuv1", "UBC_ptuv1", "UCA_ptuv1", "U2_ptuv1", "U0_ptuv1", "U1_ptuv1"
            ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ (исполнение Т) в части МТЗ, КЦН НН1, КЦН НН2, ЛО Т, CC, ПС, ТК ЗДЗ, ЛЗТ v2.0 от 21.05.25, v2.1 от 25.07.25, v3 от 2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для SGF параметров, настроек, входных и выходных значений
        self.sgf_params = {
            "LVTTOC_1_KschemeCT": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VolMod": tk.IntVar(value=1),
            "LVTTOC_1_PTOC1_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTUV1_VoltStrCond": tk.IntVar(value=0),
            "LVTTOC_1_PHAR1_RegBlock": tk.IntVar(value=0),
            "LVTTOC_1_RBLC1_StepSel": tk.IntVar(value=1),
            "LVRBVTR_1_RVTR1_EnaDis": tk.IntVar(value=0),
            "LVRBVTR_1_RVTR1_StrMod": tk.IntVar(value=0),
            "TOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0), 
            "TOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC3_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),            
            "T_LVARCTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "T_LVARCTOC_1_PTOC1_StrMod": tk.IntVar(value=0),
            "TTOCLGC_UIRZ_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TTOCLGC_UIRZ_1_PTRC1_LVTPTOC2_Ctrl": tk.IntVar(value=0),
            "TTOCLGC_UIRZ_1_PTRC1_LVTPTOC3_Ctrl": tk.IntVar(value=0), 
            #"Номинальный ток входа": tk.IntVar(value=5),            
        }

        self.settings = {
            "LVTTOC_1_PTOC1_Top": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC1_Iop": tk.DoubleVar(value=0.2),
            "LVTTOC_1_PTOC1_IopCoars": tk.DoubleVar(value=0.6),
            "LVTTOC_1_PTOC2_Top": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_Iop": tk.DoubleVar(value=0.2),
            "LVTTOC_1_PTOC2_IopCoars": tk.DoubleVar(value=0.6),
            "LVTTOC_1_PTOC3_Top": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC3_Iop": tk.DoubleVar(value=0.2),
            "LVTTOC_1_PTOC3_IopCoars": tk.DoubleVar(value=0.6),
            "LVTTOC_1_PTUV1_Uop": tk.DoubleVar(value=50),
            "LVTTOC_1_PTUV1_U2op": tk.DoubleVar(value=20),
            "LVTTOC_1_PHAR1_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PHAR1_PhStr": tk.DoubleVar(value=40),
            "LVRBVTR_1_RVTR1_Uop": tk.DoubleVar(value=50),
            "LVRBVTR_1_RVTR1_U2op": tk.DoubleVar(value=20),
            "LVRBVTR_1_RVTR1_Top": tk.DoubleVar(value=1),
            "T_LVARCTOC_1_PTOC1_Iop": tk.DoubleVar(value=1), 
            "TTOCLGC_UIRZ_1_PTRC1_Top": tk.DoubleVar(value=1),                       
        }

        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "DI_LVTTOC": tk.IntVar(value=0),
            "DI_PTOC1": tk.IntVar(value=0),
            "DI_PTOC2": tk.IntVar(value=0),
            "DI_PTOC3": tk.IntVar(value=0),
            "DI_PTOC1_Sign": tk.IntVar(value=0),
            "DI_PTOC2_Sign": tk.IntVar(value=0),
            "DI_PTOC3_Sign": tk.IntVar(value=0),
            "SBnnPosCls": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "dIA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "dIB": tk.DoubleVar(value=240),
            "IC": tk.DoubleVar(value=0),
            "dIC": tk.DoubleVar(value=120),
            "UA1": tk.DoubleVar(value=58),
            "dUA1": tk.DoubleVar(value=0),
            "UB1": tk.DoubleVar(value=58),
            "dUB1": tk.DoubleVar(value=240),
            "UC1": tk.DoubleVar(value=58),
            "dUC1": tk.DoubleVar(value=120),
            "IA2harm": tk.DoubleVar(value=0),
            "IB2harm": tk.DoubleVar(value=0),
            "IC2harm": tk.DoubleVar(value=0),
            "OutVoltStr": tk.IntVar(value=0),            
            "CBnnPosCls": tk.IntVar(value=0),
            "DI_LVRBVTR": tk.IntVar(value=0),
            "OutBlkV": tk.IntVar(value=0),
            "DI_TOFFLVLGC": tk.IntVar(value=0),
            "DI_PTRC1": tk.IntVar(value=0),
            "DI_RBRE1": tk.IntVar(value=0),
            "DI_LVCBRBLC1": tk.IntVar(value=0),
        }

        self.output_labels = {}

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        # Создание интерфейса
        self.create_widgets()

    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            if key=="LVTTOC_1_PTOC1_SBMod" or key=="LVTTOC_1_PTOC2_SBMod" or key=="LVTTOC_1_PTOC3_SBMod" or key=="LVTTOC_1_PTUV1_VoltStrCond" or key=="SGF1_ptuv2_lvttoc":
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            elif key=="LVTTOC_1_PTOC1_VolMod" or key=="LVTTOC_1_PTOC2_VolMod" or key=="LVTTOC_1_PTOC3_VolMod":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[1, 2], state="readonly").grid(row=row, column=col + 1)
            elif key=="LVTTOC_1_PTOC1_ExtVFlMod" or key=="LVTTOC_1_PTOC2_ExtVFlMod" or key=="LVTTOC_1_PTOC3_ExtVFlMod":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[1, 2], state="readonly").grid(row=row, column=col + 1)                
            elif key=="LVTTOC_1_RBLC1_StepSel":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[1, 2, 3, 4], state="readonly").grid(row=row, column=col + 1)                
            elif key=="T_LVARCTOC_1_PTOC1_StrMod":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2, 3], state="readonly").grid(row=row, column=col + 1)
            elif key=='Номинальный ток входа':
                ttk.Combobox(sgf_frame, textvariable=var, values=[1, 5], state="readonly").grid(row=row, column=col + 1)                   
            else:
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 11:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)

        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)

        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)

        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)

        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        
        # Button Load JSON
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        # Button Save JSON
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)
                
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 4:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=25, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)


            row += 1
            if row >= 28:
                row = 0
                col += 2

    def init_part(self):
        self.part = partOfFsuInTOC(
            SGF1=self.sgf_params["LVTTOC_1_KschemeCT"].get(),
            SGF1_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VolMod"].get(),
            SGF3_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_MICMod"].get(),
            SGF4_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_ExtVFlMod"].get(),
            SGF5_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VCMod"].get(),
            SGF6_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_SBMod"].get(),
            T1_ptoc1=self.settings["LVTTOC_1_PTOC1_Top"].get(),
            Iset_ptoc1=self.settings["LVTTOC_1_PTOC1_Iop"].get(),
            Icoarse_ptoc1=self.settings["LVTTOC_1_PTOC1_IopCoars"].get(),
            SGF1_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_EnaDis"].get(),
            SGF2_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VolMod"].get(),
            SGF3_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_MICMod"].get(),
            SGF4_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_ExtVFlMod"].get(),
            SGF5_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VCMod"].get(),
            SGF6_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_SBMod"].get(),
            T1_ptoc2=self.settings["LVTTOC_1_PTOC2_Top"].get(),
            Iset_ptoc2=self.settings["LVTTOC_1_PTOC2_Iop"].get(),
            Icoarse_ptoc2=self.settings["LVTTOC_1_PTOC2_IopCoars"].get(),
            SGF1_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_EnaDis"].get(),
            SGF2_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VolMod"].get(),
            SGF3_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_MICMod"].get(),
            SGF4_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_ExtVFlMod"].get(),
            SGF5_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VCMod"].get(),
            SGF6_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_SBMod"].get(),
            T1_ptoc3=self.settings["LVTTOC_1_PTOC3_Top"].get(),
            Iset_ptoc3=self.settings["LVTTOC_1_PTOC3_Iop"].get(),
            Icoarse_ptoc3=self.settings["LVTTOC_1_PTOC3_IopCoars"].get(),
            SGF1_ptuv1=self.sgf_params["LVTTOC_1_PTUV1_VoltStrCond"].get(),
            Uop_ptuv1=self.settings["LVTTOC_1_PTUV1_Uop"].get(),
            U2op_ptuv1=self.settings["LVTTOC_1_PTUV1_U2op"].get(),
            SGF1_phar1=self.sgf_params["LVTTOC_1_PHAR1_RegBlock"].get(),
            Imax_phar1=self.settings["LVTTOC_1_PHAR1_Iop"].get(),
            Ratio_phar1=self.settings["LVTTOC_1_PHAR1_PhStr"].get(),
            SGF1_rblc1=self.sgf_params["LVTTOC_1_RBLC1_StepSel"].get(),
            SGF1_lvrbvtr1=self.sgf_params["LVRBVTR_1_RVTR1_EnaDis"].get(),
            SGF2_lvrbvtr1=self.sgf_params["LVRBVTR_1_RVTR1_StrMod"].get(),
            u_min_lvrbvtr1=self.settings["LVRBVTR_1_RVTR1_Uop"].get(),
            u2_max_lvrbvtr1=self.settings["LVRBVTR_1_RVTR1_U2op"].get(),
            t1_lvrbvtr1=self.settings["LVRBVTR_1_RVTR1_Top"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_EnaDis"].get(), 
            SGF2_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC2_Ctrl"].get(), 
            SGF3_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC3_Ctrl"].get(), 
            SGF1_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_EnaDis"].get(), 
            SGF2_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl"].get(), 
            SGF3_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl"].get(),
            SGF1_ptoc1_lvarctoc=self.sgf_params["T_LVARCTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1_lvarctoc=self.sgf_params["T_LVARCTOC_1_PTOC1_StrMod"].get(),
            SGF1_ptrc1_ttoclgc=self.sgf_params["TTOCLGC_UIRZ_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_ttoclgc=self.sgf_params["TTOCLGC_UIRZ_1_PTRC1_LVTPTOC2_Ctrl"].get(),
            SGF3_ptrc1_ttoclgc=self.sgf_params["TTOCLGC_UIRZ_1_PTRC1_LVTPTOC3_Ctrl"].get(),
            Iset_ptoc1_lvarctoc=self.settings["T_LVARCTOC_1_PTOC1_Iop"].get(),
            Inom=5, #self.sgf_params["Номинальный ток входа"].get(), Убрал чтобы в уставки М300 не влезало
            T1_ptrc1_ttoclgc=self.settings["TTOCLGC_UIRZ_1_PTRC1_Top"].get() 
        )

        print("partOfFsuInTOC initialized")

    def start_polling(self):
        if self.part is None:
            print("partOfFsuInTOC not initialized")
            return

        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            #print(inputs)
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #label.config(text=f"{output}: {int(value)}")
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or round(float(value),2)!=0.0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3) # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05) # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):

        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return

        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        #output_file = "data.xlsx"
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)

        # Определяем красный цвет для заливки
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                # Устанавливаем ширину столбца
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20

                # Проверяем значения и применяем форматирование
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)  # Преобразуем значение в число
                        if value != 0:
                            cell.fill = red_fill  # Выделяем красным, если значение не равно 0
                    except (ValueError, TypeError):
                        pass  # Игнорируем ошибки преобразования

        # Применяем форматирование к каждому листу
        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        # Сохраняем изменения
        wb.save(output_file)

        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        # Выбор файла для загрузки
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        try:
            # Чтение данных из Excel
            xls = pd.ExcelFile(file_path)

            # Загрузка SGF Parameters
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            # Загрузка Inputs
            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])                    

            print("Data loaded successfully")

        except Exception as e:
            print(f"Error loading data: {e}")
            
    # === МЕТОДЫ ДЛЯ РАБОТЫ С JSON ФАЙЛАМИ УСТАВОК ===
    
    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла с суффиксом _SG1."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return
        try:
            handler = SettingsHandler.from_json_file(file_path)
            if not self.meta_handler:
                print("⚠️ Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                    elif type_str == "130":  # Integer
                        try:
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()

                            int_val = int(float(value_clean.replace(',', '.')))
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            value_clean = str(value_str).strip()
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            # --- Обновление T-параметров (settings) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                        else:
                            self.settings[key].set(float_val)
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            print("✅ Параметры обновлены из JSON (с суффиксом _SG1)")
        except FileNotFoundError:
            print(f"❌ Ошибка: Файл не найден: {file_path}")
        except Exception as e:
            print(f"❌ Ошибка загрузки JSON: {e}")
    
    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()

    def _get_output_names(self):
        return self.OUTPUT_PARAMS     


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfFsuInTOC_GUI(root)
    root.mainloop()