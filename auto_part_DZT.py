# автоматическое тестирование ФСУ в части ДЗТ, КЦТ, ПС для исполнения ДЗТ
# с графическим интерфейсом
# 2, 5 заводятся не в процентах а в rms - угол 2,5 гармоник для упрощения совпадает с 1 гармоникой
# для упрощения также 2,5 гармоники только на сторону ВН (НН1, НН2 - нет)

# ИЗМЕНЕНИЯ
# Выходы - нет пуска от ЛО Т - pusk_ptrc1_tprmofflvlgc - удалить столбец из режимов
# SGF params - появилась уставка сборка звезды со стороны Т - добавить и реализовать TDIF_1_Kp1
# базисная мощность задавать в КВА !!!!

import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension

import itertools
import openpyxl
import json
from lib2.PARTS.DZT import partDZT

from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip


class PartDZT_GUI:

    OUTPUT_PARAMS = [
        "CTR_UIRZ_1_RCTR1_FuncEnabled", "CTR_UIRZ_1_RCTR1_FuncOperDisabled", "CTR_UIRZ_1_RCTR1_StrBrkCF", "CTR_UIRZ_1_RCTR1_OpBrkCF", "CTR_UIRZ_1_RCTR1_StrAsymCF", "CTR_UIRZ_1_RCTR1_OpAsymCF", "CTR_UIRZ_1_RCTR2_FuncEnabled", "CTR_UIRZ_1_RCTR2_FuncOperDisabled", "CTR_UIRZ_1_RCTR2_StrBrkCF", "CTR_UIRZ_1_RCTR2_OpBrkCF", "CTR_UIRZ_1_RCTR2_StrAsymCF", "CTR_UIRZ_1_RCTR2_OpAsymCF", "CTR_UIRZ_1_RCTR3_FuncEnabled", "CTR_UIRZ_1_RCTR3_FuncOperDisabled", "CTR_UIRZ_1_RCTR3_StrBrkCF", "CTR_UIRZ_1_RCTR3_OpBrkCF", "CTR_UIRZ_1_RCTR3_StrAsymCF", "CTR_UIRZ_1_RCTR3_OpAsymCF", "CTR_UIRZ_1_CTR_UIRZ_CurCircFlt", "TDIF_1_PDIF2_FuncEnabled", "TDIF_1_PDIF2_FuncOperDisabled", "TDIF_1_PDIF2_StrPhA", "TDIF_1_PDIF2_OpPhA", "TDIF_1_PDIF2_OpPhAOnSignal", "TDIF_1_PDIF2_DE_IA", "TDIF_1_PDIF2_StrPhB", "TDIF_1_PDIF2_OpPhB", "TDIF_1_PDIF2_OpPhBOnSignal", "TDIF_1_PDIF2_DE_IB", "TDIF_1_PDIF2_StrPhC", "TDIF_1_PDIF2_OpPhC", "TDIF_1_PDIF2_OpPhCOnSignal", "TDIF_1_PDIF2_DE_IC", "TDIF_1_PDIF2_Str", "TDIF_1_PDIF2_OpOnSignal", "TDIF_1_PDIF2_Op", "TDIF_1_PDIF1_FuncEnabled", "TDIF_1_PDIF1_FuncOperDisabled", "TDIF_1_PDIF1_StrPhA", "TDIF_1_PDIF1_OpPhA", "TDIF_1_PDIF1_OpPhAOnSignal", "TDIF_1_PDIF1_DE_IA", "TDIF_1_PDIF1_StrPhB", "TDIF_1_PDIF1_OpPhB", "TDIF_1_PDIF1_OpPhBOnSignal", "TDIF_1_PDIF1_DE_IB", "TDIF_1_PDIF1_StrPhC", "TDIF_1_PDIF1_OpPhC", "TDIF_1_PDIF1_OpPhCOnSignal", "TDIF_1_PDIF1_DE_IC", "TDIF_1_PDIF1_Str", "TDIF_1_PDIF1_OpOnSignal", "TDIF_1_PDIF1_Op", "TDIF_1_HF2PHAR1_Str2HIa", "TDIF_1_HF2PHAR1_Str2HIb", "TDIF_1_HF2PHAR1_Str2HIc", "TDIF_1_HF2PHAR1_Str2H", "TDIF_1_HF5PHAR1_Str5HIa", "TDIF_1_HF5PHAR1_Str5HIb", "TDIF_1_HF5PHAR1_Str5HIc", "TDIF_1_HF5PHAR1_Str5H", "TDIF_1_RCTR1_FuncEnabled", "TDIF_1_RCTR1_FuncOperDisabled", "TDIF_1_RCTR1_CurCircAlmPhA", "TDIF_1_RCTR1_CurCircAlmPhB", "TDIF_1_RCTR1_CurCircAlmPhC", "TDIF_1_RCTR1_CurCircAlm", "TDIF_1_RCTR1_CurCircFltGen", "TPRMOFFLVLGC_1_PTRC1_FuncEnabled", "TPRMOFFLVLGC_1_PTRC1_FuncOperDisabled", "pusk_ptrc1_tprmofflvlgc", "TPRMOFFLVLGC_1_PTRC1_Op", "TPRMOFFLVLGC_1_RBRE1_FuncEnabled", "TPRMOFFLVLGC_1_RBRE1_FuncOperDisabled", "TPRMOFFLVLGC_1_RBRE1_BlkOp", "DZT2_LVALH_1_CALH1_Alarm", "diffA", "restA", "diffB", "restB", "diffC", "restC"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ (исполнение ДЗТ2) в части КЦТ, ДЗТ, ЛО Т, ПС, v1.4 от 27.05.26")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")


        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для SGF параметров, настроек, входных и выходных значений
        self.sgf_params = {
            "TDIF_1_Side1": tk.IntVar(value=1),
            "TDIF_1_Side2": tk.IntVar(value=1),
            "TDIF_1_Side3": tk.IntVar(value=1),
            "TDIF_1_KschemeSide1": tk.IntVar(value=0),
            "TDIF_1_KschemeSide2": tk.IntVar(value=0), 
            "TDIF_1_KschemeSide3": tk.IntVar(value=0),           # Третья сторона 03.05.2025
            "TDIF_1_Equaliz3I0s1": tk.IntVar(value=0),
            "TDIF_1_Equaliz3I0s2": tk.IntVar(value=0),
            "TDIF_1_Equaliz3I0s3": tk.IntVar(value=0),   # Третья сторона  03.05.2025         
            "CTR_UIRZ_1_RCTR1_EnaDis": tk.IntVar(value=0),
            "CTR_UIRZ_1_RCTR1_EnaDis_ctbreak": tk.IntVar(value=0),
            "CTR_UIRZ_1_RCTR2_EnaDis": tk.IntVar(value=0),
            "CTR_UIRZ_1_RCTR2_EnaDis_ctbreak": tk.IntVar(value=0),
            "CTR_UIRZ_1_RCTR3_EnaDis": tk.IntVar(value=0),
            "CTR_UIRZ_1_RCTR3_EnaDis_ctbreak": tk.IntVar(value=0),
            "TDIF_1_PDIF1_EnaDis": tk.IntVar(value=0),
            "TDIF_1_PDIF1_RstMod": tk.IntVar(value=1),
            "TDIF_1_PDIF1_EnaDisSelec": tk.IntVar(value=0),
            "TDIF_1_PDIF2_EnaDis": tk.IntVar(value=0),
            "TDIF_1_HF2PHAR1_RegBlock": tk.IntVar(value=0),
            "TDIF_1_HF5PHAR1_RegBlock": tk.IntVar(value=0),
            "TDIF_1_RCTR1_EnaDis": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
        }
        self.settings = {
            "CTR_UIRZ_1_RCTR1_Tbrk": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR1_Tasym": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR1_Inom": tk.DoubleVar(value=5),
            "CTR_UIRZ_1_RCTR1_Imin": tk.DoubleVar(value=0.05),
            "CTR_UIRZ_1_RCTR1_Ksym": tk.DoubleVar(value=0.5),
            "CTR_UIRZ_1_RCTR1_LIsym": tk.DoubleVar(value=0.02),
            "CTR_UIRZ_1_RCTR2_Tbrk": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR2_Tasym": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR2_Inom": tk.DoubleVar(value=5),
            "CTR_UIRZ_1_RCTR2_Imin": tk.DoubleVar(value=0.05),
            "CTR_UIRZ_1_RCTR2_Ksym": tk.DoubleVar(value=0.5),
            "CTR_UIRZ_1_RCTR2_LIsym": tk.DoubleVar(value=0.02),
            "CTR_UIRZ_1_RCTR3_Tbrk": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR3_Tasym": tk.DoubleVar(value=0),
            "CTR_UIRZ_1_RCTR3_Inom": tk.DoubleVar(value=5),
            "CTR_UIRZ_1_RCTR3_Imin": tk.DoubleVar(value=0.05),
            "CTR_UIRZ_1_RCTR3_Ksym": tk.DoubleVar(value=0.5),
            "CTR_UIRZ_1_RCTR3_LIsym": tk.DoubleVar(value=0.02),
            "TDIF_1_Sbase": tk.DoubleVar(value=63000),
            "TDIF_1_Unom1": tk.DoubleVar(value=35),
            "TDIF_1_Unom2": tk.DoubleVar(value=10.5),
            "TDIF_1_Unom3": tk.DoubleVar(value=10.5),  # Третья сторона  03.05.2025 
            "TDIF_1_IprimSide1": tk.DoubleVar(value=750),
            "TDIF_1_IprimSide2": tk.DoubleVar(value=3000),
            "TDIF_1_IprimSide3": tk.DoubleVar(value=3000), # Третья сторона  03.05.2025 
            "TDIF_1_Inomterm1": tk.DoubleVar(value=1),
            "TDIF_1_Inomterm2": tk.DoubleVar(value=1),
            "TDIF_1_Inomterm3": tk.DoubleVar(value=1), # Третья сторона  03.05.2025 
            "TDIF_1_IsecSide1": tk.DoubleVar(value=5),
            "TDIF_1_IsecSide2": tk.DoubleVar(value=5),
            "TDIF_1_IsecSide3": tk.DoubleVar(value=5),    # Третья сторона  03.05.2025
            "TDIF_1_ConnGr1": tk.DoubleVar(value=0),
            "TDIF_1_ConnGr2": tk.DoubleVar(value=6),
            "TDIF_1_ConnGr3": tk.DoubleVar(value=6),    # Третья сторона  03.05.2025
            "TDIF_1_PDIF1_Top": tk.DoubleVar(value=1),
            "TDIF_1_PDIF1_Iop": tk.DoubleVar(value=0.2),
            "TDIF_1_PDIF1_IopCSS": tk.DoubleVar(value=1.2),
            "TDIF_1_PDIF1_Irest1": tk.DoubleVar(value=1),
            "TDIF_1_PDIF1_Irest2": tk.DoubleVar(value=3),
            "TDIF_1_PDIF1_Krest1": tk.DoubleVar(value=0.25),
            "TDIF_1_PDIF1_Krest2": tk.DoubleVar(value=0.7), 
            "TDIF_1_PDIF2_Top": tk.DoubleVar(value=1),
            "TDIF_1_PDIF2_Iop": tk.DoubleVar(value=4),
            "TDIF_1_HF2PHAR1_Tret": tk.DoubleVar(value=0.5),
            "TDIF_1_HF2PHAR1_Tblock": tk.DoubleVar(value=0.5),
            "TDIF_1_HF2PHAR1_K2Hdiv1H": tk.DoubleVar(value=0.2),
            "TDIF_1_HF5PHAR1_Tret": tk.DoubleVar(value=0.5),
            "TDIF_1_HF5PHAR1_Tblock": tk.DoubleVar(value=0.5),
            "TDIF_1_HF5PHAR1_K5Hdiv1H": tk.DoubleVar(value=0.3), 
            "TDIF_1_RCTR1_Top": tk.DoubleVar(value=1),
            "TDIF_1_RCTR1_Iop": tk.DoubleVar(value=0.1), 
        }

        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "DI_CTR_UIRZ": tk.IntVar(value=0),
            "DI_CTR_RCTR1": tk.IntVar(value=0),
            "DI_CTR_RCTR2": tk.IntVar(value=0),
            "DI_CTR_RCTR3": tk.IntVar(value=0),
            "DI_TDIF": tk.IntVar(value=0),
            "DI_RESPDIF": tk.IntVar(value=0),
            "DI_RESPDIF_Sign": tk.IntVar(value=0),
            "DI_INSPDIF": tk.IntVar(value=0),
            "DI_INSPDIF_Sign": tk.IntVar(value=0),
            "DI_ATDIFRCTR": tk.IntVar(value=0),
            "DI_TPRMOFFLVLGC": tk.IntVar(value=0),
            "DI_TJNTPTRC": tk.IntVar(value=0),
            "DI_JNTRBRE": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "dIA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "dIB": tk.DoubleVar(value=240),
            "IC": tk.DoubleVar(value=0),
            "dIC": tk.DoubleVar(value=120),
            "IA1": tk.DoubleVar(value=0),
            "dIA1": tk.DoubleVar(value=0),
            "IB1": tk.DoubleVar(value=0),
            "dIB1": tk.DoubleVar(value=240),
            "IC1": tk.DoubleVar(value=0),
            "dIC1": tk.DoubleVar(value=120),
            "IA2harm": tk.DoubleVar(value=0),
            "IB2harm": tk.DoubleVar(value=0),
            "IC2harm": tk.DoubleVar(value=0),
            "IA5harm": tk.DoubleVar(value=0),
            "IB5harm": tk.DoubleVar(value=0),
            "IC5harm": tk.DoubleVar(value=0),
        }

        self.output_labels = {}


        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc


        # Создание интерфейса
        self.create_widgets()


    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def _get_output_names(self):
        return self.OUTPUT_PARAMS  

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0

        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            if key=="SGF6_ptoc1_lvttoc" or key=="SGF6_ptoc2_lvttoc" or key=="SGF6_ptoc3_lvttoc" or key=="SGF1_ptuv1_lvttoc" or key=="SGF1_ptuv2_lvttoc":
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            elif key=="TDIF_1_PDIF1_RstMod":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[1, 2, 3, 4], state="readonly").grid(row=row, column=col + 1)
            elif key=='Номинальный ток входа':
                ttk.Combobox(sgf_frame, textvariable=var, values=[1, 5], state="readonly").grid(row=row, column=col + 1)                   
            else:
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)

        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)

        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)

        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)

        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)        

        # Button Load JSON
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)

        # Button Save JSON
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")

                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)


            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)

                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)

            row += 1
            if row >= 4:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=25, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)

            row += 1
            if row >= 28:
                row = 0
                col += 2

    def init_part(self):
        self.part = partDZT(
            Side1=self.sgf_params["TDIF_1_Side1"].get(),
            Side2=self.sgf_params["TDIF_1_Side2"].get(),
            Side3=self.sgf_params["TDIF_1_Side3"].get(),                        
            SGF1_rctr1_ctr=self.sgf_params["CTR_UIRZ_1_RCTR1_EnaDis"].get(),
            SGF2_rctr1_ctr=self.sgf_params["CTR_UIRZ_1_RCTR1_EnaDis_ctbreak"].get(),
            SGF1_rctr2_ctr=self.sgf_params["CTR_UIRZ_1_RCTR2_EnaDis"].get(),
            SGF2_rctr2_ctr=self.sgf_params["CTR_UIRZ_1_RCTR2_EnaDis_ctbreak"].get(),
            SGF1_rctr3_ctr=self.sgf_params["CTR_UIRZ_1_RCTR3_EnaDis"].get(),
            SGF2_rctr3_ctr=self.sgf_params["CTR_UIRZ_1_RCTR3_EnaDis_ctbreak"].get(),
            SGF1_pdif1_tdif=self.sgf_params["TDIF_1_PDIF1_EnaDis"].get(),
            SGF2_pdif1_tdif=self.sgf_params["TDIF_1_PDIF1_RstMod"].get(),
            SGF3_pdif1_tdif=self.sgf_params["TDIF_1_PDIF1_EnaDisSelec"].get(),
            SGF1_pdif2_tdif=self.sgf_params["TDIF_1_PDIF2_EnaDis"].get(),
            SGF1_hf2phar1_tdif=self.sgf_params["TDIF_1_HF2PHAR1_RegBlock"].get(),
            SGF1_hf5phar1_tdif=self.sgf_params["TDIF_1_HF5PHAR1_RegBlock"].get(),
            SGF1_rctr1_tdif=self.sgf_params["TDIF_1_RCTR1_EnaDis"].get(),
            SGF1_ptrc1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_RBRE1_EnaDis"].get(),
            k_sch_vn = self.sgf_params["TDIF_1_KschemeSide1"].get(),
            k_sch_nn = self.sgf_params["TDIF_1_KschemeSide2"].get(),
            k_sch_nn2 = self.sgf_params["TDIF_1_KschemeSide3"].get(),                        # Третья сторона
            compens_3i0_vn = self.sgf_params["TDIF_1_Equaliz3I0s1"].get(),
            compens_3i0_nn = self.sgf_params["TDIF_1_Equaliz3I0s2"].get(),
            compens_3i0_nn2 = self.sgf_params["TDIF_1_Equaliz3I0s3"].get(),          # Третья сторона
            T1_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_Tbrk"].get(),
            T2_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_Tasym"].get(),
            Inom_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_Inom"].get(),
            Imin_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_Imin"].get(),
            Ksym_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_Ksym"].get(),
            LIsym_rctr1_ctr = self.settings["CTR_UIRZ_1_RCTR1_LIsym"].get(),

            T1_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_Tbrk"].get(),
            T2_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_Tasym"].get(),
            Inom_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_Inom"].get(),
            Imin_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_Imin"].get(),
            Ksym_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_Ksym"].get(),
            LIsym_rctr2_ctr = self.settings["CTR_UIRZ_1_RCTR2_LIsym"].get(),

            T1_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_Tbrk"].get(),
            T2_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_Tasym"].get(),
            Inom_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_Inom"].get(),
            Imin_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_Imin"].get(),
            Ksym_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_Ksym"].get(),
            LIsym_rctr3_ctr = self.settings["CTR_UIRZ_1_RCTR3_LIsym"].get(),

            Sbaz = self.settings["TDIF_1_Sbase"].get()*1e+3, # В киловольт*амперах!!!!
            Ubaz_vn = self.settings["TDIF_1_Unom1"].get()*1e+3,
            Ubaz_nn = self.settings["TDIF_1_Unom2"].get()*1e+3,
            Ubaz_nn2 = self.settings["TDIF_1_Unom3"].get()*1e+3,            # Третья сторона
            Iperv_vn = self.settings["TDIF_1_IprimSide1"].get(),
            Iperv_nn = self.settings["TDIF_1_IprimSide2"].get(),
            Iperv_nn2 = self.settings["TDIF_1_IprimSide3"].get(),               # Третья сторона          
            Inom_term_vn = self.settings["TDIF_1_Inomterm1"].get(),
            Inom_term_nn = self.settings["TDIF_1_Inomterm2"].get(),
            Inom_term_nn2 = self.settings["TDIF_1_Inomterm3"].get(),        # Третья сторона            
            Ivtor_vn = self.settings["TDIF_1_IsecSide1"].get(),
            Ivtor_nn = self.settings["TDIF_1_IsecSide2"].get(),
            Ivtor_nn2 = self.settings["TDIF_1_IsecSide3"].get(),               # Третья сторона 
            n_sch_vn = self.settings["TDIF_1_ConnGr1"].get(),
            n_sch_nn = self.settings["TDIF_1_ConnGr2"].get(),
            n_sch_nn2 = self.settings["TDIF_1_ConnGr3"].get(),                # Третья сторона 

            T1_pdif1_tdif = self.settings["TDIF_1_PDIF1_Top"].get(),
            Isr_pdif1_tdif = self.settings["TDIF_1_PDIF1_Iop"].get(),
            Isr_zagrub_pdif1_tdif = self.settings["TDIF_1_PDIF1_IopCSS"].get(),
            It1_pdif1_tdif = self.settings["TDIF_1_PDIF1_Irest1"].get(),
            It2_pdif1_tdif = self.settings["TDIF_1_PDIF1_Irest2"].get(),
            Kt1_pdif1_tdif = self.settings["TDIF_1_PDIF1_Krest1"].get(),
            Kt2_pdif1_tdif = self.settings["TDIF_1_PDIF1_Krest2"].get(),

            T1_pdif2_tdif = self.settings["TDIF_1_PDIF2_Top"].get(),
            Iset_pdif2_tdif = self.settings["TDIF_1_PDIF2_Iop"].get(),

            T1_hf2phar1_tdif = self.settings["TDIF_1_HF2PHAR1_Tret"].get(),
            T2_hf2phar1_tdif = self.settings["TDIF_1_HF2PHAR1_Tblock"].get(),
            Ratio_hf2phar1_tdif = self.settings["TDIF_1_HF2PHAR1_K2Hdiv1H"].get(),

            T1_hf5phar1_tdif = self.settings["TDIF_1_HF5PHAR1_Tret"].get(),
            T2_hf5phar1_tdif = self.settings["TDIF_1_HF5PHAR1_Tblock"].get(),
            Ratio_hf5phar1_tdif = self.settings["TDIF_1_HF5PHAR1_K5Hdiv1H"].get(),

            T1_rctr1_tdif = self.settings["TDIF_1_RCTR1_Top"].get(),
            Iset_rctr1_tdif = self.settings["TDIF_1_RCTR1_Iop"].get()
        )

        print("partDZT initialized")

    def start_polling(self):
        if self.part is None:
            print("partDZT not initialized")
            return

        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            #print(inputs)
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                if isinstance(value, (bool, np.bool_)):
                    label.config(text=f"{output}: {int(value)}")
                #label.config(text=f"{output}: {int(value)}")
                #print(f"{output}: {int(value)}")
                else: label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or round(float(value),2)!=0.0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3) # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05) # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):

        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return

        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        #output_file = "data.xlsx"
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)

        # Определяем красный цвет для заливки
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                # Устанавливаем ширину столбца
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20

                # Проверяем значения и применяем форматирование
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)  # Преобразуем значение в число
                        if value != 0:
                            cell.fill = red_fill  # Выделяем красным, если значение не равно 0
                    except (ValueError, TypeError):
                        pass  # Игнорируем ошибки преобразования

        # Применяем форматирование к каждому листу
        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        # Сохраняем изменения
        wb.save(output_file)

        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        # Выбор файла для загрузки
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        try:
            # Чтение данных из Excel
            xls = pd.ExcelFile(file_path)

            # Загрузка SGF Parameters
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            # Загрузка Inputs
            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])                    

            print("Data loaded successfully")

        except Exception as e:
            print(f"Error loading data: {e}")

    # === МЕТОДЫ ДЛЯ РАБОТЫ С JSON ФАЙЛАМИ УСТАВОК ===
    
    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла с суффиксом _SG1."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return
        try:
            handler = SettingsHandler.from_json_file(file_path)
            if not self.meta_handler:
                print("⚠️ Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                    elif type_str == "130":  # Integer
                        try:
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()

                            int_val = int(float(value_clean.replace(',', '.')))
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            value_clean = str(value_str).strip()
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            # --- Обновление T-параметров (settings) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па', 'ква', 'мва']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                        else:
                            self.settings[key].set(float_val)
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            print("✅ Параметры обновлены из JSON (с суффиксом _SG1)")
        except FileNotFoundError:
            print(f"❌ Ошибка: Файл не найден: {file_path}")
        except Exception as e:
            print(f"❌ Ошибка загрузки JSON: {e}")
    
    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()


if __name__ == "__main__":
    root = tk.Tk()
    app = PartDZT_GUI(root)
    root.mainloop()