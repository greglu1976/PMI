import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl

# Импортируем новый класс SWITCH3
from lib._PARTS.SWITCH3 import SWITCH


class PartOfSwitchGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ в части КСВ, КП, КА, УВ, СС, ПС и УРОВ. 31.03.25")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        self.current_mode = 1  # Текущий режим (1 или 2)
        self.mode_timer = None
        self.auto_mode = False  # Режим Авто/Ручной
        self.fixed_outputs_mode1 = {}  # Зафиксированные выходы режима 1
        self.fixed_outputs_mode2 = {}  # Зафиксированные выходы режима 2

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация переменных для параметров SGF, настроек, входных и выходных значений
        self.sgf_params = {
            "SGF1_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF2_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF3_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF4_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF5_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF6_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF7_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF8_rcbf1_lvcbsup": tk.IntVar(value=0),
            "SGF1_swctrl": tk.IntVar(value=0),
            "SGF1_cbcswi1_swctrl": tk.IntVar(value=0),
            "SGF2_cbcswi1_swctrl": tk.IntVar(value=0),
            "SGF1_cbcswi1_hvbctrl": tk.IntVar(value=0),
            "SGF1_tsd": tk.IntVar(value=0),
            "SGF1_xcbr1_tsd": tk.IntVar(value=0),
            "SGF2_xcbr1_tsd": tk.IntVar(value=0),
            "SGF3_xcbr1_tsd": tk.IntVar(value=0),
            "SGF4_xcbr1_tsd": tk.IntVar(value=0),
            "SGF5_xcbr1_tsd": tk.IntVar(value=0),
            "SGF6_xcbr1_tsd": tk.IntVar(value=0),
            "SGF1_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF2_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF3_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF4_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF5_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF6_rbrf1_tpbrf": tk.IntVar(value=0),
            "SGF1_hvcbptrc1_hvtcboff": tk.IntVar(value=0),
            "SGF13_tsa": tk.IntVar(value=0),
        }
        self.settings = {
            "T1_rcbf1_lvcbsup": tk.DoubleVar(value=1),
            "T2_rcbf1_lvcbsup": tk.DoubleVar(value=1),
            "T3_rcbf1_lvcbsup": tk.DoubleVar(value=1),
            "T1_cbcswi1_swctrl": tk.DoubleVar(value=1),
            "T2_cbcswi1_swctrl": tk.DoubleVar(value=1),
            "T3_cbcswi1_swctrl": tk.DoubleVar(value=1),
            "T4_cbcswi1_swctrl": tk.DoubleVar(value=1),
            "T1_cbcswi1_hvbctrl": tk.DoubleVar(value=1),
            "T1_xcbr1_tsd": tk.DoubleVar(value=1),
            "T2_xcbr1_tsd": tk.DoubleVar(value=1),
            "T3_xcbr1_tsd": tk.DoubleVar(value=1),
            "T4_xcbr1_tsd": tk.DoubleVar(value=1),
            "T1_rbrf1_tpbrf": tk.DoubleVar(value=1),
            "Iset_rbrf1_tpbrf": tk.DoubleVar(value=0.2),
            "T1_hvcbptrc1_hvtcboff": tk.DoubleVar(value=1),            
        }

        # Входные переменные для режима 1
        self.input_vars_mode1 = {
            "VYVOD": tk.IntVar(value=0),
            "OV_rcbf1_lvcbsup": tk.IntVar(value=0),
            "ot_emo1emv": tk.IntVar(value=0),
            "ot_emo2": tk.IntVar(value=0),
            "avar_isol_V": tk.IntVar(value=0),
            "niz_isol_V": tk.IntVar(value=0),
            "pruzh_ne_zaved": tk.IntVar(value=0),
            "Sbros": tk.IntVar(value=0),
            "otkl_ot_knopk": tk.IntVar(value=0),
            "oper_otkl_V": tk.IntVar(value=0),
            "KRV_resurs_V": tk.IntVar(value=0),
            "vnesh_blok_upr_V": tk.IntVar(value=0),
            "kontr_emv": tk.IntVar(value=0),
            "kontr_emo1": tk.IntVar(value=0),
            "kontr_emo2": tk.IntVar(value=0),
            "rabota_emv": tk.IntVar(value=0),
            "rabota_emo1": tk.IntVar(value=0),
            "rabota_emo2": tk.IntVar(value=0),
            "OV_swctrl": tk.IntVar(value=0),
            "otkl_v_ot_pu": tk.IntVar(value=0),
            "otkl_v_ichm": tk.IntVar(value=0),
            "mestnoe": tk.IntVar(value=0),
            "otkl_v_ot_tu": tk.IntVar(value=0),
            "otkl_v_asu": tk.IntVar(value=0),
            "kluch_md_priv": tk.IntVar(value=0),
            "vkl_v_ot_pu": tk.IntVar(value=0),
            "vkl_v_ichm": tk.IntVar(value=0),
            "distanz": tk.IntVar(value=0),
            "vkl_v_ot_tu": tk.IntVar(value=0),
            "vkl_v_asu": tk.IntVar(value=0),
            "v_otkl_bk": tk.IntVar(value=0),
            "v_vkl_bk": tk.IntVar(value=0),
            "OV_cbcswi1_hvbctrl": tk.IntVar(value=0),
            "oper_vkl_v": tk.IntVar(value=0),
            "OV_tsd": tk.IntVar(value=0),
            "pusk_urov_vnesh": tk.IntVar(value=0),
            "vnesh_otkl_zdz": tk.IntVar(value=0),
            "vnesh_otkl_urov": tk.IntVar(value=0),
        }

        # Входные переменные для режима 2
        self.input_vars_mode2 = {
            "VYVOD": tk.IntVar(value=0),
            "OV_rcbf1_lvcbsup": tk.IntVar(value=0),
            "ot_emo1emv": tk.IntVar(value=0),
            "ot_emo2": tk.IntVar(value=0),
            "avar_isol_V": tk.IntVar(value=0),
            "niz_isol_V": tk.IntVar(value=0),
            "pruzh_ne_zaved": tk.IntVar(value=0),
            "Sbros": tk.IntVar(value=0),
            "otkl_ot_knopk": tk.IntVar(value=0),
            "oper_otkl_V": tk.IntVar(value=0),
            "KRV_resurs_V": tk.IntVar(value=0),
            "vnesh_blok_upr_V": tk.IntVar(value=0),
            "kontr_emv": tk.IntVar(value=0),
            "kontr_emo1": tk.IntVar(value=0),
            "kontr_emo2": tk.IntVar(value=0),
            "rabota_emv": tk.IntVar(value=0),
            "rabota_emo1": tk.IntVar(value=0),
            "rabota_emo2": tk.IntVar(value=0),
            "OV_swctrl": tk.IntVar(value=0),
            "otkl_v_ot_pu": tk.IntVar(value=0),
            "otkl_v_ichm": tk.IntVar(value=0),
            "mestnoe": tk.IntVar(value=0),
            "otkl_v_ot_tu": tk.IntVar(value=0),
            "otkl_v_asu": tk.IntVar(value=0),
            "kluch_md_priv": tk.IntVar(value=0),
            "vkl_v_ot_pu": tk.IntVar(value=0),
            "vkl_v_ichm": tk.IntVar(value=0),
            "distanz": tk.IntVar(value=0),
            "vkl_v_ot_tu": tk.IntVar(value=0),
            "vkl_v_asu": tk.IntVar(value=0),
            "v_otkl_bk": tk.IntVar(value=0),
            "v_vkl_bk": tk.IntVar(value=0),
            "OV_cbcswi1_hvbctrl": tk.IntVar(value=0),
            "oper_vkl_v": tk.IntVar(value=0),
            "OV_tsd": tk.IntVar(value=0),
            "pusk_urov_vnesh": tk.IntVar(value=0),
            "vnesh_otkl_zdz": tk.IntVar(value=0),
            "vnesh_otkl_urov": tk.IntVar(value=0),
        }

        self.output_labels = {}

        # Создание интерфейса
        self.create_widgets()

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 10:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Переключатель Авто/Ручной
        self.auto_manual_var = tk.StringVar(value="Ручной")
        ttk.Button(buttons_frame, text="Переключить режим", command=self.toggle_auto_manual).grid(row=0, column=0, pady=10)

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=1, pady=10)

        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=2, pady=10)

        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=3, pady=10)

        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=4, pady=10)

        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=5, pady=10)

        # Button Fix Outputs
        ttk.Button(buttons_frame, text="Фиксация выходов", command=self.fix_outputs).grid(row=0, column=6, pady=10)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Кнопка для переключения режимов
        self.mode_button = ttk.Button(buttons_frame, text="Режим 1", command=self.toggle_mode)
        self.mode_button.grid(row=0, column=11, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=12, padx=5, pady=5)

        # Frame for input values (режим 1)
        self.input_frame_mode1 = ttk.LabelFrame(self.root, text="Inputs Mode 1")
        self.input_frame_mode1.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.input_vars_mode1.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(self.input_frame_mode1, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(self.input_frame_mode1, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(self.input_frame_mode1, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Frame for input values (режим 2) - изначально скрыт
        self.input_frame_mode2 = ttk.LabelFrame(self.root, text="Inputs Mode 2")
        row = 0
        col = 0
        for key, var in self.input_vars_mode2.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(self.input_frame_mode2, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(self.input_frame_mode2, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(self.input_frame_mode2, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")
        outputs = [
            "vvod_rcbf1_lvcbsup", "oper_vyvod_rcbf1_lvcbsup", "v_samoproisv_otkl_rcbf1_lvcbsup", "neispr_V_rcbf1_lvcbsup",
            "v_avar_otkl_rcbf1_lvcbsup", "rfk_rcbf1_lvcbsup", "blok_vkl_rcbf1_lvcbsup", "blok_otkl_rcbf1_lvcbsup",
            "neisp_emu_rcbf1_lvcbsup", "zashita_emv_rcbf1_lvcbsup", "zashita_emo1_rcbf1_lvcbsup", "zashita_emo2_rcbf1_lvcbsup",
            "vvod_swctrl", "oper_vyvod_swctrl", "vvod_cbcswi1_swctrl", "uv_otkluchit_cbcswi1_swctrl", "uv_idet_per_cbcswi1_swctrl",
            "uv_prev_vrem_per_cbcswi1_swctrl", "uv_vkluchit_cbcswi1_swctrl", "uv_ne_opredeleno_cbcswi1_swctrl",
            "uv_otklucheno_cbcswi1_swctrl", "uv_vklucheno_cbcswi1_swctrl", "uv_neispr_neopred_cbcswi1_swctrl",
            "vvod_cbcswi1_hvbctrl", "oper_vyvod_cbcswi1_hvbctrl", "uv_vkl_cbcswi1_hvbctrl",
            "vvod_tsd", "oper_vyvod_tsd", "vvod_xcbr1_tsd", "v_prom_pol_xcbr1_tsd", "v_otkluchen_xcbr1_tsd",
            "v_vkluchen_xcbr1_tsd", "v_neisp_pol_xcbr1_tsd", "v_otkluchit_rele_xcbr1_tsd", "v_vkluchit_rele_xcbr1_tsd", "ss_prev_vrem_per_ka", "pusk_t_lvalh",
            "srab_na_sebya_rbrf1_tpbrf",
            "vvod_hvcbptrc1_hvtcboff", "otkl_hvcbptrc1_hvtcboff", "otkl_avar_hvcbptrc1_hvtcboff"
        ]
        row = 0
        col = 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=35, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 32:
                row = 0
                col += 2

    def toggle_auto_manual(self):
        """Переключение между режимами Авто и Ручной."""
        if self.auto_manual_var.get() == "Ручной":
            self.auto_manual_var.set("Авто")
            print("Переключено в режим Авто")
        else:
            self.auto_manual_var.set("Ручной")
            print("Переключено в режим Ручной")

    def toggle_mode(self):
        """Переключение между режимами 1 и 2."""
        if self.current_mode == 1:
            self.current_mode = 2
            self.mode_button.config(text="Режим 2")
            self.input_frame_mode1.grid_remove()
            self.input_frame_mode2.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        else:
            self.current_mode = 1
            self.mode_button.config(text="Режим 1")
            self.input_frame_mode2.grid_remove()
            self.input_frame_mode1.grid(row=3, column=0, padx=10, pady=10, sticky="w")

    def init_part(self):
        self.part = SWITCH(
            SGF1_rcbf1_lvcbsup=self.sgf_params["SGF1_rcbf1_lvcbsup"].get(),
            SGF2_rcbf1_lvcbsup=self.sgf_params["SGF2_rcbf1_lvcbsup"].get(),
            SGF3_rcbf1_lvcbsup=self.sgf_params["SGF3_rcbf1_lvcbsup"].get(),
            SGF4_rcbf1_lvcbsup=self.sgf_params["SGF4_rcbf1_lvcbsup"].get(),
            SGF5_rcbf1_lvcbsup=self.sgf_params["SGF5_rcbf1_lvcbsup"].get(),
            SGF6_rcbf1_lvcbsup=self.sgf_params["SGF6_rcbf1_lvcbsup"].get(),
            SGF7_rcbf1_lvcbsup=self.sgf_params["SGF7_rcbf1_lvcbsup"].get(),
            SGF8_rcbf1_lvcbsup=self.sgf_params["SGF8_rcbf1_lvcbsup"].get(),
            T1_rcbf1_lvcbsup=self.settings["T1_rcbf1_lvcbsup"].get(),
            T2_rcbf1_lvcbsup=self.settings["T2_rcbf1_lvcbsup"].get(),
            T3_rcbf1_lvcbsup=self.settings["T3_rcbf1_lvcbsup"].get(),
            SGF1_swctrl=self.sgf_params["SGF1_swctrl"].get(),
            SGF1_cbcswi1_swctrl=self.sgf_params["SGF1_cbcswi1_swctrl"].get(),
            SGF2_cbcswi1_swctrl=self.sgf_params["SGF2_cbcswi1_swctrl"].get(),
            T1_cbcswi1_swctrl=self.settings["T1_cbcswi1_swctrl"].get(),
            T2_cbcswi1_swctrl=self.settings["T2_cbcswi1_swctrl"].get(),
            T3_cbcswi1_swctrl=self.settings["T3_cbcswi1_swctrl"].get(),
            T4_cbcswi1_swctr=self.settings["T4_cbcswi1_swctrl"].get(),
            SGF1_cbcswi1_hvbctrl=self.sgf_params["SGF1_cbcswi1_hvbctrl"].get(),
            T1_cbcswi1_hvbctrl=self.settings["T1_cbcswi1_hvbctrl"].get(),
            SGF1_tsd=self.sgf_params["SGF1_tsd"].get(),
            SGF1_xcbr1_tsd=self.sgf_params["SGF1_xcbr1_tsd"].get(),
            SGF2_xcbr1_tsd=self.sgf_params["SGF2_xcbr1_tsd"].get(),
            SGF3_xcbr1_tsd=self.sgf_params["SGF3_xcbr1_tsd"].get(),
            SGF4_xcbr1_tsd=self.sgf_params["SGF4_xcbr1_tsd"].get(),
            SGF5_xcbr1_tsd=self.sgf_params["SGF5_xcbr1_tsd"].get(),
            SGF6_xcbr1_tsd=self.sgf_params["SGF6_xcbr1_tsd"].get(),
            T1_xcbr1_tsd=self.settings["T1_xcbr1_tsd"].get(),
            T2_xcbr1_tsd=self.settings["T2_xcbr1_tsd"].get(),
            T3_xcbr1_tsd=self.settings["T3_xcbr1_tsd"].get(),
            T4_xcbr1_tsd=self.settings["T4_xcbr1_tsd"].get(),
            SGF13_tsa=self.sgf_params["SGF13_tsa"].get(),
            SGF1_rbrf1_tpbrf=self.sgf_params["SGF1_rbrf1_tpbrf"].get(),
            SGF2_rbrf1_tpbrf=self.sgf_params["SGF2_rbrf1_tpbrf"].get(),
            SGF3_rbrf1_tpbrf=self.sgf_params["SGF3_rbrf1_tpbrf"].get(),
            SGF4_rbrf1_tpbrf=self.sgf_params["SGF4_rbrf1_tpbrf"].get(),
            SGF5_rbrf1_tpbrf=self.sgf_params["SGF5_rbrf1_tpbrf"].get(),
            SGF6_rbrf1_tpbrf=self.sgf_params["SGF6_rbrf1_tpbrf"].get(),
            T1_rbrf1_tpbrf=self.settings["T1_rbrf1_tpbrf"].get(),
            Iset_rbrf1_tpbrf=self.settings["Iset_rbrf1_tpbrf"].get(),
            SGF1_hvcbptrc1_hvtcboff=self.sgf_params["SGF1_hvcbptrc1_hvtcboff"].get(),
            T1_hvcbptrc1_hvtcboff=self.settings["T1_hvcbptrc1_hvtcboff"].get(),
        )
        print("part_SWITCH initialized")

    def start_polling(self):
        if self.part is None:
            print("part_SWITCH not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

        if self.auto_manual_var.get() == "Авто":
            self.start_auto_mode()

    def start_auto_mode(self):
        """Запуск автоматического режима."""
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.auto_mode_logic, daemon=True)
        self.polling_thread.start()

    def auto_mode_logic(self):
        """Логика работы в режиме Авто."""
        if self.is_polling:
            # Режим 1
            self.current_mode = 1
            self.mode_button.config(text="Режим 1")
            self.input_frame_mode2.grid_remove()
            self.input_frame_mode1.grid(row=3, column=0, padx=10, pady=10, sticky="w")
            time.sleep(2.5)  # Ждем 2.5 секунды
            self.fix_outputs_mode1()  # Фиксируем выходы режима 1
            time.sleep(0.5)  # Оставшееся время режима 1

            # Режим 2
            self.current_mode = 2
            self.mode_button.config(text="Режим 2")
            self.input_frame_mode1.grid_remove()
            self.input_frame_mode2.grid(row=3, column=0, padx=10, pady=10, sticky="w")
            time.sleep(4.5)  # Ждем 4.5 секунды
            self.fix_outputs_mode2()  # Фиксируем выходы режима 2
            time.sleep(0.5)  # Оставшееся время режима 2

            self.stop_polling()  # Останавливаем вычисления

    def fix_outputs_mode1(self):
        """Фиксация выходов режима 1."""
        self.fixed_outputs_mode1 = {key: label.cget("text").split(": ")[-1] for key, label in self.output_labels.items()}
        print("Выходы режима 1 зафиксированы")

    def fix_outputs_mode2(self):
        """Фиксация выходов режима 2."""
        self.fixed_outputs_mode2 = {key: label.cget("text").split(": ")[-1] for key, label in self.output_labels.items()}
        print("Выходы режима 2 зафиксированы")

    def fix_outputs(self):
        """Фиксация выходов в ручном режиме."""
        if self.current_mode == 1:
            self.fix_outputs_mode1()
        else:
            self.fix_outputs_mode2()

    def stop_polling(self):
        """Останавливает опрос входных данных."""
        self.is_polling = False  # Устанавливаем флаг для остановки цикла
        if self.mode_timer:
            self.root.after_cancel(self.mode_timer)  # Отменяем таймер
            self.mode_timer = None
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            # Выбираем текущий набор входных переменных в зависимости от режима
            if self.current_mode == 1:
                inputs = {key: var.get() for key, var in self.input_vars_mode1.items()}
            else:
                inputs = {key: var.get() for key, var in self.input_vars_mode2.items()}
            result = self.part.Step(**inputs)
            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")
            time.sleep(0.3)  # Время шага опроса

    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })

        # Сохраняем входные данные для обоих режимов
        inputs_mode1_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars_mode1.items()
        })
        inputs_mode2_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars_mode2.items()
        })

        # Сохраняем выходы
        outputs_mode1_df = pd.DataFrame([self.fixed_outputs_mode1])
        outputs_mode2_df = pd.DataFrame([self.fixed_outputs_mode2])

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_mode1_df.to_excel(writer, sheet_name="Inputs_Mode1", index=False)
            inputs_mode2_df.to_excel(writer, sheet_name="Inputs_Mode2", index=False)
            outputs_mode1_df.to_excel(writer, sheet_name="Fixed_Outputs", startrow=0, index=False)
            outputs_mode2_df.to_excel(writer, sheet_name="Fixed_Outputs", startrow=len(outputs_mode1_df) + 2, index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs_Mode1"], inputs_mode1_df)
        format_sheet(wb["Inputs_Mode2"], inputs_mode2_df)
        format_sheet(wb["Fixed_Outputs"], pd.concat([outputs_mode1_df, outputs_mode2_df]))

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            # Загрузка параметров SGF
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])
            # Загрузка настроек
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])
            # Загрузка входных данных для режима 1
            if "Inputs_Mode1" in xls.sheet_names:
                inputs_mode1_df = pd.read_excel(xls, sheet_name="Inputs_Mode1")
                for key, var in self.input_vars_mode1.items():
                    if key in inputs_mode1_df.columns:
                        var.set(inputs_mode1_df.at[0, key])
            # Загрузка входных данных для режима 2
            if "Inputs_Mode2" in xls.sheet_names:
                inputs_mode2_df = pd.read_excel(xls, sheet_name="Inputs_Mode2")
                for key, var in self.input_vars_mode2.items():
                    if key in inputs_mode2_df.columns:
                        var.set(inputs_mode2_df.at[0, key])
            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfSwitchGUI(root)
    root.mainloop()