import tkinter as tk
from tkinter import ttk, messagebox
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension
import openpyxl
import json
# Импортируем класс partKRV
from lib2.PARTS.KRV_T2 import partKRV
from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class partKRVGUI:

    # === СПИСОК ВЫХОДНЫХ ПАРАМЕТРОВ (единое определение) ===
    OUTPUT_PARAMS = [
            # Выходы КСВ
            "T_LVCBSUP_1_RCBF1_UnpromptedCBopening", "T_LVCBSUP_1_RCBF1_FailureCB", "T_LVCBSUP_1_RCBF1_CBFailureTrip", 
            "T_LVCBSUP_1_RCBF1_BlkToCls", "T_LVCBSUP_1_RCBF1_BlkToOpn",
            # Выходы КА
            "T_SwitchDevice_1_CB1_CBPosIntermed", "T_SwitchDevice_1_CB1_CBPosOpn", "T_SwitchDevice_1_CB1_CBPosCls", 
            "T_SwitchDevice_1_CB1_OpnCB_relay", "T_SwitchDevice_1_CB1_ClsCB_relay",
            # Выходы ПС
            "T2_LVALH_1_CALH1_Alarm",
            # Выходы КРВ
            "CLS_1_CLS_FuncEnabled", "CLS_1_CLS_MDResourceExcess", "CLS_1_CLS_CBLifeExcess", 
            "CLS_1_CLS_COMMResourceExcess", "CLS_1_CLS_MDCurrentResource", 
            "CLS_1_CLS_COMMCurrResourcePhsA", "CLS_1_CLS_COMMCurrResourcePhsB", "CLS_1_CLS_COMMCurrResourcePhsC"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование partKRV (Контроль Ресурса Выключателя) T2, v1.0")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        
        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")
        
        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для параметров SGF и настроек
        self.sgf_params = {
            "CLS_1_CLS_EnaDis": tk.IntVar(value=1),
            "CLS_1_CLS_MDCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_SD_EnaDis": tk.IntVar(value=1),
            "T_SwitchDevice_1_CB1_EnaDis": tk.IntVar(value=1),
            "T_SwitchDevice_1_CB1_TPOpnResetCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_CBOSoperationCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_TPClsResetCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_CBCSoperationCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_EnaDis": tk.IntVar(value=1),
            "T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_RstFrmCLS": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_KnobCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T2_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }
        
        self.settings = {
            "CLS_1_CLS_CBpasp_Inom": tk.DoubleVar(value=1),
            "CLS_1_CLS_CBpasp_InomBr": tk.DoubleVar(value=3),
            "CLS_1_CLS_COMMpasp_Inom": tk.DoubleVar(value=7),
            "CLS_1_CLS_COMMpasp_InomBr": tk.DoubleVar(value=5),
            "CLS_1_CLS_CBpasp_MD": tk.IntVar(value=3),
            "CLS_1_CLS_InitialCOMM": tk.DoubleVar(value=100),
            "CLS_1_CLS_COMMSet": tk.DoubleVar(value=20),
            "CLS_1_CLS_InitialMD": tk.IntVar(value=0),
            "CLS_1_CLS_TmaxCB": tk.IntVar(value=500),
            "T_SwitchDevice_1_CB1_TonFaul": tk.DoubleVar(value=3000),
            "T_SwitchDevice_1_CB1_OpnTPtime": tk.DoubleVar(value=1000),
            "T_SwitchDevice_1_CB1_ClsTPtime": tk.DoubleVar(value=1000),
            "T_SwitchDevice_1_CB1_TextenCls": tk.DoubleVar(value=1000),
            "T_LVCBSUP_1_RCBF1_T_EnBlk": tk.DoubleVar(value=500),
            "T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg": tk.DoubleVar(value=1000),
            "T_LVCBSUP_1_RCBF1_T_ElmgWorking": tk.DoubleVar(value=1000),
        }
        
        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "CBPosOpn": tk.IntVar(value=1),
            "CBPosCls": tk.IntVar(value=0),
            "CLS_1_ResetCounter": tk.IntVar(value=0),
            "LocKey": tk.IntVar(value=0),
            #"CLS_1_OpnCB": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),
            "Reset": tk.IntVar(value=0),
            "OperOpnCB": tk.IntVar(value=0),
        }
        self.output_labels = {}

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc
        
        # Создание интерфейса
        self.create_widgets()
    
    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())
    
    def _get_output_names(self):
        return self.OUTPUT_PARAMS
    
    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 15:
                row = 0
                col += 2
        
        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 10:
                row = 0
                col += 2
        
        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        
        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)
        
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)
        
        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)
        
        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)
                
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2
        
        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=50, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)

            row += 1
            if row >= 16:
                row = 0
                col += 2
    
    def init_part(self):
        self.part = partKRV(
            SGF1=self.sgf_params["CLS_1_CLS_EnaDis"].get(),
            Inom_V_pasp=self.settings["CLS_1_CLS_CBpasp_Inom"].get(),
            Inom_otl_V_pasp=self.settings["CLS_1_CLS_CBpasp_InomBr"].get(),
            KRVpasp_Inom=self.settings["CLS_1_CLS_COMMpasp_Inom"].get(),
            KRVpasp_Inom_otkl=self.settings["CLS_1_CLS_COMMpasp_InomBr"].get(),
            MRVpasp=self.settings["CLS_1_CLS_CBpasp_MD"].get(),
            Nach_znach_KRV=self.settings["CLS_1_CLS_InitialCOMM"].get(),
            KRVsrab=self.settings["CLS_1_CLS_COMMSet"].get(),
            Nach_znach_MRV=self.settings["CLS_1_CLS_InitialMD"].get(),
            SGF2=self.sgf_params["CLS_1_CLS_MDCtrl"].get(),
            T1=self.settings["CLS_1_CLS_TmaxCB"].get(),
            SGF1_tsd=self.sgf_params["T_SwitchDevice_1_SD_EnaDis"].get(),
            SGF1_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_EnaDis"].get(),
            SGF2_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_TPOpnResetCtrl"].get(),
            SGF3_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_CBOSoperationCtrl"].get(),
            SGF4_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_TPClsResetCtrl"].get(),
            SGF5_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_CBCSoperationCtrl"].get(),
            T1_xcbr1_tsd=self.settings["T_SwitchDevice_1_CB1_TonFaul"].get(),
            T2_xcbr1_tsd=self.settings["T_SwitchDevice_1_CB1_OpnTPtime"].get(),
            T3_xcbr1_tsd=self.settings["T_SwitchDevice_1_CB1_ClsTPtime"].get(),
            T4_xcbr1_tsd=self.settings["T_SwitchDevice_1_CB1_TextenCls"].get(),
            SGF1_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_EnaDis"].get(),
            SGF2_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol"].get(),
            SGF3_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault"].get(),
            SGF4_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_RstFrmCLS"].get(),
            SGF5_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl"].get(),
            SGF6_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault"].get(),
            SGF7_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_KnobCtrl"].get(),
            SGF8_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm"].get(),
            T1_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),
            T2_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(),
            T3_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),
            SGF1_lvalh=self.sgf_params["T2_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["T2_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["T2_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["T2_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["T2_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["T2_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["T2_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["T2_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["T2_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["T2_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
        )
        print("✅ partKRV initialized")
    
    def start_polling(self):
        if self.part is None:
            print("⚠️ partKRV not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()
    
    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("🛑 Polling stopped")
    
    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)
            
            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                if isinstance(value, float):
                    label.config(text=f"{output}: {round(value, 2)}")
                else:
                    label.config(text=f"{output}: {value}")
                if value != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")
            
            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")
    
    def save_to_excel(self):
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        sgf_df = pd.DataFrame({key: [var.get()] for key, var in self.sgf_params.items()})
        settings_df = pd.DataFrame({key: [var.get()] for key, var in self.settings.items()})
        inputs_df = pd.DataFrame({key: [var.get()] for key, var in self.input_vars.items()})
        outputs_df = pd.DataFrame({key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()})

        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            for col_num, column in enumerate(sheet.columns, start=1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:
                        continue
                    try:
                        if float(cell.value) != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        for sheet_name in ["SGF_Parameters", "Settings", "Inputs", "Outputs"]:
            format_sheet(wb[sheet_name], None)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")
    
    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name, var_dict in [("SGF_Parameters", self.sgf_params),
                                          ("Settings", self.settings),
                                          ("Inputs", self.input_vars)]:
                if sheet_name in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    for key, var in var_dict.items():
                        if key in df.columns:
                            var.set(df.at[0, key])
            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")




    
    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла с суффиксом _SG1."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return
        try:
            handler = SettingsHandler.from_json_file(file_path)
            if not self.meta_handler:
                print("⚠️ Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                    elif type_str == "130":  # Integer
                        try:
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()

                            int_val = int(float(value_clean.replace(',', '.')))
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            value_clean = str(value_str).strip()
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            # --- Обновление T-параметров (settings) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                        else:
                            self.settings[key].set(float_val)
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            print("✅ Параметры обновлены из JSON (с суффиксом _SG1)")
        except FileNotFoundError:
            print(f"❌ Ошибка: Файл не найден: {file_path}")
        except Exception as e:
            print(f"❌ Ошибка загрузки JSON: {e}")
    
    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()




if __name__ == "__main__":
    root = tk.Tk()
    app = partKRVGUI(root)
    root.mainloop()