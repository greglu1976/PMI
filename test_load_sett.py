########################
### Версия для 2026 года
########################

import tkinter as tk
from tkinter import ttk, messagebox
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl

import json
from lib._PARTS.TECH_T import part_TECH_T 

from SettingsHandler import SettingsHandler

class PartOfTECH_T_GUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование Газовых и Технологических Защит М300-Т. вер1 от 2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация SGF-параметров
        self.sgf_params = {
            "APTTECHLGC_1_OILPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "APTTECHLGC_1_WINPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "APTTECHLGC_1_VLVPTRC1_EnaDis": tk.IntVar(value=0),
            "APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis": tk.IntVar(value=0),
            "ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis": tk.IntVar(value=0),
            "TALMGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TALMGASLGC_1_PTRC1_LowIsolSignCtrl": tk.IntVar(value=0),
            "TTRGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TTRGASLGC_1_PTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "TLTCGASLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TLTCGASLGC_1_PTRC1_LowIsolTripCtrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC3_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA1": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA3": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA4": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA5": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG1": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_GAS_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_TECH_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_OCcir_CB": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_ARCnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_CBFPnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_IEDvt_OCControl": tk.IntVar(value=0),
        }

        # Настройки (T-параметры)
        self.settings = {
            #"T1_oilptrc1_apttechlgc": tk.DoubleVar(value=1),
            #"T1_winptrc1_apttechlgc": tk.DoubleVar(value=1),
            #"T1_vlvptrc1_apttechlgc": tk.DoubleVar(value=1),
            "APTTECHLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),            
            "TALMGASLGC_1_PTRC1_TopOnBlk": tk.DoubleVar(value=1),
            "TTRGASLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),
            "TLTCGASLGC_1_LLN0_TopOnBlk": tk.DoubleVar(value=1),
        }

        # Входные параметры для Step()
        self.input_vars = {
            "VYVOD": tk.IntVar(value=0),
            "OV_tz": tk.IntVar(value=0),
            "OV_dtm": tk.IntVar(value=0),
            "OV_dto": tk.IntVar(value=0),
            "OV_rd": tk.IntVar(value=0),
            "NaSign_dtm": tk.IntVar(value=0),
            "NaSign_dto": tk.IntVar(value=0),
            "NaSign_rd": tk.IntVar(value=0),
            "srabKontOtkl_m": tk.IntVar(value=0),
            "srabKontSign_m": tk.IntVar(value=0),
            "srabKontOtkl_o": tk.IntVar(value=0),
            "srabKontSign_o": tk.IntVar(value=0),
            "srabKontOtkl_rd": tk.IntVar(value=0),
            "srabKI_m": tk.IntVar(value=0),
            "srabKI_o": tk.IntVar(value=0),
            "srabKI_rd": tk.IntVar(value=0),
            "Sbros": tk.IntVar(value=0),
            "OV_ts": tk.IntVar(value=0),
            "OV_pk": tk.IntVar(value=0),
            "OV_ok": tk.IntVar(value=0),
            "OV_lev": tk.IntVar(value=0),
            "NaSign_pk": tk.IntVar(value=0),
            "NaSign_ok": tk.IntVar(value=0),
            "NaSign_lev": tk.IntVar(value=0),
            "srabKontOtkl_pk": tk.IntVar(value=0),
            "srabKontOtkl_ok": tk.IntVar(value=0),
            "srabKontOtkl_lev": tk.IntVar(value=0),
            "OV_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "NaOtkl_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_talmgaslgc": tk.IntVar(value=0),
            "OV_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "NaSign_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_ttrgaslgc": tk.IntVar(value=0),
            "OV_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "NaSign_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "srabKont_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "srabKI_ptrc1_tltcgaslgc": tk.IntVar(value=0),
            "OV_tofflvlg": tk.IntVar(value=0),
            "OVlo_tofflvlg": tk.IntVar(value=0),
            "OVzapv_tofflvlg": tk.IntVar(value=0),
            "OVzavr_tofflvlg": tk.IntVar(value=0),
            "oil_t_hi_level": tk.IntVar(value=0),
            "oil_ltc_hi_level": tk.IntVar(value=0),
            "oil_ltc_lo_level": tk.IntVar(value=0),
            "oil_ltc_lo_temp": tk.IntVar(value=0),
        }

        self.output_labels = {}  # Для вывода результатов
        self.create_widgets()

    def create_widgets(self):
        # Фрейм для SGF-параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 14:
                row = 0
                col += 2

        # Фрейм для настроек (T-параметры)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 3:
                row = 0
                col += 2

        # Фрейм для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)

        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Фрейм для входных параметров
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(input_frame, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(input_frame, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Фрейм для выходных параметров
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")
        outputs = [
            "vvod_oilptrc1_apttechlgc", "oper_vyvod_oilptrc1_apttechlgc", "srab_oilptrc1_apttechlgc", "srabsign_oilptrc1_apttechlgc",
            "zablok_oilptrc1_apttechlgc", "ET_oilptrc1_apttechlgc", "vvod_winptrc1_apttechlgc", "oper_vyvod_winptrc1_apttechlgc",
            "srab_winptrc1_apttechlgc", "srabsign_winptrc1_apttechlgc", "zablok_winptrc1_apttechlgc", "ET_winptrc1_apttechlgc",
            "vvod_vlvptrc1_apttechlgc", "oper_vyvod_vlvptrc1_apttechlgc", "srab_vlvptrc1_apttechlgc", "srabsign_vlvptrc1_apttechlgc",
            "zablok_vlvptrc1_apttechlgc", "ET_vlvptrc1_apttechlgc", "vvod_prvlvptrc1_almtechlgc", "oper_vyvod_prvlvptrc1_almtechlgc",
            "srab_prvlvptrc1_almtechlgc", "srabsign_prvlvptrc1_almtechlgc", "vvod_shvlvptrc1_almtechlgc", "oper_vyvod_shvlvptrc1_almtechlgc",
            "srab_shvlvptrc1_almtechlgc", "srabsign_shvlvptrc1_almtechlgc", "vvod_levptrc1_almtechlgc", "oper_vyvod_levptrc1_almtechlgc",
            "srab_levptrc1_almtechlgc", "srabsign_levptrc1_almtechlgc", "vvod_ptrc1_talmgaslgc", "oper_vyvod_ptrc1_talmgaslgc",
            "srab_ptrc1_talmgaslgc", "srabsign_ptrc1_talmgaslgc", "zablok_ptrc1_talmgaslgc", "ET_ptrc1_talmgaslgc",
            "vvod_ptrc1_ttrgaslgc", "oper_vyvod_ptrc1_ttrgaslgc", "srab_ptrc1_ttrgaslgc", "srabsign_ptrc1_ttrgaslgc",
            "zablok_ptrc1_ttrgaslgc", "ET_ptrc1_ttrgaslgc", "vvod_ptrc1_tltcgaslgc", "oper_vyvod_ptrc1_tltcgaslgc",
            "srab_ptrc1_tltcgaslgc", "srabsign_ptrc1_tltcgaslgc", "zablok_ptrc1_tltcgaslgc", "ET_ptrc1_tltcgaslgc",
            "vvod_ptrc1_tofflvlgc", "oper_vyvod_ptrc1_tofflvlgc", "pusk_ptrc1_tofflvlgc", "srab_ptrc1_tofflvlgc",
            "vvod_rblc1_tofflvlgc", "oper_vyvod_rblc1_tofflvlgc", "zapret_rblc1_tofflvlgc",
            "vvod_rbre1_tofflvlgc", "oper_vyvod_rbre1_tofflvlgc", "zapret_rbre1_tofflvlgc",
            "SS_gz_sign", "SS_gz_zablok", "SS_gz_nizk_isol", "SS_tz_sign", "SS_tz_nizk_isol", "SS_tz_zablok",
            "SS_ts_sign", "SS_vnesh_otkl", "SS_vyh_zepi_razobr", "SS_bi_vyved", "SS_ot_sign",
            "SS_neispr_ot_gz", "SS_neispr_ot_tz", "SS_neispr_ot_v", "SS_ot_nn_sign",
            "SS_prev_vrem_per_ka", "SS_obsh_vnesh_sign", "pusk_lvalv"
        ]

        row, col = 0, 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=33, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_TECH_T(
            # Передаем SGF-параметры из self.sgf_params
            SGF1_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_EnaDis"].get(),
            SGF2_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl"].get(),
            #T1_oilptrc1_apttechlgc=self.settings["T1_oilptrc1_apttechlgc"].get(),
            SGF1_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_EnaDis"].get(),
            SGF2_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl"].get(),
            #T1_winptrc1_apttechlgc=self.settings["T1_winptrc1_apttechlgc"].get(),
            SGF1_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_EnaDis"].get(),
            SGF2_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl"].get(),
            #T1_vlvptrc1_apttechlgc=self.settings["T1_vlvptrc1_apttechlgc"].get(),
            T1_apttechlgc=self.settings["APTTECHLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_prvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis"].get(),
            SGF1_shvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis"].get(),
            SGF1_levptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis"].get(),
            SGF1_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_LowIsolSignCtrl"].get(),
            T1_ptrc1_talmgaslgc=self.settings["TALMGASLGC_1_PTRC1_TopOnBlk"].get(),
            SGF1_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_ttrgaslgc=self.settings["TTRGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_tltcgaslgc=self.settings["TLTCGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF2_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC2_Ctrl"].get(),
            SGF3_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC3_Ctrl"].get(),
            SGF1_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_EnaDis"].get(),
            SGF2_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl"].get(),
            SGF3_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl"].get(),
            SGF1_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF14_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA5"].get(),
            SGF6_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SG1"].get(),
            SGF7_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SG2"].get(),
            SGF8_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_GAS_OCControl"].get(),
            SGF9_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_TECH_OCControl"].get(),
            SGF10_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_OCcir_CB"].get(),
            SGF11_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_ARCnn_OCControl"].get(),
            SGF12_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_CBFPnn_OCControl"].get(),
            SGF13_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_IEDvt_OCControl"].get(),
        )
        print("part_TECH_T initialized")

    def start_polling(self):
        if not self.part:
            print("part_TECH_T not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")


    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла, добавляя '_SG1' к именам."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return

        try:
            handler = SettingsHandler.from_json_file(file_path)

            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is not None:
                    try:
                        if value_str.lower() in ("true", "1", "on"):
                            self.sgf_params[key].set(1)
                        elif value_str.lower() in ("false", "0", "off"):
                            self.sgf_params[key].set(0)
                        else:
                            print(f"⚠️ Неизвестное значение для {json_key}: '{value_str}'")
                    except Exception as e:
                        print(f"Ошибка при обработке {json_key}: {e}")

            # --- Обновление T-параметров (тоже с _SG1, если требуется) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is not None:
                    try:
                        val = float(value_str)
                        self.settings[key].set(val)
                    except ValueError:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = {value_str}")

            messagebox.showinfo("Успех", "Уставки успешно загружены из JSON-файла.")
            print("Параметры обновлены из JSON (с суффиксом _SG1)")

        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить уставки:\n{e}")
            print(f"Ошибка загрузки JSON: {e}")


    def save_settings_to_json(self):
        """Сохраняет текущие значения GUI в JSON-файл через SettingsHandler (все параметры с _SG1)."""
        file_path = askopenfilename(
            title="Выберите JSON-файл уставок для обновления",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return

        try:
            # 1. Загружаем исходный файл
            handler = SettingsHandler.from_json_file(file_path)

            # 2. Обновляем ВСЕ параметры с суффиксом _SG1
            all_params = {}
            # Добавляем SGF-параметры
            for key, var in self.sgf_params.items():
                all_params[key] = var.get()
            # Добавляем T-параметры
            for key, var in self.settings.items():
                all_params[key] = var.get()

            # 3. Пробуем обновить каждый с _SG1
            for key, value in all_params.items():
                json_key = key + "_SG1"
                if json_key in handler._settings:
                    handler.set_value_to_parameter(json_key, str(value))
                else:
                    print(f"⚠️ Параметр {json_key} отсутствует в файле")

            # 4. Сохраняем
            handler.save_to_json_file(file_path)

            messagebox.showinfo("Успех", f"Уставки сохранены в:\n{file_path}")
            print("Настройки успешно обновлены в JSON")

        except FileNotFoundError:
            messagebox.showerror("Ошибка", "Файл не найден.")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить:\n{e}")
            print(f"Ошибка: {e}")




if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfTECH_T_GUI(root)
    root.mainloop()