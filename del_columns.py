import os
import openpyxl

def delete_columns_from_excel_files():
    """
    Удаляет столбцы, содержащие указанные подстроки в заголовках,
    во всех xlsx файлах в текущей папке
    """
    # Список столбцов для удаления (по частичному совпадению в названии)
    #cols_to_delete = ["T4_cbcswi1_swctrl", "T3_cbcswi1_swctrl","SGF6_xcbr1_tsd", "T1_cbcswi1_hvbctrl"]
    cols_to_delete = ["mestnoe",]  
    # Получаем все xlsx файлы в текущей папке
    current_dir = os.getcwd()
    xlsx_files = [f for f in os.listdir(current_dir) if f.endswith('.xlsx')]
    
    print(f"Найдено {len(xlsx_files)} xlsx файлов")
    print(f"Будут удалены столбцы, содержащие: {cols_to_delete}")
    print("-" * 50)
    
    for file_name in xlsx_files:
        file_path = os.path.join(current_dir, file_name)
        print(f"\nОбработка: {file_name}")
        
        try:
            # Загружаем книгу
            wb = openpyxl.load_workbook(file_path)
            changes_made = False
            
            # Обрабатываем каждый лист
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                
                # Получаем заголовки из первой строки
                headers = []
                for col_idx in range(1, sheet.max_column + 1):
                    header = sheet.cell(row=1, column=col_idx).value
                    headers.append((col_idx, header))
                
                # Определяем столбцы для удаления (идем с конца)
                cols_to_remove = []
                for col_idx, header in reversed(headers):
                    if header:  # Проверяем, что заголовок не None
                        for pattern in cols_to_delete:
                            if pattern in str(header):
                                cols_to_remove.append(col_idx)
                                print(f"  Лист '{sheet_name}': удален столбец {col_idx} ('{header}')")
                                break
                
                # Удаляем столбцы
                if cols_to_remove:
                    for col_idx in sorted(cols_to_remove, reverse=True):
                        sheet.delete_cols(col_idx)
                    changes_made = True
            
            # Сохраняем изменения
            if changes_made:
                wb.save(file_path)
                print(f"  ✅ Изменения сохранены")
            else:
                print(f"  ℹ️  Столбцы для удаления не найдены")
            
            wb.close()
            
        except Exception as e:
            print(f"  ❌ Ошибка: {e}")
    
    print("\n" + "-" * 50)
    print("Готово!")

if __name__ == "__main__":
    delete_columns_from_excel_files()

    print("-" * 50)
    input("Нажмите Enter, чтобы выйти...")