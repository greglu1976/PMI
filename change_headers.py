

# Скрипт меняет заголовки в xlsx на заголовки из файла example.xlsx - только листы SGF_Parameters и Settings


import os
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

class ExcelHeadersReplacer:
    def __init__(self, example_file: str, target_folder: str = None):
        """
        Инициализирует замену заголовков в Excel файлах.
        
        Args:
            example_file: путь к файлу example.xlsx с эталонными заголовками
            target_folder: папка для поиска файлов (по умолчанию текущая папка)
        """
        self.example_file = example_file
        self.target_folder = target_folder or os.getcwd()
        self.headers_data = {}
        self.sheet_names = ["SGF_Parameters", "Settings"]
        
    def load_headers_from_example(self) -> bool:
        """Загружает заголовки из example.xlsx"""
        if not os.path.exists(self.example_file):
            print(f"❌ Файл не найден: {self.example_file}")
            return False
        
        try:
            workbook = load_workbook(self.example_file)
            
            for sheet_name in self.sheet_names:
                if sheet_name not in workbook.sheetnames:
                    print(f"⚠️  Лист '{sheet_name}' не найден в {self.example_file}")
                    continue
                
                sheet = workbook[sheet_name]
                headers = []
                
                # Получаем первую строку с заголовками
                for cell in sheet[1]:
                    headers.append(cell.value)
                
                self.headers_data[sheet_name] = {
                    "values": headers,
                    "formatting": self._extract_formatting(sheet, 1)
                }
                print(f"✅ Загружены заголовки из '{sheet_name}': {len(headers)} столбцов")
            
            workbook.close()
            return True
            
        except Exception as e:
            print(f"❌ Ошибка при загрузке {self.example_file}: {e}")
            return False
    
    def _extract_formatting(self, sheet, row: int):
        """Извлекает форматирование ячеек из строки"""
        formatting = {}
        for col_idx, cell in enumerate(sheet[row], 1):
            formatting[col_idx] = {
                "font": cell.font.copy() if cell.font else None,
                "fill": cell.fill.copy() if cell.fill else None,
                "border": cell.border.copy() if cell.border else None,
                "alignment": cell.alignment.copy() if cell.alignment else None,
                "number_format": cell.number_format
            }
        return formatting
    
    def _apply_formatting(self, sheet, row: int, formatting: dict):
        """Применяет форматирование к строке"""
        for col_idx, format_dict in formatting.items():
            cell = sheet.cell(row=row, column=col_idx)
            if format_dict["font"]:
                cell.font = format_dict["font"]
            if format_dict["fill"]:
                cell.fill = format_dict["fill"]
            if format_dict["border"]:
                cell.border = format_dict["border"]
            if format_dict["alignment"]:
                cell.alignment = format_dict["alignment"]
            cell.number_format = format_dict["number_format"]
    
    def find_xlsx_files(self) -> list:
        """Находит все xlsx файлы в папке (кроме example.xlsx)"""
        xlsx_files = []
        example_name = os.path.basename(self.example_file)
        
        try:
            for file in os.listdir(self.target_folder):
                if file.endswith(".xlsx") and file != example_name:
                    xlsx_files.append(os.path.join(self.target_folder, file))
            
            print(f"📁 Найдено {len(xlsx_files)} файлов для обработки")
            return xlsx_files
            
        except Exception as e:
            print(f"❌ Ошибка при поиске файлов: {e}")
            return []
    
    def replace_headers_in_file(self, file_path: str) -> bool:
        """Заменяет заголовки в одном файле"""
        if not os.path.exists(file_path):
            print(f"❌ Файл не найден: {file_path}")
            return False
        
        try:
            workbook = load_workbook(file_path)
            changes_made = False
            
            for sheet_name in self.sheet_names:
                if sheet_name not in self.headers_data:
                    continue
                
                if sheet_name not in workbook.sheetnames:
                    print(f"   ⚠️  Лист '{sheet_name}' не найден, пропускаем")
                    continue
                
                sheet = workbook[sheet_name]
                headers = self.headers_data[sheet_name]["values"]
                formatting = self.headers_data[sheet_name]["formatting"]
                
                # Очищаем первую строку
                for col_idx in range(1, sheet.max_column + 1):
                    sheet.cell(row=1, column=col_idx).value = None
                
                # Записываем новые заголовки
                for col_idx, value in enumerate(headers, 1):
                    sheet.cell(row=1, column=col_idx).value = value
                
                # Применяем форматирование
                self._apply_formatting(sheet, 1, formatting)
                
                print(f"   ✅ Заголовки заменены в листе '{sheet_name}'")
                changes_made = True
            
            if changes_made:
                workbook.save(file_path)
                print(f"💾 Файл сохранен: {os.path.basename(file_path)}")
            
            workbook.close()
            return True
            
        except Exception as e:
            print(f"❌ Ошибка при обработке {file_path}: {e}")
            return False
    
    def process_all_files(self) -> dict:
        """Обрабатывает все найденные файлы"""
        if not self.headers_data:
            print("❌ Заголовки не загружены. Запустите load_headers_from_example() первым.")
            return {"success": 0, "failed": 0, "skipped": 0}
        
        xlsx_files = self.find_xlsx_files()
        
        if not xlsx_files:
            print("❌ Файлы для обработки не найдены")
            return {"success": 0, "failed": 0, "skipped": 0}
        
        stats = {"success": 0, "failed": 0, "skipped": 0}
        
        print("\n" + "="*60)
        print("🔄 Начало замены заголовков...")
        print("="*60 + "\n")
        
        for file_path in xlsx_files:
            file_name = os.path.basename(file_path)
            print(f"📄 Обработка: {file_name}")
            
            if self.replace_headers_in_file(file_path):
                stats["success"] += 1
            else:
                stats["failed"] += 1
            print()
        
        return stats
    
    def print_summary(self, stats: dict):
        """Выводит итоговый отчет"""
        print("="*60)
        print("📊 ИТОГОВЫЙ ОТЧЕТ")
        print("="*60)
        print(f"✅ Успешно обработано: {stats['success']}")
        print(f"❌ Ошибок: {stats['failed']}")
        print(f"⏭️  Пропущено: {stats['skipped']}")
        total = stats['success'] + stats['failed'] + stats['skipped']
        print(f"📌 Всего: {total}")
        print("="*60)


def main():
    """Главная функция"""
    example_file = "example.xlsx"
    target_folder = os.getcwd()
    
    replacer = ExcelHeadersReplacer(example_file, target_folder)
    
    if not replacer.load_headers_from_example():
        print("❌ Не удалось загрузить заголовки. Завершение работы.")
        return 1
    
    stats = replacer.process_all_files()
    replacer.print_summary(stats)
    
    if stats["success"] > 0:
        print("\n✨ Замена заголовков успешно завершена!")
        return 0
    else:
        print("\n⚠️  Не удалось обработать файлы")
        return 1


if __name__ == "__main__":
    exit_code = main()
    sys.exit(exit_code)