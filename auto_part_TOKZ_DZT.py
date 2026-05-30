import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension
import openpyxl

# Импортируем класс partTOKZ
from lib2.PARTS.TOKZ_DZT import partTOKZ

from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartOfTOKGUI:

    OUTPUT_PARAMS = [
                "TOVCTOC_1_HVPTOC1_FuncEnabled", "TOVCTOC_1_HVPTOC1_FuncOperDisabled", "TOVCTOC_1_HVPTOC1_Str", "TOVCTOC_1_HVPTOC1_DE_I", "TOVCTOC_1_HVPTOC1_Op", "TOVCTOC_1_HVPTOC1_OpOnTrip", "TOVCTOC_1_PTOC1_FuncEnabled", "TOVCTOC_1_PTOC1_FuncOperDisabled", "TOVCTOC_1_PTOC1_Str", "TOVCTOC_1_PTOC1_DE_I", "TOVCTOC_1_PTOC1_Op", "TOVCTOC_1_PTOC1_OpOnTrip", "TOVCTOC_1_PTOC2_FuncEnabled", "TOVCTOC_1_PTOC2_FuncOperDisabled", "TOVCTOC_1_PTOC2_Str", "TOVCTOC_1_PTOC2_DE_I", "TOVCTOC_1_PTOC2_Op", "TOVCTOC_1_PTOC2_OpOnTrip", "TOVCTOC_1_TOVCTOC_Op",
                "LVARCTOC_1_PTOC1_FuncEnabled", "LVARCTOC_1_PTOC1_FuncOperDisabled", "LVARCTOC_1_PTOC1_Str", "LVARCTOC_1_PTOC1_DE_I",
                "LTCBLKTOC_1_PTOC1_FuncEnabled", "LTCBLKTOC_1_PTOC1_FuncOperDisabled", "LTCBLKTOC_1_PTOC1_Str", "LTCBLKTOC_1_PTOC1_DE_I",
                "STRTPALC_1_HVPTOC1_FuncEnabled", "STRTPALC_1_HVPTOC1_FuncOperDisabled", "STRTPALC_1_HVPTOC1_Str", "STRTPALC_1_HVPTOC1_DE_I", "STRTPALC_1_LVPTOC1_FuncEnabled", "STRTPALC_1_LVPTOC1_FuncOperDisabled", "STRTPALC_1_LVPTOC1_Str", "STRTPALC_1_LVPTOC1_DE_I", "STRTPALC_1_LVPTOC2_FuncEnabled", "STRTPALC_1_LVPTOC2_FuncOperDisabled", "STRTPALC_1_LVPTOC2_Str", "STRTPALC_1_LVPTOC2_DE_I", "STRTPALC_1_PTOC1_Str", "STRTPALC_1_PTOC1_FuncEnabled",
                "TPALC_1_HVPTOC1_FuncEnabled", "TPALC_1_HVPTOC1_FuncOperDisabled", "TPALC_1_HVPTOC1_Str", "TPALC_1_HVPTOC1_DE_I", "TPALC_1_LVPTOC1_FuncEnabled", "TPALC_1_LVPTOC1_FuncOperDisabled", "TPALC_1_LVPTOC1_Str", "TPALC_1_LVPTOC1_DE_I", "TPALC_1_LVPTOC2_FuncEnabled", "TPALC_1_LVPTOC2_FuncOperDisabled", "TPALC_1_LVPTOC2_Str", "TPALC_1_LVPTOC2_DE_I", "TPALC_1_PTOC1_Str", "TPALC_1_PTOC1_FuncEnabled",
                "EQPALC_1_PALC1_FuncEnabled", "EQPALC_1_PALC1_FuncOperDisabled", "EQPALC_1_PALC1_Str", "EQPALC_1_PALC1_OpOnSignal", "EQPALC_1_PALC1_Op",
                "TPRMOFFLVLGC_1_PTRC1_FuncEnabled", "TPRMOFFLVLGC_1_PTRC1_FuncOperDisabled", "pusk_ptrc1_tprmofflvlgc", "TPRMOFFLVLGC_1_PTRC1_Op", "TPRMOFFLVLGC_1_RBRE1_FuncEnabled", "TPRMOFFLVLGC_1_RBRE1_FuncOperDisabled", "TPRMOFFLVLGC_1_RBRE1_BlkOp",
                "DZT2_LVALH_1_CALH1_Alarm"
    ]


    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование токовых функций ФСУ ДЗТ2, v0.1 от 15.05.25, v0.2 от 25.07.25, v0.3 от 29.05.26")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для параметров SGF и настроек
        self.sgf_params = {
            "TOVCTOC_1_HVPTOC1_EnaDis": tk.IntVar(value=0),
            "TOVCTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "TOVCTOC_1_PTOC2_EnaDis": tk.IntVar(value=0),
            "LVARCTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "LTCBLKTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "STRTPALC_1_HVPTOC1_EnaDis": tk.IntVar(value=0),
            "STRTPALC_1_LVPTOC1_EnaDis": tk.IntVar(value=0),
            "STRTPALC_1_LVPTOC2_EnaDis": tk.IntVar(value=0),
            "TPALC_1_HVPTOC1_EnaDis": tk.IntVar(value=0),
            "TPALC_1_LVPTOC1_EnaDis": tk.IntVar(value=0),
            "TPALC_1_LVPTOC2_EnaDis": tk.IntVar(value=0),
            "EQPALC_1_PALC1_EnaDis": tk.IntVar(value=0),
            "EQPALC_1_PALC1_CurrCtrlEna": tk.IntVar(value=0),
            "EQPALC_1_PALC1_OilTmpCtrlEna": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "Номинальный ток входа ВН": tk.IntVar(value=5),
            "Номинальный ток входа НН1": tk.IntVar(value=5),
            "Номинальный ток входа НН2": tk.IntVar(value=5),                                    
        }

        self.settings = {
            "TOVCTOC_1_HVPTOC1_Top": tk.DoubleVar(value=1),
            "TOVCTOC_1_HVPTOC1_Iop": tk.DoubleVar(value=0.2),
            "TOVCTOC_1_PTOC1_Top": tk.DoubleVar(value=1),
            "TOVCTOC_1_PTOC1_Iop": tk.DoubleVar(value=0.2),
            "TOVCTOC_1_PTOC2_Top": tk.DoubleVar(value=1),
            "TOVCTOC_1_PTOC2_Iop": tk.DoubleVar(value=0.2),
            "LVARCTOC_1_PTOC1_Iop": tk.DoubleVar(value=0.2),
            "LTCBLKTOC_1_PTOC1_Iop": tk.DoubleVar(value=0.2),
            "STRTPALC_1_HVPTOC1_Iop": tk.DoubleVar(value=0.2),
            "STRTPALC_1_LVPTOC1_Iop": tk.DoubleVar(value=0.2),
            "STRTPALC_1_LVPTOC2_Iop": tk.DoubleVar(value=0.2),
            "TPALC_1_HVPTOC1_Iop": tk.DoubleVar(value=0.2),
            "TPALC_1_LVPTOC1_Iop": tk.DoubleVar(value=0.2),
            "TPALC_1_LVPTOC2_Iop": tk.DoubleVar(value=0.2),
            "EQPALC_1_PALC1_Top": tk.DoubleVar(value=1),
        }

        self.input_vars = {
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),
            "IA1": tk.DoubleVar(value=0),
            "IB1": tk.DoubleVar(value=0),
            "IC1": tk.DoubleVar(value=0),
            "DI_ControllerDisable": tk.IntVar(value=0),
            "DI_TOVCTOC": tk.IntVar(value=0),
            "DI_TOVCTOC_HVPTOC1": tk.IntVar(value=0),
            "DI_TOVCTOC_Sign": tk.IntVar(value=0),            
            #"NaOtkl_hvptoc1_tovctoc": tk.IntVar(value=0),
            "DI_TOVCTOC_LVPTOC1": tk.IntVar(value=0),
            #"NaOtkl_ptoc1_tovctoc": tk.IntVar(value=0),
            "DI_TOVCTOC_LVPTOC2": tk.IntVar(value=0),
            #"NaOtkl_ptoc2_tovctoc": tk.IntVar(value=0),
            "DI_LVARCTOC": tk.IntVar(value=0),
            "DI_LTCBLKTOC": tk.IntVar(value=0),            
            "DI_STRTPALC": tk.IntVar(value=0),
            "OV_hvptoc1_strpalc": tk.IntVar(value=0), # 30.05.2026 Сигнал отсутствует в ФСУ !!!!
            "DI_STRTPALC_LVPTOC1": tk.IntVar(value=0),
            "DI_STRTPALC_LVPTOC2": tk.IntVar(value=0),
            "DI_TPALC": tk.IntVar(value=0),
            "DI_TPALC_HVPTOC1": tk.IntVar(value=0),
            "DI_TPALC_LVPTOC1": tk.IntVar(value=0),
            "DI_TPALC_LVPTOC2": tk.IntVar(value=0),
            "DI_EQPALC": tk.IntVar(value=0),
            "DI_EQPALC_Sign": tk.IntVar(value=0),
            "FailCoolSys": tk.IntVar(value=0),
            "AlcOilTmp": tk.IntVar(value=0),
            "DI_TPRMOFFLVLGC": tk.IntVar(value=0),
            "DI_TJNTPTRC": tk.IntVar(value=0),
            "DI_JNTRBRE": tk.IntVar(value=0),
        }

        self.output_labels = {}

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc


        # Создание интерфейса
        self.create_widgets()


    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def _get_output_names(self):
        return self.OUTPUT_PARAMS  


    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.sgf_params.items():

            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            if 'Номинальный ток входа' in key:
                ttk.Combobox(sgf_frame, textvariable=var, values=[1, 5], state="readonly").grid(row=row, column=col + 1)
            else:
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 8:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.settings.items():

            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2


        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")

                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)

            elif isinstance(var, tk.DoubleVar):

                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)

                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:

            label = ttk.Label(output_frame, text=output, width=25, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)


            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = partTOKZ(
            SGF1_hvptoc1_tovctoc=self.sgf_params["TOVCTOC_1_HVPTOC1_EnaDis"].get(),
            SGF1_ptoc1_tovctoc=self.sgf_params["TOVCTOC_1_PTOC1_EnaDis"].get(),
            SGF1_ptoc2_tovctoc=self.sgf_params["TOVCTOC_1_PTOC2_EnaDis"].get(),
            SGF1_ptoc1_lvarctoc=self.sgf_params["LVARCTOC_1_PTOC1_EnaDis"].get(),
            SGF1_ptoc1_ltcblktoc=self.sgf_params["LTCBLKTOC_1_PTOC1_EnaDis"].get(),
            SGF1_hvptoc1_strpalc=self.sgf_params["STRTPALC_1_HVPTOC1_EnaDis"].get(),
            SGF1_lvptoc1_strpalc=self.sgf_params["STRTPALC_1_LVPTOC1_EnaDis"].get(),
            SGF1_lvptoc2_strpalc=self.sgf_params["STRTPALC_1_LVPTOC2_EnaDis"].get(),
            SGF1_hvptoc1_tpalc=self.sgf_params["TPALC_1_HVPTOC1_EnaDis"].get(),
            SGF1_lvptoc1_tpalc=self.sgf_params["TPALC_1_LVPTOC1_EnaDis"].get(),
            SGF1_lvptoc2_tpalc=self.sgf_params["TPALC_1_LVPTOC2_EnaDis"].get(),
            SGF1_lvoileqpalc_eqpalc=self.sgf_params["EQPALC_1_PALC1_EnaDis"].get(),
            SGF2_lvoileqpalc_eqpalc=self.sgf_params["EQPALC_1_PALC1_CurrCtrlEna"].get(),
            SGF3_lvoileqpalc_eqpalc=self.sgf_params["EQPALC_1_PALC1_OilTmpCtrlEna"].get(),
            SGF1_ptrc1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_RBRE1_EnaDis"].get(),
            T1_hvptoc1_tovctoc=self.settings["TOVCTOC_1_HVPTOC1_Top"].get(),
            Iset_hvptoc1_tovctoc=self.settings["TOVCTOC_1_HVPTOC1_Iop"].get(),
            T1_ptoc1_tovctoc=self.settings["TOVCTOC_1_PTOC1_Top"].get(),
            Iset_ptoc1_tovctoc=self.settings["TOVCTOC_1_PTOC1_Iop"].get(),
            T1_ptoc2_tovctoc=self.settings["TOVCTOC_1_PTOC2_Top"].get(),
            Iset_ptoc2_tovctoc=self.settings["TOVCTOC_1_PTOC2_Iop"].get(),
            Iset_ptoc1_lvarctoc=self.settings["LVARCTOC_1_PTOC1_Iop"].get(),
            Iset_ptoc1_ltcblktoc=self.settings["LTCBLKTOC_1_PTOC1_Iop"].get(),                        
            Iset_hvptoc1_strpalc=self.settings["STRTPALC_1_HVPTOC1_Iop"].get(),
            Iset_lvptoc1_strpalc=self.settings["STRTPALC_1_LVPTOC1_Iop"].get(),
            Iset_lvptoc2_strpalc=self.settings["STRTPALC_1_LVPTOC2_Iop"].get(),
            Iset_hvptoc1_tpalc=self.settings["TPALC_1_HVPTOC1_Iop"].get(),
            Iset_lvptoc1_tpalc=self.settings["TPALC_1_LVPTOC1_Iop"].get(),
            Iset_lvptoc2_tpalc=self.settings["TPALC_1_LVPTOC2_Iop"].get(),
            T1_lvoileqpalc_eqpalc=self.settings["EQPALC_1_PALC1_Top"].get(),
            Inom=self.sgf_params["Номинальный ток входа ВН"].get(),
            Inom2=self.sgf_params["Номинальный ток входа НН1"].get(),
            Inom3=self.sgf_params["Номинальный ток входа НН2"].get(),
        )

        print("partTOKZ initialized")

    def start_polling(self):
        if self.part is None:
            print("partTOKZ not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or float(value)!=0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])                    

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfTOKGUI(root)
    root.mainloop()