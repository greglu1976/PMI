import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl
import json

from lib2.PARTS.LO_T import part_LO
from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartLO_GUI:

    # === СПИСОК ВЫХОДНЫХ ПАРАМЕТРОВ (единое определение) ===
    OUTPUT_PARAMS = [
        "LVTTOC_1_LVTTOC_Str", "LVTTOC_1_PTOC1_Op", "LVTTOC_1_PTOC2_Op", "LVTTOC_1_PTOC3_Op", "T_LVTOC_1_PTOC1_Str", "T_LVTOC_1_PTOC1_Op",
        "TOFFLVLGC_1_PTRC1_FuncEnabled", "TOFFLVLGC_1_PTRC1_FuncOperDisabled", "TOFFLVLGC_1_PTRC1_Str", "TOFFLVLGC_1_PTRC1_Op", "TOFFLVLGC_1_LVCBRBLC1_FuncEnabled", "TOFFLVLGC_1_LVCBRBLC1_FuncOperDisabled", "TOFFLVLGC_1_LVCBRBLC1_BlkOp", "TOFFLVLGC_1_RBRE1_FuncEnabled", "TOFFLVLGC_1_RBRE1_FuncOperDisabled", "TOFFLVLGC_1_RBRE1_BlkOp", "T_HVTCBOFF_1_HVCBPTRC1_FuncEnabled", "T_HVTCBOFF_1_HVCBPTRC1_FuncOperDisabled", "T_HVTCBOFF_1_HVCBPTRC1_Op", "T_HVTCBOFF_1_HVCBPTRC1_Tr", "LVTCBOFF_1_LVCBPTRC1_FuncEnabled", "LVTCBOFF_1_LVCBPTRC1_FuncOperDisabled", "LVTCBOFF_1_LVCBPTRC1_Op", "LVTCBOFF_1_LVCBPTRC1_Tr", "LVTCBOFF_1_LVCBRECRBRE1_FuncEnabled", "LVTCBOFF_1_LVCBRECRBRE1_FuncOperDisabled", "LVTCBOFF_1_LVCBRECRBRE1_BlkOp", "LVTCBOFF_1_LVBTSRBLC1_FuncEnabled", "LVTCBOFF_1_LVBTSRBLC1_FuncOperDisabled", "LVTCBOFF_1_LVBTSRBLC1_BlkOp", "T_LVCBSUP_1_RCBF1_BlkToOpn",
       "T_TPBRF_1_GENRBRF1_FuncEnabled", "T_TPBRF_1_GENRBRF1_FuncOperDisabled", "T_TPBRF_1_GENRBRF1_Acc", "T_TPBRF_1_GENRBRF1_OpEx", "T_TPBRF_1_GENRBRF1_Str", "T_TPBRF_1_GENRBRF1_DE_I", "T_TPBRF_1_GENRBRF1_OpIn",
        "T_SignAssembly_1_GASSign", "T_SignAssembly_1_GASBlock", "T_SignAssembly_1_LowIsolGAS", "T_SignAssembly_1_TECHSign", "T_SignAssembly_1_LowIsolTECH", "T_SignAssembly_1_TECHBlock", "T_SignAssembly_1_ALMSign", "T_SignAssembly_1_OpExt", "T_SignAssembly_1_CtlCir", "T_SignAssembly_1_TestBlock", "T_SignAssembly_1_OCSign", "T_SignAssembly_1_GAS_OCControlSignAssem",  "T_SignAssembly_1_TECH_OCControlSignAssem", "T_SignAssembly_1_OCcir_CBSignAssem", "T_SignAssembly_1_OCnnSign", "T_SignAssembly_1_SwOperExcTim", "T_SignAssembly_1_ExtSignGen", "T_LVALH_1_CALH1_Alarm"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ЛО, УРОВ, СС, ПС, ЛО ВН, ЛО НН М300-Т. вер.0 от 07.04.25, вер.1 - 2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация SGF-параметров
        self.sgf_params = {
            "LVTTOC_1_KschemeCT": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC1_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC2_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_EnaDis": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VolMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_MICMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_ExtVFlMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_VCMod": tk.IntVar(value=0),
            "LVTTOC_1_PTOC3_SBMod": tk.IntVar(value=0),
            "LVTTOC_1_PTUV1_VoltStrCond": tk.IntVar(value=0),
            "LVTTOC_1_PHAR1_RegBlock": tk.IntVar(value=0),
            "LVTTOC_1_RBLC1_StepSel": tk.IntVar(value=0),
            "T_LVTOC_1_PTOC1_EnaDis": tk.IntVar(value=0),
            "T_LVTOC_1_PTOC1_KschemeCT": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_EnaDis": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_RstFrmCLS": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_KnobCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_EnaDis": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CurrentPickUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CBOSTypeCtrl": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_ActUpSwitch": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent": tk.IntVar(value=0),
            "TOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_RBRE1_PVOC3_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_EnaDis": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl": tk.IntVar(value=0),
            "TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl": tk.IntVar(value=0),
            "T_HVTCBOFF_1_HVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVCBRECRBRE1_EnaDis": tk.IntVar(value=0),
            "LVTCBOFF_1_LVBTSRBLC1_EnaDis": tk.IntVar(value=0), 
            "T_SignAssembly_1_Ctl_SA1": tk.IntVar(value=0),                       
            "T_SignAssembly_1_Ctl_SA2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA3": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA4": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SA5": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG1": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_SG2": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_GAS_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_TECH_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_OCcir_CB": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_ARCnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_CBFPnn_OCControl": tk.IntVar(value=0),
            "T_SignAssembly_1_Ctl_IEDvt_OCControl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }

        # Настройки (T-параметры)
        self.settings = {
            "LVTTOC_1_PTOC1_Top": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTOC1_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC1_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_Top": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC2_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC3_Top": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTOC3_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PTOC3_IopCoars": tk.DoubleVar(value=1),
            "LVTTOC_1_PTUV1_Uop": tk.DoubleVar(value=1),            
            "LVTTOC_1_PTUV1_U2op": tk.DoubleVar(value=1),
            "LVTTOC_1_PHAR1_Iop": tk.DoubleVar(value=1),
            "LVTTOC_1_PHAR1_PhStr": tk.DoubleVar(value=40),
            "T_LVTOC_1_PTOC1_Top": tk.DoubleVar(value=1),
            "T_LVTOC_1_PTOC1_Iop": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_EnBlk": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_ElmgWorking": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Top": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Iop": tk.DoubleVar(value=1),
            "T_HVTCBOFF_1_HVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
            "LVTCBOFF_1_LVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
        }

        # Входные параметры для Step()
        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "DI_LVTTOC": tk.IntVar(value=0),
            "DI_LVTOC": tk.IntVar(value=0),
            "DI_TOFFLVLGC": tk.IntVar(value=0),
            "DI_PTRC1": tk.IntVar(value=0),
            "DI_RBRE1": tk.IntVar(value=0),
            "DI_LVCBRBLC1": tk.IntVar(value=0),
            "DI_HVTCBOFF": tk.IntVar(value=0),
            "OpExtOfARC_NN": tk.IntVar(value=0),
            "OpExtOfCBFP_NN": tk.IntVar(value=0),
            "DI_LVTCBOFF": tk.IntVar(value=0),
            "DI_LVCBPTRC1": tk.IntVar(value=0),
            "DI_LVCBRECRBRE1": tk.IntVar(value=0),
            "DI_LVBTSRBLC1": tk.IntVar(value=0),
            "DI_LVCBSUP": tk.IntVar(value=0),
            "ExternalBlkCB": tk.IntVar(value=0),
            "DI_TPBRF": tk.IntVar(value=0),
            "ExternalRBRFStart": tk.IntVar(value=0),
            "T_TPBRF_1_CBOS1Supervision": tk.IntVar(value=0),
            "T_TPBRF_1_CBOS2Supervision": tk.IntVar(value=0),                        
            "CtlCirSwPos1": tk.IntVar(value=0),
            "CtlCirSwPos2": tk.IntVar(value=0),
            "CtlCirSwPos3": tk.IntVar(value=0),
            "CtlCirSwPos4": tk.IntVar(value=0),
            "CtlCirSwPos5": tk.IntVar(value=0),
            "TestBlockPos1": tk.IntVar(value=0),
            "TestBlockPos2": tk.IntVar(value=0),
            "GAS_OCControl": tk.IntVar(value=0),
            "TECH_OCControl": tk.IntVar(value=0),
            "OCcir_CB": tk.IntVar(value=0),
            "ARCnn_OCControl": tk.IntVar(value=0),
            "CBFPnn_OCControl": tk.IntVar(value=0),
            "IEDvt_OCControl": tk.IntVar(value=0),
            "ExtSignal1": tk.IntVar(value=0),
            "ExtSignal2": tk.IntVar(value=0),
            "ExtSignal3": tk.IntVar(value=0),
            "ExtSignal4": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),            
        }

        self.output_labels = {}  # Для вывода результатов

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        self.create_widgets()

    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def create_widgets(self):
        # Фрейм для SGF-параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 16:
                row = 0
                col += 2

        # Фрейм для настроек (T-параметры)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Фрейм для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)
        
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Фрейм для входных параметров
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)
                
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Фрейм для выходных параметров
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row, col = 0, 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=33, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)

            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_LO(
            # Передаем SGF-параметры из self.sgf_params
            SGF1=self.sgf_params["LVTTOC_1_KschemeCT"].get(),
            SGF1_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VolMod"].get(),
            SGF3_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_MICMod"].get(),
            SGF4_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_ExtVFlMod"].get(),
            SGF5_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_VCMod"].get(),
            SGF6_ptoc1=self.sgf_params["LVTTOC_1_PTOC1_SBMod"].get(),
            SGF1_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_EnaDis"].get(),
            SGF2_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VolMod"].get(),
            SGF3_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_MICMod"].get(),
            SGF4_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_ExtVFlMod"].get(),
            SGF5_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_VCMod"].get(),
            SGF6_ptoc2=self.sgf_params["LVTTOC_1_PTOC2_SBMod"].get(),
            SGF1_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_EnaDis"].get(),
            SGF2_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VolMod"].get(),
            SGF3_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_MICMod"].get(),
            SGF4_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_ExtVFlMod"].get(),
            SGF5_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_VCMod"].get(),
            SGF6_ptoc3=self.sgf_params["LVTTOC_1_PTOC3_SBMod"].get(),
            SGF1_ptuv1=self.sgf_params["LVTTOC_1_PTUV1_VoltStrCond"].get(),
            SGF1_phar1=self.sgf_params["LVTTOC_1_PHAR1_RegBlock"].get(),
            SGF1_rblc1=self.sgf_params["LVTTOC_1_RBLC1_StepSel"].get(),
            SGF1_ptoc1_lvtoc=self.sgf_params["T_LVTOC_1_PTOC1_EnaDis"].get(),
            SGF2_ptoc1_lvtoc=self.sgf_params["T_LVTOC_1_PTOC1_KschemeCT"].get(),
            SGF1_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_EnaDis"].get(),
            SGF2_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol"].get(),
            SGF3_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault"].get(),
            SGF4_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_RstFrmCLS"].get(),
            SGF5_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl"].get(),
            SGF6_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault"].get(),
            SGF7_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_KnobCtrl"].get(),
            SGF8_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm"].get(),
            SGF1_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_EnaDis"].get(),
            SGF2_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp"].get(),
            SGF3_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CurrentPickUp"].get(),
            SGF4_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CBOSTypeCtrl"].get(),
            SGF5_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_ActUpSwitch"].get(),
            SGF6_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF2_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC2_Ctrl"].get(),
            SGF3_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC3_Ctrl"].get(),
            SGF1_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_EnaDis"].get(),
            SGF2_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl"].get(),
            SGF3_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl"].get(),
            SGF1_hvcbptrc1_hvtcboff=self.sgf_params["T_HVTCBOFF_1_HVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbptrc1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbrecrbre1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVCBRECRBRE1_EnaDis"].get(),
            SGF1_lvbtsrblc1_lvtcboff=self.sgf_params["LVTCBOFF_1_LVBTSRBLC1_EnaDis"].get(),
            SGF1_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["T_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["T_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_lvalh=self.sgf_params["T_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_lvalh=self.sgf_params["T_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_lvalh=self.sgf_params["T_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF14_lvalh=self.sgf_params["T_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SA5"].get(),
            SGF6_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SG1"].get(),
            SGF7_tsa=self.sgf_params["T_SignAssembly_1_Ctl_SG2"].get(),
            SGF8_tsa=self.sgf_params["T_SignAssembly_1_Ctl_GAS_OCControl"].get(),
            SGF9_tsa=self.sgf_params["T_SignAssembly_1_Ctl_TECH_OCControl"].get(),
            SGF10_tsa=self.sgf_params["T_SignAssembly_1_Ctl_OCcir_CB"].get(),
            SGF11_tsa=self.sgf_params["T_SignAssembly_1_Ctl_ARCnn_OCControl"].get(),
            SGF12_tsa=self.sgf_params["T_SignAssembly_1_Ctl_CBFPnn_OCControl"].get(),
            SGF13_tsa=self.sgf_params["T_SignAssembly_1_Ctl_IEDvt_OCControl"].get(),

            T1_ptoc1 = self.settings["LVTTOC_1_PTOC1_Top"].get()/1000 if self.settings["LVTTOC_1_PTOC1_Top"].get()>=100 else self.settings["LVTTOC_1_PTOC1_Top"].get(), #T1_ptoc1=self.settings["LVTTOC_1_PTOC1_Top"].get(),

            Iset_ptoc1=self.settings["LVTTOC_1_PTOC1_Iop"].get(),
            Icoarse_ptoc1=self.settings["LVTTOC_1_PTOC1_IopCoars"].get(),

            T1_ptoc2= self.settings["LVTTOC_1_PTOC2_Top"].get()/1000 if self.settings["LVTTOC_1_PTOC2_Top"].get()>=100 else self.settings["LVTTOC_1_PTOC2_Top"].get(),  #                  self.settings["LVTTOC_1_PTOC2_Top"].get(),

            Iset_ptoc2=self.settings["LVTTOC_1_PTOC2_Iop"].get(),
            Icoarse_ptoc2=self.settings["LVTTOC_1_PTOC2_IopCoars"].get(),

            T1_ptoc3=  self.settings["LVTTOC_1_PTOC3_Top"].get()/1000 if self.settings["LVTTOC_1_PTOC3_Top"].get()>=100 else self.settings["LVTTOC_1_PTOC3_Top"].get(),                   #self.settings["LVTTOC_1_PTOC3_Top"].get(),

            Iset_ptoc3=self.settings["LVTTOC_1_PTOC3_Iop"].get(),
            Icoarse_ptoc3=self.settings["LVTTOC_1_PTOC3_IopCoars"].get(),
            Uop_ptuv1=self.settings["LVTTOC_1_PTUV1_Uop"].get(),
            U2op_ptuv1=self.settings["LVTTOC_1_PTUV1_U2op"].get(),
            Imax_phar1=self.settings["LVTTOC_1_PHAR1_Iop"].get(),
            Ratio_phar1=self.settings["LVTTOC_1_PHAR1_PhStr"].get(),

            T1_ptoc1_lvtoc= self.settings["T_LVTOC_1_PTOC1_Top"].get()/1000 if self.settings["T_LVTOC_1_PTOC1_Top"].get()>=100 else self.settings["T_LVTOC_1_PTOC1_Top"].get(),                          #self.settings["T_LVTOC_1_PTOC1_Top"].get(),

            Iset_ptoc1_lvtoc=self.settings["T_LVTOC_1_PTOC1_Iop"].get(),
            T1_rcbf1_lvcbsup= self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),    #self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),
            T2_rcbf1_lvcbsup= self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(), #self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(),
            T3_rcbf1_lvcbsup= self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),   #self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),
            T1_rbrf1_tpbrf= self.settings["T_TPBRF_1_GENRBRF1_Top"].get()/1000 if self.settings["T_TPBRF_1_GENRBRF1_Top"].get()>=100 else self.settings["T_TPBRF_1_GENRBRF1_Top"].get(),      #self.settings["T_TPBRF_1_GENRBRF1_Top"].get(),
            Iset_rbrf1_tpbrf=self.settings["T_TPBRF_1_GENRBRF1_Iop"].get(),
            T1_hvcbptrc1_hvtcboff=self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get()/1000 if self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get()>=100 else self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(), #   self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(),
            T1_lvcbptrc1_lvtcboff=self.settings["LVTCBOFF_1_LVCBPTRC1_Tpulse"].get()/1000 if self.settings["LVTCBOFF_1_LVCBPTRC1_Tpulse"].get()>=100 else self.settings["LVTCBOFF_1_LVCBPTRC1_Tpulse"].get(),        #self.settings["LVTCBOFF_1_LVCBPTRC1_Tpulse"].get(),
        )
        print("part_LO initialized")

    def start_polling(self):
        if not self.part:
            print("part_LO not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")


    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

    # === МЕТОДЫ ДЛЯ РАБОТЫ С JSON ФАЙЛАМИ УСТАВОК ===
    
    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла с суффиксом _SG1."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return
        try:
            handler = SettingsHandler.from_json_file(file_path)
            if not self.meta_handler:
                print("⚠️ Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                    elif type_str == "130":  # Integer
                        try:
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()

                            int_val = int(float(value_clean.replace(',', '.')))
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            value_clean = str(value_str).strip()
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            # --- Обновление T-параметров (settings) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                        else:
                            self.settings[key].set(float_val)
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            print("✅ Параметры обновлены из JSON (с суффиксом _SG1)")
        except FileNotFoundError:
            print(f"❌ Ошибка: Файл не найден: {file_path}")
        except Exception as e:
            print(f"❌ Ошибка загрузки JSON: {e}")
    
    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()

    def _get_output_names(self):
        return self.OUTPUT_PARAMS   


if __name__ == "__main__":
    root = tk.Tk()
    app = PartLO_GUI(root)
    root.mainloop()