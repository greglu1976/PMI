# Скопировано без изменений из auto_part_LO_T

import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl

from lib2.PARTS.LO_DZT import part_LO 

from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartLO_GUI:

    OUTPUT_PARAMS = [
            "TPRMOFFLVLGC_1_PTRC1_FuncEnabled", "TPRMOFFLVLGC_1_PTRC1_FuncOperDisabled", "pusk_ptrc1_tprmofflvlgc", "TPRMOFFLVLGC_1_PTRC1_Op", "TPRMOFFLVLGC_1_RBRE1_FuncEnabled", "TPRMOFFLVLGC_1_RBRE1_FuncOperDisabled", "TPRMOFFLVLGC_1_RBRE1_BlkOp",
            "HVTCBOFF_1_HVCBPTRC1_FuncEnabled", "HVTCBOFF_1_HVCBPTRC1_FuncOperDisabled", "HVTCBOFF_1_HVCBPTRC1_Op", "HVTCBOFF_1_HVCBPTRC1_Tr",
            "LVTPRMCBOFF1_1_LVCBPTRC1_FuncEnabled", "LVTPRMCBOFF1_1_LVCBPTRC1_FuncOperDisabled", "LVTPRMCBOFF1_1_LVCBPTRC1_Op", "LVTPRMCBOFF1_1_LVCBPTRC1_Tr", "LVTPRMCBOFF1_1_LVCBRECRBRE1_FuncEnabled", "LVTPRMCBOFF1_1_LVCBRECRBRE1_FuncOperDisabled", "LVTPRMCBOFF1_1_LVCBRECRBRE1_BlkOp",
            "LVTPRMCBOFF2_1_LVCBPTRC1_FuncEnabled", "LVTPRMCBOFF2_1_LVCBPTRC1_FuncOperDisabled", "LVTPRMCBOFF2_1_LVCBPTRC1_Op", "LVTPRMCBOFF2_1_LVCBPTRC1_Tr", "LVTPRMCBOFF2_1_LVCBRECRBRE1_FuncEnabled", "LVTPRMCBOFF2_1_LVCBRECRBRE1_FuncOperDisabled", "LVTPRMCBOFF2_1_LVCBRECRBRE1_BlkOp",
            "DZT2_TPBRF_1_GENRBRF1_FuncEnabled", "DZT2_TPBRF_1_GENRBRF1_FuncOperDisabled", "DZT2_TPBRF_1_GENRBRF1_OpEx", "DZT2_TPBRF_1_GENRBRF1_Str", "DZT2_TPBRF_1_GENRBRF1_DE_I", "DZT2_TPBRF_1_GENRBRF1_OpIn",
            "DZT2_SignAssembly_1_GASSign", "DZT2_SignAssembly_1_LowIsolGAS", "DZT2_SignAssembly_1_GASBlock", "DZT2_SignAssembly_1_TECHSign", "DZT2_SignAssembly_1_LowIsolTECH", "DZT2_SignAssembly_1_TECHBlock", "DZT2_SignAssembly_1_ALMSign", "DZT2_SignAssembly_1_OpExt", "DZT2_SignAssembly_1_CtlCir", "DZT2_SignAssembly_1_TestBlock", "DZT2_SignAssembly_1_OCSign", "DZT2_SignAssembly_1_GAS_OCControlSignAssem", "DZT2_SignAssembly_1_TECH_OCControlSignAssem", "DZT2_SignAssembly_1_OCnnSign", "DZT2_SignAssembly_1_ExtSignGen",
            "DZT2_LVALH_1_CALH1_Alarm"
        ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ЛО, УРОВ, СС, ПС, ЛО ВН, ЛО НН М300-ДЗТ2. вер.3 от 29.05.26")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")


        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None


        # Инициализация SGF-параметров
        self.sgf_params = {
            "DZT2_TPBRF_1_GENRBRF1_EnaDis": tk.IntVar(value=0),
            "DZT2_TPBRF_1_GENRBRF1_CurrentPickUp": tk.IntVar(value=0),
            "DZT2_TPBRF_1_GENRBRF1_ActUpSwitch": tk.IntVar(value=0),
            "DZT2_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_PTRC1_EnaDis": tk.IntVar(value=0),
            "TPRMOFFLVLGC_1_RBRE1_EnaDis": tk.IntVar(value=0),
            "HVTCBOFF_1_HVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTPRMCBOFF1_1_LVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTPRMCBOFF1_1_LVCBRECRBRE1_EnaDis": tk.IntVar(value=0),
            "LVTPRMCBOFF2_1_LVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "LVTPRMCBOFF2_1_LVCBRECRBRE1_EnaDis": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA1": tk.IntVar(value=0),                       
            "DZT2_SignAssembly_1_Ctl_SA2": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA3": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SA4": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG2": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_SG3": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_GAS_OCControl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_TECH_OCControl_SG1": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_ARCnn1_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_ARCnn2_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_CBFPnn1_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_Ctl_CBFPnn2_OCControl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "DZT2_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }

        # Настройки (T-параметры)
        self.settings = {
            "DZT2_TPBRF_1_GENRBRF1_Top": tk.DoubleVar(value=1),
            "DZT2_TPBRF_1_GENRBRF1_Iop": tk.DoubleVar(value=1),
            "HVTCBOFF_1_HVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
            "LVTPRMCBOFF1_1_LVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
            "LVTPRMCBOFF2_1_LVCBPTRC1_Tpulse": tk.DoubleVar(value=1),
        }

        # Входные параметры для Step()
        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "OpExtOfCoolSys": tk.IntVar(value=0),
            "DI_TPRMOFFLVLGC": tk.IntVar(value=0),
            "DI_TJNTPTRC": tk.IntVar(value=0),
            "DI_JNTRBRE": tk.IntVar(value=0),
            "DI_HVTCBOFF": tk.IntVar(value=0),
            "OpExtOfARC_NN1": tk.IntVar(value=0),
            "OpExtOfARC_NN2": tk.IntVar(value=0),
            "OpExtOfCBFP_NN1": tk.IntVar(value=0),
            "OpExtOfCBFP_NN2": tk.IntVar(value=0),
            "DI_LVTPRMCBOFF1": tk.IntVar(value=0),
            "DI_LVCBPTRC1": tk.IntVar(value=0),
            "DI_LVCBRECRBRE1": tk.IntVar(value=0),
            "DI_LVTPRMCBOFF2": tk.IntVar(value=0),
            "DI_LVCBPTRC2": tk.IntVar(value=0),
            "DI_LVCBRECRBRE2": tk.IntVar(value=0),
            "DI_TPBRF": tk.IntVar(value=0),
            "ExternalRBRFStart": tk.IntVar(value=0),
            "CtlCirSwPos1": tk.IntVar(value=0),
            "CtlCirSwPos2": tk.IntVar(value=0),
            "CtlCirSwPos3": tk.IntVar(value=0),
            "CtlCirSwPos4": tk.IntVar(value=0),
            "TestBlockPos1": tk.IntVar(value=0),
            "TestBlockPos2": tk.IntVar(value=0),
            "TestBlockPos3": tk.IntVar(value=0),
            "GAS_OCControl": tk.IntVar(value=0),
            "DZT2_SignAssembly_1_TECH_OCControlSignAssem": tk.IntVar(value=0),
            "ARCnn1_OCControl": tk.IntVar(value=0),
            "ARCnn2_OCControl": tk.IntVar(value=0),            
            "CBFPnn1_OCControl": tk.IntVar(value=0),
            "CBFPnn2_OCControl": tk.IntVar(value=0),
            "ExtSignal1": tk.IntVar(value=0),
            "ExtSignal2": tk.IntVar(value=0),
            "ExtSignal3": tk.IntVar(value=0),
            "ExtSignal4": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),            
        }

        self.output_labels = {}  # Для вывода результатов

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        self.create_widgets()

    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def _get_output_names(self):
        return self.OUTPUT_PARAMS 


    def create_widgets(self):
        # Фрейм для SGF-параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row, col = 0, 0

        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 10:
                row = 0
                col += 2

        # Фрейм для настроек (T-параметры)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)

            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Фрейм для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Фрейм для входных параметров
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")

                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)

            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)

                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Фрейм для выходных параметров
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row, col = 0, 0
        for output in self.OUTPUT_PARAMS:
            
            label = ttk.Label(output_frame, text=output, width=33, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)


            row += 1
            if row >= 25:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_LO(
            # Передаем SGF-параметры из self.sgf_params
            SGF1_rbrf1_tpbrf=self.sgf_params["DZT2_TPBRF_1_GENRBRF1_EnaDis"].get(),
            SGF2_rbrf1_tpbrf=self.sgf_params["DZT2_TPBRF_1_GENRBRF1_CurrentPickUp"].get(),
            SGF3_rbrf1_tpbrf=self.sgf_params["DZT2_TPBRF_1_GENRBRF1_ActUpSwitch"].get(),
            SGF4_rbrf1_tpbrf=self.sgf_params["DZT2_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent"].get(),
            SGF1_ptrc1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tprmofflvlgc=self.sgf_params["TPRMOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF1_hvcbptrc1_hvtcboff=self.sgf_params["HVTCBOFF_1_HVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbptrc1_lvtcboff1=self.sgf_params["LVTPRMCBOFF1_1_LVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbrecrbre1_lvtcboff1=self.sgf_params["LVTPRMCBOFF1_1_LVCBRECRBRE1_EnaDis"].get(),
            SGF1_lvcbptrc1_lvtcboff2=self.sgf_params["LVTPRMCBOFF2_1_LVCBPTRC1_EnaDis"].get(),
            SGF1_lvcbrecrbre1_lvtcboff2=self.sgf_params["LVTPRMCBOFF2_1_LVCBRECRBRE1_EnaDis"].get(),
            SGF1_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_lvalh=self.sgf_params["DZT2_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG1"].get(),
            SGF6_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG2"].get(),
            SGF7_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_SG3"].get(),
            SGF8_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_GAS_OCControl_SG1"].get(),
            SGF9_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_TECH_OCControl_SG1"].get(),
            SGF10_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_ARCnn1_OCControl"].get(),
            SGF11_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_ARCnn2_OCControl"].get(),
            SGF12_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_CBFPnn1_OCControl"].get(),
            SGF13_tsa=self.sgf_params["DZT2_SignAssembly_1_Ctl_CBFPnn2_OCControl"].get(),
            T1_rbrf1_tpbrf=self.settings["DZT2_TPBRF_1_GENRBRF1_Top"].get(),
            Iset_rbrf1_tpbrf=self.settings["DZT2_TPBRF_1_GENRBRF1_Iop"].get(),
            T1_hvcbptrc1_hvtcboff=self.settings["HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(),
            T1_lvcbptrc1_lvtcboff1=self.settings["LVTPRMCBOFF1_1_LVCBPTRC1_Tpulse"].get(),
            T1_lvcbptrc1_lvtcboff2=self.settings["LVTPRMCBOFF2_1_LVCBPTRC1_Tpulse"].get(),
        )
        print("part_LO initialized")

    def start_polling(self):
        if not self.part:
            print("part_LO not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")


    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = PartLO_GUI(root)
    root.mainloop()