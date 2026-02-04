import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image
import numpy as np
import os

# Читаем исходный файл
df = pd.read_excel('input.xlsx')  # замените 'input.xlsx' на имя вашего файла

# Преобразуем время
df['Время'] = pd.to_datetime(df['Время'], format='%d.%m.%Y %H:%M:%S,%f')
df['Время_формат'] = df['Время'].dt.strftime('%S.%f').str[:6]  # СС.МММ

# ==================== СОЗДАНИЕ ТАБЛИЦЫ ====================

# Создаем сводную таблицу, группируя события по времени
pivot_data = []
time_objects = []

for time_str, group in df.groupby('Время_формат'):
    time_dt = group.iloc[0]['Время']
    time_objects.append(time_dt)
    
    events = []
    has_non_zero = False
    
    for _, row in group.iterrows():
        event_name = row['Наименование события']
        value = row['Значение']
        
        if pd.notna(value) and value != 0:
            events.append(f"{event_name} [{value}]")
            has_non_zero = True
        else:
            events.append(event_name)
    
    pivot_data.append({
        'Время': time_str,
        'Время_dt': time_dt,
        'События': '\n'.join(events),
        'Есть_ненулевые': has_non_zero
    })

# Создаем DataFrame
pivot_df = pd.DataFrame(pivot_data)
pivot_df = pivot_df.sort_values('Время_dt')

# Рассчитываем временные интервалы
time_intervals = []
time_intervals_numeric = []
for i in range(len(pivot_df)):
    if i == 0:
        time_intervals.append('0.000')
        time_intervals_numeric.append(0.0)
    else:
        time_diff = (pivot_df.iloc[i]['Время_dt'] - pivot_df.iloc[i-1]['Время_dt']).total_seconds()
        time_intervals.append(f"{time_diff:.3f}")
        time_intervals_numeric.append(time_diff)

# Создаем Excel файл
output_file = 'events_timeline_output.xlsx'
wb = Workbook()
ws = wb.active
ws.title = "Таймлайн событий"

# Заполняем данные
time_row = pivot_df['Время'].tolist()
events_row = pivot_df['События'].tolist()
flags_row = pivot_df['Есть_ненулевые'].tolist()

for col_idx, (time_val, events_val, interval_val, flag_val) in enumerate(
    zip(time_row, events_row, time_intervals, flags_row), 1):
    
    ws.cell(row=1, column=col_idx, value=time_val)          # Время
    ws.cell(row=2, column=col_idx, value=events_val)        # События
    ws.cell(row=3, column=col_idx, value=interval_val)      # Интервалы
    ws.cell(row=4, column=col_idx, value=flag_val)          # Флаги

# Стилизация таблицы
time_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
time_font = Font(color="FFFFFF", bold=True, size=10)
interval_fill = PatternFill(start_color="F2E0D6", end_color="F2E0D6", fill_type="solid")
interval_font = Font(color="000000", size=9)
normal_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
normal_font = Font(color="000000", size=9)
red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
red_font = Font(color="FF0000", bold=True, size=9)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Форматирование
for col_idx in range(1, len(time_row) + 1):
    # Время
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = time_fill
    cell.font = time_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
    
    # Интервалы
    cell = ws.cell(row=3, column=col_idx)
    cell.fill = interval_fill
    cell.font = interval_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
    
    if col_idx == 1:
        cell.value = f"0.000 (начало)"
    else:
        interval_val = float(cell.value)
        if interval_val > 0.1:
            cell.font = Font(color="FF0000", bold=True, size=9)
    
    # События
    cell = ws.cell(row=2, column=col_idx)
    flag_cell = ws.cell(row=4, column=col_idx)
    has_non_zero = flag_cell.value
    
    if has_non_zero:
        cell.fill = red_fill
        cell.font = red_font
    else:
        cell.fill = normal_fill
        cell.font = normal_font
    
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    cell.border = thin_border

# Скрываем строку с флагами
ws.row_dimensions[4].hidden = True

# Настройка размеров
for col_idx in range(1, len(time_row) + 1):
    max_length = 0
    for row in range(1, 4):
        cell_value = ws.cell(row=row, column=col_idx).value
        if cell_value:
            if "\n" in str(cell_value):
                lines = str(cell_value).split("\n")
                max_line_length = max(len(line) for line in lines)
                max_length = max(max_length, max_line_length)
            else:
                max_length = max(max_length, len(str(cell_value)))
    
    column_letter = ws.cell(row=1, column=col_idx).column_letter
    ws.column_dimensions[column_letter].width = min(max_length + 3, 40)

ws.row_dimensions[1].height = 25
ws.row_dimensions[2].height = 80
ws.row_dimensions[3].height = 25

# ==================== СОЗДАНИЕ ГРАФИКА С ВРЕМЕНЕМ НА ЛИНИЯХ ====================

print("\nПодготовка данных для графика...")

# Получаем уникальные события
unique_events = df['Наименование события'].unique()
print(f"Найдено {len(unique_events)} уникальных событий")

# Создаем временную шкалу
times = pivot_df['Время_dt'].tolist()
time_str_labels = [t.strftime('%S.%f')[:-3] for t in times]

# Подготовка данных для каждого события
event_data = {}
for event in unique_events:
    event_data[event] = {
        'times': [],
        'values': [],
        'colors': [],
        'x_indices': [],
        'time_diffs': []  # Время между точками для этого события
    }

# Заполняем данные
for idx, time_dt in enumerate(times):
    time_str = time_str_labels[idx]
    events_at_time = df[df['Время_формат'] == time_str]
    
    for event in unique_events:
        event_rows = events_at_time[events_at_time['Наименование события'] == event]
        
        if not event_rows.empty:
            value = event_rows.iloc[0]['Значение']
            event_data[event]['times'].append(time_dt)
            event_data[event]['values'].append(value if pd.notna(value) else 0)
            event_data[event]['x_indices'].append(idx)
            event_data[event]['colors'].append('red' if pd.notna(value) and value != 0 else 'blue')
            
            # Рассчитываем время от предыдущей точки этого события
            if len(event_data[event]['times']) > 1:
                prev_time = event_data[event]['times'][-2]
                time_diff = (time_dt - prev_time).total_seconds()
                event_data[event]['time_diffs'].append(time_diff)
            else:
                event_data[event]['time_diffs'].append(0)  # Для первой точки

# Создаем график
plt.figure(figsize=(18, 12))

# Цвета для разных событий
colors = plt.cm.tab20(np.linspace(0, 1, len(unique_events)))

# Для легенды
legend_patches = []

# Рисуем линии и добавляем время на линии
for idx, (event_name, data) in enumerate(event_data.items()):
    if len(data['times']) > 1:
        x = data['x_indices']
        y = [idx] * len(x)
        
        # Линия, соединяющая точки
        line = plt.plot(x, y, 
                       color=colors[idx % len(colors)], 
                       linewidth=3, 
                       alpha=0.8,
                       marker='',
                       zorder=1,
                       solid_capstyle='round')[0]
        
        # Добавляем время на линию между точками
        for i in range(len(x) - 1):
            # Средняя точка между двумя точками
            mid_x = (x[i] + x[i + 1]) / 2
            mid_y = idx
            
            # Время между точками (из time_diffs)
            time_diff = data['time_diffs'][i + 1]  # i+1 потому что первая точка имеет diff=0
            
            # Форматируем время
            if time_diff < 0.001:  # Менее 1 мс
                time_text = f"{time_diff*1000:.1f}мс"
            elif time_diff < 1:  # Менее 1 секунды
                time_text = f"{time_diff*1000:.0f}мс"
            else:
                time_text = f"{time_diff:.3f}с"
            
            # Создаем фоновую подложку для текста
            plt.text(mid_x, mid_y + 0.15, time_text,
                    fontsize=9,
                    fontweight='bold',
                    ha='center',
                    va='center',
                    bbox=dict(boxstyle="round,pad=0.3",
                             facecolor="white",
                             edgecolor=colors[idx % len(colors)],
                             alpha=0.9,
                             linewidth=2),
                    zorder=3,
                    color=colors[idx % len(colors)])
            
            # Также добавляем стрелки для указания направления
            plt.annotate('',
                        xy=(x[i+1], mid_y),
                        xytext=(x[i], mid_y),
                        arrowprops=dict(arrowstyle='->',
                                       color=colors[idx % len(colors)],
                                       linewidth=1,
                                       alpha=0.5),
                        zorder=2)
    
    # Точки (маркеры)
    for i, (x_idx, value, color) in enumerate(zip(data['x_indices'], data['values'], data['colors'])):
        marker_size = 120 if (pd.notna(value) and value != 0) else 80
        marker_shape = 'o' if (pd.notna(value) and value == 0) else 's'
        
        plt.scatter(x_idx, idx,
                   s=marker_size,
                   c=color,
                   marker=marker_shape,
                   edgecolors='black',
                   linewidth=2,
                   zorder=4,
                   alpha=0.95)
        
        # Подпись значения рядом с точкой
        if pd.notna(value) and value != 0:
            value_text = f"={value}"
            plt.text(x_idx + 0.15, idx + 0.15, value_text,
                    fontsize=9,
                    fontweight='bold',
                    ha='left',
                    va='bottom',
                    bbox=dict(boxstyle="round,pad=0.3",
                             facecolor="yellow",
                             edgecolor="black",
                             alpha=0.8,
                             linewidth=1),
                    zorder=5)
        
        # Подпись времени события под точкой (для первой и последней точки каждого события)
        if i == 0 or i == len(data['x_indices']) - 1:
            time_str_short = time_str_labels[x_idx]
            plt.text(x_idx, idx - 0.35, time_str_short,
                    fontsize=8,
                    ha='center',
                    va='top',
                    bbox=dict(boxstyle="round,pad=0.2",
                             facecolor="lightgray",
                             alpha=0.7),
                    zorder=3)

# Настройка осей
plt.yticks(range(len(unique_events)), unique_events, fontsize=10, fontweight='bold')
plt.xticks(range(len(time_str_labels)), time_str_labels, rotation=45, ha='right', fontsize=9)

# Сетка
plt.grid(True, alpha=0.2, linestyle='--', which='both', zorder=0)

# Заголовок и подписи
plt.title('Временная диаграмма событий с интервалами на линиях', 
          fontsize=16, fontweight='bold', pad=25)
plt.xlabel('Время (секунды.миллисекунды)', fontsize=12)
plt.ylabel('События', fontsize=12)

# Легенда
zero_patch = mpatches.Patch(color='blue', label='Значение = 0', alpha=0.8)
nonzero_patch = mpatches.Patch(color='red', label='Значение ≠ 0', alpha=0.8)
circle_patch = mpatches.Patch(color='white', label='Круг = 0', 
                              edgecolor='black', linewidth=1.5)
square_patch = mpatches.Patch(color='white', label='Квадрат ≠ 0', 
                              edgecolor='black', linewidth=1.5)

# Добавляем легенду
first_legend = plt.legend(handles=[zero_patch, nonzero_patch], 
                         loc='upper left', 
                         bbox_to_anchor=(1.02, 1),
                         borderaxespad=0.,
                         title='Значения событий',
                         fontsize=10,
                         title_fontsize=11)

plt.gca().add_artist(first_legend)
plt.legend(handles=[circle_patch, square_patch], 
          loc='upper left', 
          bbox_to_anchor=(1.02, 0.85),
          borderaxespad=0.,
          title='Формы маркеров',
          fontsize=10,
          title_fontsize=11)

# Добавляем легенду для времени
time_example = mpatches.Patch(color='white', label='На линиях: время между\nсобытиями в мс/сек',
                             edgecolor='none')
plt.legend(handles=[time_example], 
          loc='upper left', 
          bbox_to_anchor=(1.02, 0.7),
          borderaxespad=0.,
          title='Обозначения времени',
          fontsize=9,
          title_fontsize=10)

# Добавляем статистику в заголовок
total_time = (times[-1] - times[0]).total_seconds() if len(times) > 1 else 0
plt.figtext(0.02, 0.98, 
           f"Всего событий: {len(unique_events)} | Временных точек: {len(times)} | Общее время: {total_time:.3f} сек",
           fontsize=10,
           bbox=dict(boxstyle="round,pad=0.3", facecolor="lightblue", alpha=0.8))

# Настройка layout
plt.tight_layout(rect=[0, 0, 0.82, 0.95])  # Оставляем больше места для легенды

# Сохраняем график
graph_filename = 'events_timeline_with_intervals.png'
plt.savefig(graph_filename, dpi=200, bbox_inches='tight')
print(f"График с временем на линиях сохранен как: {graph_filename}")

# Добавляем график в Excel
if os.path.exists(graph_filename):
    img = Image(graph_filename)
    img.width = 900
    img.height = 600
    ws.add_image(img, 'A15')  # Помещаем ниже таблицы
    print("График добавлен в Excel файл")

# Сохраняем Excel файл
wb.save(output_file)

# Выводим статистику
print(f"\nТаблица успешно создана: {output_file}")
print(f"Количество временных меток: {len(pivot_df)}")
print(f"Колонок с ненулевыми значениями: {pivot_df['Есть_ненулевые'].sum()}")

# Статистика по интервалам
if len(time_intervals_numeric) > 1:
    intervals_numeric = time_intervals_numeric[1:]  # Пропускаем первый (0)
    max_interval = max(intervals_numeric)
    min_interval = min(intervals_numeric)
    avg_interval = sum(intervals_numeric) / len(intervals_numeric)
    
    print(f"\nСтатистика временных интервалов:")
    print(f"  Максимальный интервал: {max_interval*1000:.1f} мс ({max_interval:.3f} сек)")
    print(f"  Минимальный интервал: {min_interval*1000:.1f} мс ({min_interval:.3f} сек)")
    print(f"  Средний интервал: {avg_interval*1000:.1f} мс ({avg_interval:.3f} сек)")
    
    long_intervals = [x for x in intervals_numeric if x > 0.1]
    print(f"  Интервалов > 100 мс: {len(long_intervals)}")

print("\nГрафик показывает:")
print("  • Время между событиями прямо на соединительных линиях")
print("  • Синие круги = значение 0")
print("  • Красные квадраты = значение ≠ 0")
print("  • Время на линиях в миллисекундах (мс) или секундах (с)")
print("  • Стрелки показывают направление времени")
print("  • Временные метки под первой и последней точкой каждого события")

# Показываем график
show_graph = input("\nПоказать график? (y/n): ").lower()
if show_graph == 'y':
    plt.show()

# Дополнительно: создаем текстовый файл со статистикой по каждому событию
print("\nСоздание детальной статистики...")
stats_filename = 'events_statistics.txt'
with open(stats_filename, 'w', encoding='utf-8') as f:
    f.write("СТАТИСТИКА ПО СОБЫТИЯМ\n")
    f.write("=" * 50 + "\n\n")
    
    for event_name, data in event_data.items():
        f.write(f"Событие: {event_name}\n")
        f.write(f"  Количество записей: {len(data['times'])}\n")
        
        if len(data['times']) > 1:
            first_time = data['times'][0]
            last_time = data['times'][-1]
            total_duration = (last_time - first_time).total_seconds()
            
            # Время между событиями
            if len(data['time_diffs']) > 1:
                actual_diffs = [d for d in data['time_diffs'] if d > 0]
                if actual_diffs:
                    avg_diff = sum(actual_diffs) / len(actual_diffs)
                    max_diff = max(actual_diffs)
                    f.write(f"  Среднее время между состояниями: {avg_diff*1000:.1f} мс\n")
                    f.write(f"  Макс. время между состояниями: {max_diff*1000:.1f} мс\n")
            
            f.write(f"  Общая продолжительность: {total_duration:.3f} сек\n")
        
        # Статистика по значениям
        nonzero_count = sum(1 for v in data['values'] if pd.notna(v) and v != 0)
        zero_count = len(data['values']) - nonzero_count
        f.write(f"  Состояний '0': {zero_count}\n")
        f.write(f"  Состояний '≠0': {nonzero_count}\n")
        
        if nonzero_count > 0:
            nonzero_values = [v for v in data['values'] if pd.notna(v) and v != 0]
            f.write(f"  Значения ≠0: {', '.join(map(str, nonzero_values))}\n")
        
        f.write("\n")

print(f"Детальная статистика сохранена в: {stats_filename}")