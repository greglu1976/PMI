# автоматическое тестирование ФСУ в части МТЗ, КЦН, ПС Т2
# с графическим интерфейсом

import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension

import itertools
import openpyxl
from lib._PARTS.MTZ_T2 import partOfFsuInTOC

class PartOfFsuInTOC_GUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ Т2 в части МТЗ, КЦН НН1, КЦН НН2, ЛО Т, CC, ПС, v2.0 от 26.03.25, v2.1 01.07.25, v2.2 25.07.25")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация переменных для SGF параметров, настроек, входных и выходных значений
        self.sgf_params = {
            "SGF1_lvttoc": tk.IntVar(value=0),
            "SGF1_ptoc1_lvttoc": tk.IntVar(value=1),
            "SGF2_ptoc1_lvttoc": tk.IntVar(value=0),
            "SGF3_ptoc1_lvttoc": tk.IntVar(value=0),
            "SGF4_ptoc1_lvttoc": tk.IntVar(value=0),
            "SGF5_ptoc1_lvttoc": tk.IntVar(value=1),
            "SGF6_ptoc1_lvttoc": tk.IntVar(value=0),
            "SGF7_ptoc1_lvttoc": tk.IntVar(value=0),
            "SGF1_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF2_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF3_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF4_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF5_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF6_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF7_ptoc2_lvttoc": tk.IntVar(value=0),
            "SGF1_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF2_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF3_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF4_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF5_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF6_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF7_ptoc3_lvttoc": tk.IntVar(value=0),
            "SGF1_ptuv1_lvttoc": tk.IntVar(value=0),
            "SGF1_ptuv2_lvttoc": tk.IntVar(value=0),
            "SGF1_phar1_lvttoc": tk.IntVar(value=0),
            "SGF1_rblc1_lvttoc": tk.IntVar(value=0),
            "SGF1_rvtr1_lvrbvtr1": tk.IntVar(value=0),
            "SGF2_rvtr1_lvrbvtr1": tk.IntVar(value=0),
            "SGF1_rvtr1_lvrbvtr2": tk.IntVar(value=0),
            "SGF2_rvtr1_lvrbvtr2": tk.IntVar(value=0),
            "SGF1_ptrc1_tresofflvlgc": tk.IntVar(value=0), 
            "SGF1_rbre1_tresofflvlgc": tk.IntVar(value=0),
            "SGF2_rbre1_tresofflvlgc": tk.IntVar(value=0),
            "SGF3_rbre1_tresofflvlgc": tk.IntVar(value=0),
            "SGF1_lvcbrblc1_tresofflvlgc": tk.IntVar(value=0),
            "SGF2_lvcbrblc1_tresofflvlgc": tk.IntVar(value=0),
            "SGF3_lvcbrblc1_tresofflvlgc": tk.IntVar(value=0),
            "SGF1_lvalh": tk.IntVar(value=0),
            "SGF2_lvalh": tk.IntVar(value=0),
            "SGF3_lvalh": tk.IntVar(value=0),
            "SGF4_lvalh": tk.IntVar(value=0),
            "SGF5_lvalh": tk.IntVar(value=0),
            "SGF6_lvalh": tk.IntVar(value=0),
            "SGF7_lvalh": tk.IntVar(value=0),
            "SGF8_lvalh": tk.IntVar(value=0),
            "SGF9_lvalh": tk.IntVar(value=0),
            "SGF10_lvalh": tk.IntVar(value=0),
            "SGF1_ptoc1_lvarctoc": tk.IntVar(value=0),
            "SGF2_ptoc1_lvarctoc": tk.IntVar(value=0),
            "SGF1_ptrc1_ttoclgc": tk.IntVar(value=0),
            "SGF2_ptrc1_ttoclgc": tk.IntVar(value=0),
            "SGF3_ptrc1_ttoclgc": tk.IntVar(value=0),              
            "Номинальный ток входа": tk.IntVar(value=5),            
        }
            #"SGF1_ptrc1_tofflvlgc":tk.IntVar(value=0),
            #"SGF1_rbre1_tofflvlgc":tk.IntVar(value=0),
            #"SGF2_rbre1_tofflvlgc":tk.IntVar(value=0),
            #"SGF3_rbre1_tofflvlgc":tk.IntVar(value=0), 
            #"SGF1_rblc1_tofflvlgc":tk.IntVar(value=0), 
            #"SGF2_rblc1_tofflvlgc":tk.IntVar(value=0),
            #"SGF3_rblc1_tofflvlgc":tk.IntVar(value=0),
        #}

        self.settings = {
            "T1_ptoc1_lvttoc": tk.DoubleVar(value=1),
            "Iset_ptoc1_lvttoc": tk.DoubleVar(value=0.2),
            "Icoarse_ptoc1_lvttoc": tk.DoubleVar(value=0.6),
            "T1_ptoc2_lvttoc": tk.DoubleVar(value=1),
            "Iset_ptoc2_lvttoc": tk.DoubleVar(value=0.2),
            "Icoarse_ptoc2_lvttoc": tk.DoubleVar(value=0.6),
            "T1_ptoc3_lvttoc": tk.DoubleVar(value=1),
            "Iset_ptoc3_lvttoc": tk.DoubleVar(value=0.2),
            "Icoarse_ptoc3_lvttoc": tk.DoubleVar(value=0.6),
            "Uop_ptuv1_lvttoc": tk.DoubleVar(value=50),
            "U2op_ptuv1_lvttoc": tk.DoubleVar(value=15),
            "Uop_ptuv2_lvttoc": tk.DoubleVar(value=50),
            "U2op_ptuv2_lvttoc": tk.DoubleVar(value=15),
            "Imax_phar1_lvttoc": tk.DoubleVar(value=1),
            "Ratio_phar1_lvttoc": tk.DoubleVar(value=40),
            "Umin_rvtr1_lvrbvtr1": tk.DoubleVar(value=50),
            "U2max_rvtr1_lvrbvtr1": tk.DoubleVar(value=15),
            "T1_rvtr1_lvrbvtr1": tk.DoubleVar(value=1),
            "Umin_rvtr1_lvrbvtr2": tk.DoubleVar(value=50),
            "U2max_rvtr1_lvrbvtr2": tk.DoubleVar(value=15),
            "T1_rvtr1_lvrbvtr2": tk.DoubleVar(value=1),
            "Iset_ptoc1_lvarctoc": tk.DoubleVar(value=1),
            "T1_ptrc1_ttoclgc": tk.DoubleVar(value=1),             
        }

        self.input_vars = {
            "VYVOD": tk.IntVar(value=0),
            "OV_lvttoc": tk.IntVar(value=0),
            "OVst_ptoc1": tk.IntVar(value=0),
            "OVst_ptoc2": tk.IntVar(value=0),
            "OVst_ptoc3": tk.IntVar(value=0),
            "NaSign_ptoc1": tk.IntVar(value=0),
            "NaSign_ptoc2": tk.IntVar(value=0),
            "NaSign_ptoc3": tk.IntVar(value=0),
            "SV1vkl": tk.IntVar(value=0),
            "SV2vkl": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=1),
            "dIA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=1),
            "dIB": tk.DoubleVar(value=240),
            "IC": tk.DoubleVar(value=1),
            "dIC": tk.DoubleVar(value=120),
            "UA1": tk.DoubleVar(value=50),
            "dUA1": tk.DoubleVar(value=0),
            "UB1": tk.DoubleVar(value=50),
            "dUB1": tk.DoubleVar(value=240),
            "UC1": tk.DoubleVar(value=50),
            "dUC1": tk.DoubleVar(value=120),
            "UA2": tk.DoubleVar(value=50),
            "dUA2": tk.DoubleVar(value=0),
            "UB2": tk.DoubleVar(value=50),
            "dUB2": tk.DoubleVar(value=240),
            "UC2": tk.DoubleVar(value=50),
            "dUC2": tk.DoubleVar(value=120),
            "KPONvnesh_ptuv1": tk.IntVar(value=0),
            "KPONvnesh_ptuv2": tk.IntVar(value=0),
            "IA2harm": tk.DoubleVar(value=0),
            "IB2harm": tk.DoubleVar(value=0),
            "IC2harm": tk.DoubleVar(value=0),
            "VNN1vkl": tk.IntVar(value=0),
            "VNN2vkl": tk.IntVar(value=0),
            "OV_lvrbvtr1": tk.IntVar(value=0),
            "vnesh_bnn_srab_lvrbvtr1": tk.IntVar(value=0),
            "OV_lvrbvtr2": tk.IntVar(value=0),
            "vnesh_bnn_srab_lvrbvtr2": tk.IntVar(value=0),
            "OVlot": tk.IntVar(value=0),
            "OVlo": tk.IntVar(value=0),
            "OVzapv": tk.IntVar(value=0),
            "OVzavr": tk.IntVar(value=0),
        }

        self.output_labels = {}

        # Создание интерфейса
        self.create_widgets()

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            if key=="SGF7_ptoc1_lvttoc" or key=="SGF7_ptoc2_lvttoc" or key=="SGF7_ptoc3_lvttoc" or key=="SGF1_ptuv1_lvttoc" or key=="SGF1_ptuv2_lvttoc":
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            elif key=="SGF1_rblc1_lvttoc":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2, 3], state="readonly").grid(row=row, column=col + 1)
            elif key=='Номинальный ток входа':
                ttk.Combobox(sgf_frame, textvariable=var, values=[1, 5], state="readonly").grid(row=row, column=col + 1)                   
            else:
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 11:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)

        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)

        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)

        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)

        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)        


        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)

        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(input_frame, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(input_frame, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        outputs = [
            "vvod_lvrbvtr1", "oper_vyvod_lvrbvtr1", "u_lin_pusk_lvrbvtr1", "u2_pusk_lvrbvtr1", "pusk_lvrbvtr1", "neispr_zn_lvrbvtr1",
            "vvod_lvrbvtr2", "oper_vyvod_lvrbvtr2", "u_lin_pusk_lvrbvtr2", "u2_pusk_lvrbvtr2", "pusk_lvrbvtr2", "neispr_zn_lvrbvtr2",
            "vvod_ptoc1_lvttoc", "oper_vyvod_ptoc1_lvttoc", "mtzA_pusk_ptoc1_lvttoc", "mtzB_pusk_ptoc1_lvttoc", "mtzC_pusk_ptoc1_lvttoc",
            "gen_pusk_ptoc_lvttoc", "mtz_srabsign_ptoc1_lvttoc", "mtz_srab_ptoc1_lvttoc", "io_A_ptoc1_lvttoc", "io_B_ptoc1_lvttoc",
            "io_C_ptoc1_lvttoc", "vvod_ptoc2_lvttoc", "oper_vyvod_ptoc2_lvttoc", "mtzA_pusk_ptoc2_lvttoc", "mtzB_pusk_ptoc2_lvttoc",
            "mtzC_pusk_ptoc2_lvttoc", "gen_pusk_ptoc2_lvttoc", "mtz_srabsign_ptoc2_lvttoc", "mtz_srab_ptoc2_lvttoc", "io_A_ptoc2_lvttoc",
            "io_B_ptoc2_lvttoc", "io_C_ptoc2_lvttoc", "vvod_ptoc3_lvttoc", "oper_vyvod_ptoc3_lvttoc", "mtzA_pusk_ptoc3_lvttoc",
            "mtzB_pusk_ptoc3_lvttoc", "mtzC_pusk_ptoc3_lvttoc", "gen_pusk_ptoc3_lvttoc", "mtz_srabsign_ptoc3_lvttoc", "mtz_srab_ptoc3_lvttoc",
            "io_A_ptoc3_lvttoc", "io_B_ptoc3_lvttoc", "io_C_ptoc3_lvttoc", "kpon_pusk_ptuv1_lvttoc", "kpon_pusk_ptuv2_lvttoc",
            "ia_start_out_phar1_lvttoc", "ib_start_out_phar1_lvttoc", "ic_start_out_phar1_lvttoc", "start_phar1_lvttoc", "blok_rblc1_lvttoc",
            "mtz_pusk_lvttoc", "vvod_ptrc1", "oper_vyvod_ptrc1", "pusk_ptrc1", "srab_ptrc1", "vvod_rblc1", "oper_vyvod_rblc1", "zapret_rblc1",
            "vvod_rbre1", "oper_vyvod_rbre1", "zapret_rbre1", "pusk_lvalh", 
            "pusk_ptoc1_lvarctoc", 
            "pusk_ptrc1_ttoclgc", 
            "IAB", "dIAB", "IBC", "dIBC", "ICA", "dICA", "I2", "I0", "I1", "UAB_ptuv1", "UBC_ptuv1", "UCA_ptuv1", "U2_ptuv1", "U0_ptuv1", "U1_ptuv1", "UAB_ptuv2", "UBC_ptuv2", "UCA_ptuv2", "U2_ptuv2",  "U0_ptuv2", "U1_ptuv2"
        ]

        row = 0
        col = 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=25, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = partOfFsuInTOC(
            SGF1=self.sgf_params["SGF1_lvttoc"].get(),
            SGF1_ptoc1=self.sgf_params["SGF1_ptoc1_lvttoc"].get(),
            SGF2_ptoc1=self.sgf_params["SGF2_ptoc1_lvttoc"].get(),
            SGF3_ptoc1=self.sgf_params["SGF3_ptoc1_lvttoc"].get(),
            SGF4_ptoc1=self.sgf_params["SGF4_ptoc1_lvttoc"].get(),
            SGF5_ptoc1=self.sgf_params["SGF5_ptoc1_lvttoc"].get(),
            SGF6_ptoc1=self.sgf_params["SGF6_ptoc1_lvttoc"].get(),
            SGF7_ptoc1=self.sgf_params["SGF7_ptoc1_lvttoc"].get(),
            T1_ptoc1=self.settings["T1_ptoc1_lvttoc"].get(),
            Iset_ptoc1=self.settings["Iset_ptoc1_lvttoc"].get(),
            Icoarse_ptoc1=self.settings["Icoarse_ptoc1_lvttoc"].get(),
            SGF1_ptoc2=self.sgf_params["SGF1_ptoc2_lvttoc"].get(),
            SGF2_ptoc2=self.sgf_params["SGF2_ptoc2_lvttoc"].get(),
            SGF3_ptoc2=self.sgf_params["SGF3_ptoc2_lvttoc"].get(),
            SGF4_ptoc2=self.sgf_params["SGF4_ptoc2_lvttoc"].get(),
            SGF5_ptoc2=self.sgf_params["SGF5_ptoc2_lvttoc"].get(),
            SGF6_ptoc2=self.sgf_params["SGF6_ptoc2_lvttoc"].get(),
            SGF7_ptoc2=self.sgf_params["SGF7_ptoc2_lvttoc"].get(),
            T1_ptoc2=self.settings["T1_ptoc2_lvttoc"].get(),
            Iset_ptoc2=self.settings["Iset_ptoc2_lvttoc"].get(),
            Icoarse_ptoc2=self.settings["Icoarse_ptoc2_lvttoc"].get(),
            SGF1_ptoc3=self.sgf_params["SGF1_ptoc3_lvttoc"].get(),
            SGF2_ptoc3=self.sgf_params["SGF2_ptoc3_lvttoc"].get(),
            SGF3_ptoc3=self.sgf_params["SGF3_ptoc3_lvttoc"].get(),
            SGF4_ptoc3=self.sgf_params["SGF4_ptoc3_lvttoc"].get(),
            SGF5_ptoc3=self.sgf_params["SGF5_ptoc3_lvttoc"].get(),
            SGF6_ptoc3=self.sgf_params["SGF6_ptoc3_lvttoc"].get(),
            SGF7_ptoc3=self.sgf_params["SGF7_ptoc3_lvttoc"].get(),
            T1_ptoc3=self.settings["T1_ptoc3_lvttoc"].get(),
            Iset_ptoc3=self.settings["Iset_ptoc3_lvttoc"].get(),
            Icoarse_ptoc3=self.settings["Icoarse_ptoc3_lvttoc"].get(),
            SGF1_ptuv1=self.sgf_params["SGF1_ptuv1_lvttoc"].get(),
            Uop_ptuv1=self.settings["Uop_ptuv1_lvttoc"].get(),
            U2op_ptuv1=self.settings["U2op_ptuv1_lvttoc"].get(),
            SGF1_ptuv2=self.sgf_params["SGF1_ptuv2_lvttoc"].get(),
            Uop_ptuv2=self.settings["Uop_ptuv2_lvttoc"].get(),
            U2op_ptuv2=self.settings["U2op_ptuv2_lvttoc"].get(),
            SGF1_phar1=self.sgf_params["SGF1_phar1_lvttoc"].get(),
            Imax_phar1=self.settings["Imax_phar1_lvttoc"].get(),
            Ratio_phar1=self.settings["Ratio_phar1_lvttoc"].get(),
            SGF1_rblc1=self.sgf_params["SGF1_rblc1_lvttoc"].get(),
            SGF1_lvrbvtr1=self.sgf_params["SGF1_rvtr1_lvrbvtr1"].get(),
            SGF2_lvrbvtr1=self.sgf_params["SGF2_rvtr1_lvrbvtr1"].get(),
            u_min_lvrbvtr1=self.settings["Umin_rvtr1_lvrbvtr1"].get(),
            u2_max_lvrbvtr1=self.settings["U2max_rvtr1_lvrbvtr1"].get(),
            t1_lvrbvtr1=self.settings["T1_rvtr1_lvrbvtr1"].get(),
            SGF1_lvrbvtr2=self.sgf_params["SGF1_rvtr1_lvrbvtr2"].get(),
            SGF2_lvrbvtr2=self.sgf_params["SGF2_rvtr1_lvrbvtr2"].get(),
            u_min_lvrbvtr2=self.settings["Umin_rvtr1_lvrbvtr2"].get(),
            u2_max_lvrbvtr2=self.settings["U2max_rvtr1_lvrbvtr2"].get(),
            t1_lvrbvtr2=self.settings["T1_rvtr1_lvrbvtr2"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["SGF1_ptrc1_tresofflvlgc"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["SGF1_rbre1_tresofflvlgc"].get(), 
            SGF2_rbre1_tofflvlgc=self.sgf_params["SGF2_rbre1_tresofflvlgc"].get(), 
            SGF3_rbre1_tofflvlgc=self.sgf_params["SGF3_rbre1_tresofflvlgc"].get(), 
            SGF1_rblc1_tofflvlgc=self.sgf_params["SGF1_lvcbrblc1_tresofflvlgc"].get(), 
            SGF2_rblc1_tofflvlgc=self.sgf_params["SGF2_lvcbrblc1_tresofflvlgc"].get(), 
            SGF3_rblc1_tofflvlgc=self.sgf_params["SGF3_lvcbrblc1_tresofflvlgc"].get(),
            SGF1_ptoc1_lvarctoc=self.sgf_params["SGF1_ptoc1_lvarctoc"].get(),
            SGF2_ptoc1_lvarctoc=self.sgf_params["SGF2_ptoc1_lvarctoc"].get(),
            SGF1_ptrc1_ttoclgc=self.sgf_params["SGF1_ptrc1_ttoclgc"].get(),
            SGF2_ptrc1_ttoclgc=self.sgf_params["SGF2_ptrc1_ttoclgc"].get(),
            SGF3_ptrc1_ttoclgc=self.sgf_params["SGF3_ptrc1_ttoclgc"].get(),
            Iset_ptoc1_lvarctoc=self.settings["Iset_ptoc1_lvarctoc"].get(),
            Inom=self.sgf_params["Номинальный ток входа"].get(),
            T1_ptrc1_ttoclgc=self.settings["T1_ptrc1_ttoclgc"].get() 
        )

        print("partOfFsuInTOC initialized")

    def start_polling(self):
        if self.part is None:
            print("partOfFsuInTOC not initialized")
            return

        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            #print(inputs)
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #label.config(text=f"{output}: {int(value)}")
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or float(value)!=0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3) # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05) # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):

        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return

        output_file = f"{function}_{mode}.xlsx"


        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        #output_file = "data.xlsx"
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)

        # Определяем красный цвет для заливки
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                # Устанавливаем ширину столбца
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20

                # Проверяем значения и применяем форматирование
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)  # Преобразуем значение в число
                        if value != 0:
                            cell.fill = red_fill  # Выделяем красным, если значение не равно 0
                    except (ValueError, TypeError):
                        pass  # Игнорируем ошибки преобразования

        # Применяем форматирование к каждому листу
        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        # Сохраняем изменения
        wb.save(output_file)

        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        # Выбор файла для загрузки
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        try:
            # Чтение данных из Excel
            xls = pd.ExcelFile(file_path)

            # Загрузка SGF Parameters
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            # Загрузка Inputs
            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])                    

            print("Data loaded successfully")

        except Exception as e:
            print(f"Error loading data: {e}")        


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfFsuInTOC_GUI(root)
    root.mainloop()