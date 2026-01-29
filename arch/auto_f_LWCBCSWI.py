import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension

import itertools
import openpyxl
from lib._FUNCS.LWCBCSWI import LWCBCSWI  # Импортируем новый класс

class LWCBCSWI_GUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование функции УВ в составе КП")
        self.lwcbcswi = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация переменных для SGF параметров
        self.sgf_params = {
            "SGF1": tk.IntVar(value=0),  # Ввод функции в работу
            "SGF2": tk.IntVar(value=0),  # Блокировка включения от аварийного отключения В
        }

        # Инициализация переменных для настроек таймеров
        self.settings = {
            "T1": tk.DoubleVar(value=1),  # Тперекл (Допустимое время переключения)
            "T2": tk.DoubleVar(value=2),  # Тблк (Выдержка времени подавления выдачи положения "Не определено")
            "T3": tk.DoubleVar(value=0.5),  # Таймер 3
            "T4": tk.DoubleVar(value=0.5),  # Таймер 4
        }

        # Инициализация переменных для входных значений
        self.input_vars = {
            "VVOD": tk.IntVar(value=0),  # Ввод
            "KSV_blk_otkl": tk.IntVar(value=0),  # Блокировка отключения КСВ
            "otkl_v_ot_pu": tk.IntVar(value=0),  # Отключение В от ПУ
            "otkl_v_ichm": tk.IntVar(value=0),  # Отключение В от ИЧМ
            "mestnoe": tk.IntVar(value=0),  # Местное управление
            "otkl_v_ot_tu": tk.IntVar(value=0),  # Отключение В от ТУ
            "otkl_v_asu": tk.IntVar(value=0),  # Отключение В от АСУ
            "kluch_md_priv": tk.IntVar(value=0),  # Ключ МД привода
            "vkl_v_ot_pu": tk.IntVar(value=0),  # Включение В от ПУ
            "vkl_v_ichm": tk.IntVar(value=0),  # Включение В от ИЧМ
            "distanz": tk.IntVar(value=0),  # Дистанционное управление
            "vkl_v_ot_tu": tk.IntVar(value=0),  # Включение В от ТУ
            "vkl_v_asu": tk.IntVar(value=0),  # Включение В от АСУ
            "ksv_blok_vkl": tk.IntVar(value=0),  # Блокировка включения КСВ
            "ksv_v_avar_otkl": tk.IntVar(value=0),  # Аварийное отключение КСВ
            "v_otkl_bk": tk.IntVar(value=0),  # В отключено (БК)
            "v_vkl_bk": tk.IntVar(value=0),  # В включено (БК)
        }

        # Инициализация переменных для выходных значений
        self.output_labels = {}

        # Создание интерфейса
        self.create_widgets()

    def create_widgets(self):
        # Frame для SGF параметров
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1

        # Frame для настроек таймеров
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1

        # Frame для кнопок
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Кнопка Init
        ttk.Button(buttons_frame, text="Init", command=self.init_lwcbcswi).grid(row=0, column=0, pady=10)

        # Кнопка Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)

        # Кнопка Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)

        # Кнопка Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)

        # Кнопка Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)

        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Frame для входных значений
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.input_vars.items():
            ttk.Checkbutton(input_frame, text=key, variable=var).grid(row=row, column=col, sticky="w")
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame для выходных значений
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        outputs = [
            "vvod", "uv_otkluchit", "uv_idet_per", "uv_prev_vrem_per", "uv_vkluchit",
            "uv_ne_opredeleno", "uv_otklucheno", "uv_vklucheno", "uv_neispr_neopred",
            "ET_t1", "ET_t21", "ET_t22", "ET_t3", "ET_t4", "_p001", "_p002", "_p003",
            "_p004", "_p005", "_p006"
        ]

        row = 0
        col = 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=20, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 10:
                row = 0
                col += 2

    def init_lwcbcswi(self):
        self.lwcbcswi = LWCBCSWI(
            SGF1=self.sgf_params["SGF1"].get(),
            SGF2=self.sgf_params["SGF2"].get(),
            T1=self.settings["T1"].get(),
            T2=self.settings["T2"].get(),
            T3=self.settings["T3"].get(),
            T4=self.settings["T4"].get(),
        )
        print("LWCBCSWI initialized")

    def start_polling(self):
        if self.lwcbcswi is None:
            print("LWCBCSWI not initialized")
            return

        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.lwcbcswi.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0:
                    label.config(background="red")
                else:
                    label.config(background="green")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="red", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="green", foreground="white")

    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return

        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        # Применяем форматирование к каждому листу
        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        # Сохраняем изменения
        wb.save(output_file)

        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        # Выбор файла для загрузки
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        try:
            # Чтение данных из Excel
            xls = pd.ExcelFile(file_path)

            # Загрузка SGF Parameters
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            # Загрузка Inputs
            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")

        except Exception as e:
            print(f"Error loading data: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = LWCBCSWI_GUI(root)
    root.mainloop()