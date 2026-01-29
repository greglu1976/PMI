# автоматическое тестирование ФСУ в части ДЗТ, КЦТ, ПС для исполнения ДЗТ
# с графическим интерфейсом
# 2, 5 заводятся не в процентах а в rms - угол 2,5 гармоник для упрощения совпадает с 1 гармоникой
# для упрощения также 2,5 гармоники только на сторону ВН (НН1, НН2 - нет)

import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension

import itertools
import openpyxl
from lib._PARTS.DZT import partDZT

import numpy as np

class PartDZT_GUI:

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ (исполнение ДЗТ2) в части КЦТ, ДЗТ, ЛО Т, ПС, v1.3 от 05.05.25")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # Инициализация переменных для SGF параметров, настроек, входных и выходных значений
        self.sgf_params = {
            "Side1_tdif": tk.IntVar(value=1),
            "Side2_tdif": tk.IntVar(value=1),
            "Side3_tdif": tk.IntVar(value=1),
            "ksch1_tdif": tk.IntVar(value=0),
            "ksch2_tdif": tk.IntVar(value=0), 
            "ksch3_tdif": tk.IntVar(value=0),           # Третья сторона 03.05.2025
            "comp3i0_rmxu1_tdif": tk.IntVar(value=0),
            "comp3i0_rmxu2_tdif": tk.IntVar(value=0),
            "comp3i0_rmxu3_tdif": tk.IntVar(value=0),   # Третья сторона  03.05.2025         
            "SGF1_rctr1_ctr": tk.IntVar(value=0),
            "SGF2_rctr1_ctr": tk.IntVar(value=0),
            "SGF1_rctr2_ctr": tk.IntVar(value=0),
            "SGF2_rctr2_ctr": tk.IntVar(value=0),
            "SGF1_rctr3_ctr": tk.IntVar(value=0),
            "SGF2_rctr3_ctr": tk.IntVar(value=0),
            "SGF1_pdif1_tdif": tk.IntVar(value=0),
            "SGF2_pdif1_tdif": tk.IntVar(value=1),
            "SGF3_pdif1_tdif": tk.IntVar(value=0),
            "SGF1_pdif2_tdif": tk.IntVar(value=0),
            "SGF1_hf2phar1_tdif": tk.IntVar(value=0),
            "SGF1_hf5phar1_tdif": tk.IntVar(value=0),
            "SGF1_rctr1_tdif": tk.IntVar(value=0),
            "SGF1_ptrc1_tprmofflvlgc": tk.IntVar(value=0),
            "SGF1_rbre1_tprmofflvlgc": tk.IntVar(value=0),
        }
        self.settings = {
            "T1_rctr1_ctr": tk.DoubleVar(value=0),
            "T2_rctr1_ctr": tk.DoubleVar(value=0),
            "Inom_rctr1_ctr": tk.DoubleVar(value=5),
            "Imin_rctr1_ctr": tk.DoubleVar(value=0.05),
            "Ksym_rctr1_ctr": tk.DoubleVar(value=0.5),
            "LIsym_rctr1_ctr": tk.DoubleVar(value=0.02),
            "T1_rctr2_ctr": tk.DoubleVar(value=0),
            "T2_rctr2_ctr": tk.DoubleVar(value=0),
            "Inom_rctr2_ctr": tk.DoubleVar(value=5),
            "Imin_rctr2_ctr": tk.DoubleVar(value=0.05),
            "Ksym_rctr2_ctr": tk.DoubleVar(value=0.5),
            "LIsym_rctr2_ctr": tk.DoubleVar(value=0.02),
            "T1_rctr3_ctr": tk.DoubleVar(value=0),
            "T2_rctr3_ctr": tk.DoubleVar(value=0),
            "Inom_rctr3_ctr": tk.DoubleVar(value=5),
            "Imin_rctr3_ctr": tk.DoubleVar(value=0.05),
            "Ksym_rctr3_ctr": tk.DoubleVar(value=0.5),
            "LIsym_rctr3_ctr": tk.DoubleVar(value=0.02),
            "Sbaz_tdif": tk.DoubleVar(value=10),
            "Ubaz_rmxu1_tdif": tk.DoubleVar(value=35),
            "Ubaz_rmxu2_tdif": tk.DoubleVar(value=10.5),
            "Ubaz_rmxu3_tdif": tk.DoubleVar(value=10.5),  # Третья сторона  03.05.2025 
            "Iperv_rmxu1_tdif": tk.DoubleVar(value=750),
            "Iperv_rmxu2_tdif": tk.DoubleVar(value=3000),
            "Iperv_rmxu3_tdif": tk.DoubleVar(value=3000), # Третья сторона  03.05.2025 
            "Inomterm_rmxu1_tdif": tk.DoubleVar(value=1),
            "Inomterm_rmxu2_tdif": tk.DoubleVar(value=1),
            "Inomterm_rmxu3_tdif": tk.DoubleVar(value=1), # Третья сторона  03.05.2025 
            "Ivtor_rmxu1_tdif": tk.DoubleVar(value=5),
            "Ivtor_rmxu2_tdif": tk.DoubleVar(value=5),
            "Ivtor_rmxu3_tdif": tk.DoubleVar(value=5),    # Третья сторона  03.05.2025
            "Nsch_rmxu1_tdif": tk.DoubleVar(value=0),
            "Nsch_rmxu2_tdif": tk.DoubleVar(value=6),
            "Nsch_rmxu3_tdif": tk.DoubleVar(value=6),    # Третья сторона  03.05.2025
            "T1_pdif1_tdif": tk.DoubleVar(value=1),
            "Isr_pdif1_tdif": tk.DoubleVar(value=0.2),
            "Isrzagrub_pdif1_tdif": tk.DoubleVar(value=1.2),
            "It1_pdif1_tdif": tk.DoubleVar(value=1),
            "It2_pdif1_tdif": tk.DoubleVar(value=3),
            "Kt1_pdif1_tdif": tk.DoubleVar(value=0.25),
            "Kt2_pdif1_tdif": tk.DoubleVar(value=0.7), 
            "T1_pdif2_tdif": tk.DoubleVar(value=1),
            "Iset_pdif2_tdif": tk.DoubleVar(value=4),
            "T1_hf2phar1_tdif": tk.DoubleVar(value=0.5),
            "T2_hf2phar1_tdif": tk.DoubleVar(value=0.5),
            "Ratio_hf2phar1_tdif": tk.DoubleVar(value=0.2),
            "T1_hf5phar1_tdif": tk.DoubleVar(value=0.5),
            "T2_hf5phar1_tdif": tk.DoubleVar(value=0.5),
            "Ratio_hf5phar1_tdif": tk.DoubleVar(value=0.3), 
            "T1_rctr1_tdif": tk.DoubleVar(value=1),
            "Iset_rctr1_tdif": tk.DoubleVar(value=0.1), 
        }

        self.input_vars = {
            "VYVOD": tk.IntVar(value=0),
            "OV_ctr": tk.IntVar(value=0),
            "OVst_rctr1": tk.IntVar(value=0),
            "OVst_rctr2": tk.IntVar(value=0),
            "OVst_rctr3": tk.IntVar(value=0),
            "OV_tdif": tk.IntVar(value=0),
            "OV_pdif1_tdif": tk.IntVar(value=0),
            "NaSign_pdif1_tdif": tk.IntVar(value=0),
            "OV_pdif2_tdif": tk.IntVar(value=0),
            "NaSign_pdif2_tdif": tk.IntVar(value=0),
            "OV_rctr1_tdif": tk.IntVar(value=0),
            "OV_tprmofflvlgc": tk.IntVar(value=0),
            "OV_ptrc1_tprmofflvlgc": tk.IntVar(value=0),
            "OV_rbre1_tprmofflvlgc": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "dIA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "dIB": tk.DoubleVar(value=240),
            "IC": tk.DoubleVar(value=0),
            "dIC": tk.DoubleVar(value=120),
            "IA1": tk.DoubleVar(value=0),
            "dIA1": tk.DoubleVar(value=0),
            "IB1": tk.DoubleVar(value=0),
            "dIB1": tk.DoubleVar(value=240),
            "IC1": tk.DoubleVar(value=0),
            "dIC1": tk.DoubleVar(value=120),
            "IA2harm": tk.DoubleVar(value=0),
            "IB2harm": tk.DoubleVar(value=0),
            "IC2harm": tk.DoubleVar(value=0),
            "IA5harm": tk.DoubleVar(value=0),
            "IB5harm": tk.DoubleVar(value=0),
            "IC5harm": tk.DoubleVar(value=0),
        }

        self.output_labels = {}

        # Создание интерфейса
        self.create_widgets()

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            ttk.Label(sgf_frame, text=key).grid(row=row, column=col, sticky="w")
            if key=="SGF6_ptoc1_lvttoc" or key=="SGF6_ptoc2_lvttoc" or key=="SGF6_ptoc3_lvttoc" or key=="SGF1_ptuv1_lvttoc" or key=="SGF1_ptuv2_lvttoc":
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            elif key=="SGF2_pdif1_tdif":
                 ttk.Combobox(sgf_frame, textvariable=var, values=[1, 2, 3, 4], state="readonly").grid(row=row, column=col + 1)
            elif key=='Номинальный ток входа':
                ttk.Combobox(sgf_frame, textvariable=var, values=[1, 5], state="readonly").grid(row=row, column=col + 1)                   
            else:
                ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.settings.items():
            ttk.Label(settings_frame, text=key).grid(row=row, column=col, sticky="w")
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)

        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)

        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)

        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)

        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)        


        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)

        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")

        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                ttk.Checkbutton(input_frame, text=key, variable=var).grid(row=row, column=col, sticky="w")
            elif isinstance(var, tk.DoubleVar):
                ttk.Label(input_frame, text=key).grid(row=row, column=col, sticky="w")
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 4:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        outputs = [
            "vvod_rctr1_ctr", "oper_vyvod_rctr1_ctr", "pusk_obryv_rctr1_ctr", "srab_obryv_rctr1_ctr", "pusk_assym_rctr1_ctr", "srab_assym_rctr1_ctr", "vvod_rctr2_ctr", "oper_vyvod_rctr2_ctr", "pusk_obryv_rctr2_ctr", "srab_obryv_rctr2_ctr", "pusk_assym_rctr2_ctr", "srab_assym_rctr2_ctr", "vvod_rctr3_ctr", "oper_vyvod_rctr3_ctr", "pusk_obryv_rctr3_ctr", "srab_obryv_rctr3_ctr", "pusk_assym_rctr3_ctr", "srab_assym_rctr3_ctr", "srab_ctr", "vvod_pdif2_tdif", "oper_vyvod_pdif2_tdif", "pusk_A_pdif2_tdif", "srab_A_pdif2_tdif", "srabsign_A_pdif2_tdif", "io_A_pdif2_tdif", "pusk_B_pdif2_tdif", "srab_B_pdif2_tdif", "srabsign_B_pdif2_tdif", "io_B_pdif2_tdif", "pusk_C_pdif2_tdif", "srab_C_pdif2_tdif", "srabsign_C_pdif2_tdif", "io_C_pdif2_tdif", "pusk_pdif2_tdif", "srabsign_pdif2_tdif", "srab_pdif2_tdif", "vvod_pdif1_tdif", "oper_vyvod_pdif1_tdif", "pusk_A_pdif1_tdif", "srab_A_pdif1_tdif", "srabsign_A_pdif1_tdif", "io_A_pdif1_tdif", "pusk_B_pdif1_tdif", "srab_B_pdif1_tdif", "srabsign_B_pdif1_tdif", "io_B_pdif1_tdif", "pusk_C_pdif1_tdif", "srab_C_pdif1_tdif", "srabsign_C_pdif1_tdif", "io_C_pdif1_tdif", "pusk_pdif1_tdif", "srabsign_pdif1_tdif", "srab_pdif1_tdif", "pusk_A_hf2phar1_tdif", "pusk_B_hf2phar1_tdif", "pusk_C_hf2phar1_tdif", "pusk_hf2phar1_tdif", "pusk_A_hf5phar1_tdif", "pusk_B_hf5phar1_tdif", "pusk_C_hf5phar1_tdif", "pusk_hf5phar1_tdif", "vvod_rctr1_tdif", "oper_vyvod_rctr1_tdif", "srab_A_rctr1_tdif", "srab_B_rctr1_tdif", "srab_C_rctr1_tdif", "srab_rctr1_tdif", "neispr_rctr1_tdif", "vvod_ptrc1_tprmofflvlgc", "oper_vyvod_ptrc1_tprmofflvlgc", "pusk_ptrc1_tprmofflvlgc", "srab_ptrc1_tprmofflvlgc", "vvod_rbre1_tprmofflvlgc", "oper_vyvod_rbre1_tprmofflvlgc", "zapret_rbre1_tprmofflvlgc", "pusk_lvalh", "diffA", "restA", "diffB", "restB", "diffC", "restC"
        ]

        row = 0
        col = 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=25, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            row += 1
            if row >= 28:
                row = 0
                col += 2

    def init_part(self):
        self.part = partDZT(
            Side1=self.sgf_params["Side1_tdif"].get(),
            Side2=self.sgf_params["Side2_tdif"].get(),
            Side3=self.sgf_params["Side3_tdif"].get(),                        
            SGF1_rctr1_ctr=self.sgf_params["SGF1_rctr1_ctr"].get(),
            SGF2_rctr1_ctr=self.sgf_params["SGF2_rctr1_ctr"].get(),
            SGF1_rctr2_ctr=self.sgf_params["SGF1_rctr2_ctr"].get(),
            SGF2_rctr2_ctr=self.sgf_params["SGF2_rctr2_ctr"].get(),
            SGF1_rctr3_ctr=self.sgf_params["SGF1_rctr3_ctr"].get(),
            SGF2_rctr3_ctr=self.sgf_params["SGF2_rctr3_ctr"].get(),
            SGF1_pdif1_tdif=self.sgf_params["SGF1_pdif1_tdif"].get(),
            SGF2_pdif1_tdif=self.sgf_params["SGF2_pdif1_tdif"].get(),
            SGF3_pdif1_tdif=self.sgf_params["SGF3_pdif1_tdif"].get(),
            SGF1_pdif2_tdif=self.sgf_params["SGF1_pdif2_tdif"].get(),
            SGF1_hf2phar1_tdif=self.sgf_params["SGF1_hf2phar1_tdif"].get(),
            SGF1_hf5phar1_tdif=self.sgf_params["SGF1_hf5phar1_tdif"].get(),
            SGF1_rctr1_tdif=self.sgf_params["SGF1_rctr1_tdif"].get(),
            SGF1_ptrc1_tprmofflvlgc=self.sgf_params["SGF1_ptrc1_tprmofflvlgc"].get(),
            SGF1_rbre1_tprmofflvlgc=self.sgf_params["SGF1_rbre1_tprmofflvlgc"].get(),
            k_sch_vn = self.sgf_params["ksch1_tdif"].get(),
            k_sch_nn = self.sgf_params["ksch2_tdif"].get(),
            k_sch_nn2 = self.sgf_params["ksch3_tdif"].get(),                        # Третья сторона
            compens_3i0_vn = self.sgf_params["comp3i0_rmxu1_tdif"].get(),
            compens_3i0_nn = self.sgf_params["comp3i0_rmxu2_tdif"].get(),
            compens_3i0_nn2 = self.sgf_params["comp3i0_rmxu3_tdif"].get(),          # Третья сторона
            T1_rctr1_ctr = self.settings["T1_rctr1_ctr"].get(),
            T2_rctr1_ctr = self.settings["T2_rctr1_ctr"].get(),
            Inom_rctr1_ctr = self.settings["Inom_rctr1_ctr"].get(),
            Imin_rctr1_ctr = self.settings["Imin_rctr1_ctr"].get(),
            Ksym_rctr1_ctr = self.settings["Ksym_rctr1_ctr"].get(),
            LIsym_rctr1_ctr = self.settings["LIsym_rctr1_ctr"].get(),

            T1_rctr2_ctr = self.settings["T1_rctr2_ctr"].get(),
            T2_rctr2_ctr = self.settings["T2_rctr2_ctr"].get(),
            Inom_rctr2_ctr = self.settings["Inom_rctr2_ctr"].get(),
            Imin_rctr2_ctr = self.settings["Imin_rctr2_ctr"].get(),
            Ksym_rctr2_ctr = self.settings["Ksym_rctr2_ctr"].get(),
            LIsym_rctr2_ctr = self.settings["LIsym_rctr2_ctr"].get(),

            T1_rctr3_ctr = self.settings["T1_rctr3_ctr"].get(),
            T2_rctr3_ctr = self.settings["T2_rctr3_ctr"].get(),
            Inom_rctr3_ctr = self.settings["Inom_rctr3_ctr"].get(),
            Imin_rctr3_ctr = self.settings["Imin_rctr3_ctr"].get(),
            Ksym_rctr3_ctr = self.settings["Ksym_rctr3_ctr"].get(),
            LIsym_rctr3_ctr = self.settings["LIsym_rctr3_ctr"].get(),

            Sbaz = self.settings["Sbaz_tdif"].get()*1e+6,
            Ubaz_vn = self.settings["Ubaz_rmxu1_tdif"].get()*1e+3,
            Ubaz_nn = self.settings["Ubaz_rmxu2_tdif"].get()*1e+3,
            Ubaz_nn2 = self.settings["Ubaz_rmxu3_tdif"].get()*1e+3,            # Третья сторона
            Iperv_vn = self.settings["Iperv_rmxu1_tdif"].get(),
            Iperv_nn = self.settings["Iperv_rmxu2_tdif"].get(),
            Iperv_nn2 = self.settings["Iperv_rmxu3_tdif"].get(),               # Третья сторона          
            Inom_term_vn = self.settings["Inomterm_rmxu1_tdif"].get(),
            Inom_term_nn = self.settings["Inomterm_rmxu2_tdif"].get(),
            Inom_term_nn2 = self.settings["Inomterm_rmxu3_tdif"].get(),        # Третья сторона            
            Ivtor_vn = self.settings["Ivtor_rmxu1_tdif"].get(),
            Ivtor_nn = self.settings["Ivtor_rmxu2_tdif"].get(),
            Ivtor_nn2 = self.settings["Ivtor_rmxu3_tdif"].get(),               # Третья сторона 
            n_sch_vn = self.settings["Nsch_rmxu1_tdif"].get(),
            n_sch_nn = self.settings["Nsch_rmxu2_tdif"].get(),
            n_sch_nn2 = self.settings["Nsch_rmxu3_tdif"].get(),                # Третья сторона 

            T1_pdif1_tdif = self.settings["T1_pdif1_tdif"].get(),
            Isr_pdif1_tdif = self.settings["Isr_pdif1_tdif"].get(),
            Isr_zagrub_pdif1_tdif = self.settings["Isrzagrub_pdif1_tdif"].get(),
            It1_pdif1_tdif = self.settings["It1_pdif1_tdif"].get(),
            It2_pdif1_tdif = self.settings["It2_pdif1_tdif"].get(),
            Kt1_pdif1_tdif = self.settings["Kt1_pdif1_tdif"].get(),
            Kt2_pdif1_tdif = self.settings["Kt2_pdif1_tdif"].get(),

            T1_pdif2_tdif = self.settings["T1_pdif2_tdif"].get(),
            Iset_pdif2_tdif = self.settings["Iset_pdif2_tdif"].get(),

            T1_hf2phar1_tdif = self.settings["T1_hf2phar1_tdif"].get(),
            T2_hf2phar1_tdif = self.settings["T2_hf2phar1_tdif"].get(),
            Ratio_hf2phar1_tdif = self.settings["Ratio_hf2phar1_tdif"].get(),

            T1_hf5phar1_tdif = self.settings["T1_hf5phar1_tdif"].get(),
            T2_hf5phar1_tdif = self.settings["T2_hf5phar1_tdif"].get(),
            Ratio_hf5phar1_tdif = self.settings["Ratio_hf5phar1_tdif"].get(),

            T1_rctr1_tdif = self.settings["T1_rctr1_tdif"].get(),
            Iset_rctr1_tdif = self.settings["Iset_rctr1_tdif"].get()
        )

        print("partDZT initialized")

    def start_polling(self):
        if self.part is None:
            print("partDZT not initialized")
            return

        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            #print(inputs)
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                if isinstance(value, (bool, np.bool_)):
                    label.config(text=f"{output}: {int(value)}")
                #label.config(text=f"{output}: {int(value)}")
                #print(f"{output}: {int(value)}")
                else: label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or round(float(value),2)!=0.0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3) # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05) # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):

        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return

        output_file = f"{function}_{mode}.xlsx"


        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        #output_file = "data.xlsx"
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)

        # Определяем красный цвет для заливки
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                # Устанавливаем ширину столбца
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20

                # Проверяем значения и применяем форматирование
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)  # Преобразуем значение в число
                        if value != 0:
                            cell.fill = red_fill  # Выделяем красным, если значение не равно 0
                    except (ValueError, TypeError):
                        pass  # Игнорируем ошибки преобразования

        # Применяем форматирование к каждому листу
        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        # Сохраняем изменения
        wb.save(output_file)

        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        # Выбор файла для загрузки
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        try:
            # Чтение данных из Excel
            xls = pd.ExcelFile(file_path)

            # Загрузка SGF Parameters
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            # Загрузка Inputs
            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            # Загрузка Settings
            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])                    

            print("Data loaded successfully")

        except Exception as e:
            print(f"Error loading data: {e}")        


if __name__ == "__main__":
    root = tk.Tk()
    app = PartDZT_GUI(root)
    root.mainloop()