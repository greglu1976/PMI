import tkinter as tk
from tkinter import ttk, messagebox
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension
import openpyxl
import json
# Импортируем класс CLS
from lib2.FUNCS.CLS import CLS
from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class CLSGUI:

    # === СПИСОК ВЫХОДНЫХ ПАРАМЕТРОВ (единое определение) ===
    OUTPUT_PARAMS = [
            "CLS_1_CLS_FuncEnabled", "CLS_1_CLS_MDResourceExcess", "CLS_1_CLS_CBLifeExcess", "CLS_1_CLS_COMMResourceExcess",
            "CLS_1_CLS_MDCurrentResource", "CLS_1_CLS_COMMCurrResourcePhsA", "CLS_1_CLS_COMMCurrResourcePhsB", "CLS_1_CLS_COMMCurrResourcePhsC"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование Контроля Ресурса Выключателя, v1.0")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        
        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")
        
        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для параметров SGF и настроек
        self.sgf_params = {
            "CLS_1_CLS_EnaDis": tk.IntVar(value=0),
            "CLS_1_CLS_MDCtrl": tk.IntVar(value=0),
        }
        
        self.settings = {
            "CLS_1_CLS_CBpasp_Inom": tk.DoubleVar(value=1000),
            "CLS_1_CLS_CBpasp_InomBr": tk.DoubleVar(value=31500),
            "CLS_1_CLS_COMMpasp_Inom": tk.DoubleVar(value=4000),
            "CLS_1_CLS_COMMpasp_InomBr": tk.DoubleVar(value=100),
            "CLS_1_CLS_CBpasp_MD": tk.IntVar(value=200000),
            "CLS_1_CLS_InitialCOMM": tk.DoubleVar(value=100),
            "CLS_1_CLS_COMMSet": tk.DoubleVar(value=20),
            "CLS_1_CLS_InitialMD": tk.IntVar(value=0),
            "CLS_1_CLS_TmaxCB": tk.IntVar(value=1), 
        }
        
        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "CLS_1_CBPosCls": tk.IntVar(value=0),
            "CLS_1_CBPosOpn": tk.IntVar(value=0),
            "CLS_1_ResetCounter": tk.IntVar(value=0),
            "CLS_1_OpnCB": tk.IntVar(value=0),
            "IA": tk.DoubleVar(value=0),
            "IB": tk.DoubleVar(value=0),
            "IC": tk.DoubleVar(value=0),
        }
        self.output_labels = {}

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc
        
        # Создание интерфейса
        self.create_widgets()
    
    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())
    
    def _get_output_names(self):
        return self.OUTPUT_PARAMS
    
    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 10:
                row = 0
                col += 2
        
        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.settings.items():
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2
        
        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        
        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        
        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=5, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=6, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        
        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=9, padx=5, pady=5)
        
        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)
                
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 5:
                row = 0
                col += 2
        
        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=75, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label

            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)

            row += 1
            if row >= 8:
                row = 0
                col += 2
    
    def init_part(self):
        self.part = CLS(
            SGF1=self.sgf_params["CLS_1_CLS_EnaDis"].get(),
            Inom_V_pasp=self.settings["CLS_1_CLS_CBpasp_Inom"].get(),
            Inom_otl_V_pasp=self.settings["CLS_1_CLS_CBpasp_InomBr"].get(),
            KRVpasp_Inom=self.settings["CLS_1_CLS_COMMpasp_Inom"].get(),
            KRVpasp_Inom_otkl=self.settings["CLS_1_CLS_COMMpasp_InomBr"].get(),
            MRVpasp=self.settings["CLS_1_CLS_CBpasp_MD"].get(),
            Nach_znach_KRV=self.settings["CLS_1_CLS_InitialCOMM"].get(),
            KRVsrab=self.settings["CLS_1_CLS_COMMSet"].get(),
            Nach_znach_MRV=self.settings["CLS_1_CLS_InitialMD"].get(),
            SGF2=self.sgf_params["CLS_1_CLS_MDCtrl"].get(),
            T1=self.settings["CLS_1_CLS_TmaxCB"].get(),
        )
        print("✅ CLS initialized")
    
    def start_polling(self):
        if self.part is None:
            print("⚠️ CLS not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()
    
    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("🛑 Polling stopped")
    
    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)
            
            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                label.config(text=f"{output}: {round(value, 2)}")
                if int(value) != 0 or float(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")
            
            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")
    
    def save_to_excel(self):
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        sgf_df = pd.DataFrame({key: [var.get()] for key, var in self.sgf_params.items()})
        settings_df = pd.DataFrame({key: [var.get()] for key, var in self.settings.items()})
        inputs_df = pd.DataFrame({key: [var.get()] for key, var in self.input_vars.items()})
        outputs_df = pd.DataFrame({key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()})

        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            for col_num, column in enumerate(sheet.columns, start=1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:
                        continue
                    try:
                        if float(cell.value) != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        for sheet_name in ["SGF_Parameters", "Settings", "Inputs", "Outputs"]:
            format_sheet(wb[sheet_name], None)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")
    
    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name, var_dict in [("SGF_Parameters", self.sgf_params),
                                          ("Settings", self.settings),
                                          ("Inputs", self.input_vars)]:
                if sheet_name in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    for key, var in var_dict.items():
                        if key in df.columns:
                            var.set(df.at[0, key])
            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = CLSGUI(root)
    root.mainloop()