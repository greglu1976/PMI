# автоматическое тестирование ФСУ в части КСВ, КП, КА, УВ + СС, ПС
# ДОБАВЛЕН УРОВ
# использовался для генерации Тестов от 02.04.25 Версия 1.
# доработанная версия от 21.07.25 Версия 2.
# доработанная версия от 24.02.2026 Версия 3.0 (добавлены подсказки, загрузка/сохранение JSON)


import tkinter as tk
from tkinter import ttk
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.dimensions import ColumnDimension
import openpyxl
import json

# Импортируем новый класс SWITCH3
from lib2.PARTS.SWITCH3_inout import SWITCH
from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler
from ToolTip import ToolTip

class PartOfSwitchGUI:

    # === СПИСОК ВЫХОДНЫХ ПАРАМЕТРОВ (единое определение) ===
    OUTPUT_PARAMS = [
        "T_LVCBSUP_1_RCBF1_FuncEnabled", "T_LVCBSUP_1_RCBF1_FuncOperDisabled",
        "T_LVCBSUP_1_RCBF1_UnpromptedCBopening", "T_LVCBSUP_1_RCBF1_FailureCB",
        "T_LVCBSUP_1_RCBF1_CBFailureTrip", "T_LVCBSUP_1_RCBF1_FixingContacts", "T_LVCBSUP_1_RCBF1_BlkToCls",
        "T_LVCBSUP_1_RCBF1_BlkToOpn", "T_LVCBSUP_1_RCBF1_ElmgLaunchFault", "T_LVCBSUP_1_RCBF1_ProtectCBCS",
        "T_LVCBSUP_1_RCBF1_ProtectCBOS1", "T_LVCBSUP_1_RCBF1_ProtectCBOS2", "T_SWCTRL_1_SWCTRL_FuncEnabled",
        "T_SWCTRL_1_SWCTRL_FuncOperDisabled", "T_SWCTRL_1_CBCSWI1_FuncEnabled", "T_SWCTRL_1_CBCSWI1_OpOpn",
        "T_SWCTRL_1_CBCSWI1_SwitchInProgress", "T_SWCTRL_1_CBCSWI1_OpTmAlm",
        "T_SWCTRL_1_CBCSWI1_OpCls", "T_SWCTRL_1_CBCSWI1_CBPosInterm",
        "T_SWCTRL_1_CBCSWI1_PosOpn", "T_SWCTRL_1_CBCSWI1_PosCls",
        "T_SWCTRL_1_CBCSWI1_CBPosFault", "T_HVBCTRL_1_CBCSWI1_FuncEnabled",
        "T_HVBCTRL_1_CBCSWI1_FuncOperDisabled", "T_HVBCTRL_1_CBCSWI1_OpCls", "T_SwitchDevice_1_SD_FuncEnabled",
        "T_SwitchDevice_1_SD_FuncOperDisabled", "T_SwitchDevice_1_CB1_FuncEnabled", "T_SwitchDevice_1_CB1_CBPosIntermed",
        "T_SwitchDevice_1_CB1_CBPosOpn", "T_SwitchDevice_1_CB1_CBPosCls", "T_SwitchDevice_1_CB1_CBPosFaul",
        "T_SwitchDevice_1_CB1_OpnCB_relay", "T_SwitchDevice_1_CB1_ClsCB_relay",
        "T_SignAssembly_1_SwOperExcTim", "T_LVALH_1_CALH1_Alarm", "T_TPBRF_1_GENRBRF1_OpIn",
        "T_HVTCBOFF_1_HVCBPTRC1_FuncEnabled", "T_HVTCBOFF_1_HVCBPTRC1_Op",
        "T_HVTCBOFF_1_HVCBPTRC1_Tr"
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование ФСУ в части КСВ, КП, КА, УВ, СС, ПС и УРОВ. v2.0 21.07.25, v3.0 24.02.2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False

        # Инициализация переменных для имени файла
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # Инициализация переменных для параметров SGF, настроек, входных и выходных значений
        self.sgf_params = {
            "T_LVCBSUP_1_RCBF1_EnaDis": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_RstFrmCLS": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_KnobCtrl": tk.IntVar(value=0),
            "T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm": tk.IntVar(value=0),
            "T_SWCTRL_1_SWCTRL_EnaDis": tk.IntVar(value=0),
            "T_SWCTRL_1_CBCSWI1_EnaDis": tk.IntVar(value=0),
            "T_SWCTRL_1_CBCSWI1_BlkToClsFrmFailureTrip": tk.IntVar(value=0),
            "T_HVBCTRL_1_CBCSWI1_EnaDis": tk.IntVar(value=0),
            "T_SwitchDevice_1_SD_EnaDis": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_EnaDis": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_TPOpnResetCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_CBOSoperationCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_TPClsResetCtrl": tk.IntVar(value=0),
            "T_SwitchDevice_1_CB1_CBCSoperationCtrl": tk.IntVar(value=0),
            #"SGF6_xcbr1_tsd": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_EnaDis": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CurrentPickUp": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_CBOSTypeCtrl": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_ActUpSwitch": tk.IntVar(value=0),
            "T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent": tk.IntVar(value=0),
            "T_HVTCBOFF_1_HVCBPTRC1_EnaDis": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_GASBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TECHBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ALMSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OCnnSign_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_OpExt_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_CtlCir_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_TestBlock_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl": tk.IntVar(value=0),
            "T_LVALH_1_CALH1_ExtSignGen_Ctl": tk.IntVar(value=0),
        }

        self.settings = {
            "T_LVCBSUP_1_RCBF1_T_EnBlk": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg": tk.DoubleVar(value=1),
            "T_LVCBSUP_1_RCBF1_T_ElmgWorking": tk.DoubleVar(value=1),
            "T_SWCTRL_1_CBCSWI1_TchangeCB": tk.DoubleVar(value=1),
            "T_SWCTRL_1_CBCSWI1_Tblk": tk.DoubleVar(value=1),
            #"T3_cbcswi1_swctrl": tk.DoubleVar(value=1),
            #"T4_cbcswi1_swctrl": tk.DoubleVar(value=1),
            #"T1_cbcswi1_hvbctrl": tk.DoubleVar(value=1),
            "T_SwitchDevice_1_CB1_TonFaul": tk.DoubleVar(value=1),
            "T_SwitchDevice_1_CB1_OpnTPtime": tk.DoubleVar(value=1),
            "T_SwitchDevice_1_CB1_ClsTPtime": tk.DoubleVar(value=1),
            "T_SwitchDevice_1_CB1_TextenCls": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Top": tk.DoubleVar(value=1),
            "T_TPBRF_1_GENRBRF1_Iop": tk.DoubleVar(value=0.2),
            "T_HVTCBOFF_1_HVCBPTRC1_Tpulse": tk.DoubleVar(value=1),            
        }
        
        self.input_vars = {
            "DI_ControllerDisable": tk.IntVar(value=0),
            "DI_LVCBSUP": tk.IntVar(value=0),
            "CBCS_CBOS1_OCControl": tk.IntVar(value=0),
            "CBOS2_OCControl": tk.IntVar(value=0),
            #"lovn_otkl": tk.IntVar(value=0),
            #"urov_nasebya": tk.IntVar(value=0),
            "InsTr": tk.IntVar(value=0),
            "LowIns": tk.IntVar(value=0),
            "EnBlk": tk.IntVar(value=0),
            "Reset": tk.IntVar(value=0),
            "OpnCBFrmKnob": tk.IntVar(value=0),
            "OperOpnCB": tk.IntVar(value=0),
            "T_LVCBSUP_1_ClsResourceExcess": tk.IntVar(value=0),
            "ExternalBlkCB": tk.IntVar(value=0),
            "CBCSCtrl": tk.IntVar(value=0),
            "CBOS1Ctrl": tk.IntVar(value=0),
            "CBOS2Ctrl": tk.IntVar(value=0),
            "CBCSWorking": tk.IntVar(value=0),
            "CBOS1Working": tk.IntVar(value=0),
            "CBOS2Working": tk.IntVar(value=0),
            "DI_SWCTRL": tk.IntVar(value=0),
            "OpnCBFrmCtrlPanel": tk.IntVar(value=0),
            "OpnCBFrm_HMI": tk.IntVar(value=0),
            "LocKey": tk.IntVar(value=0),
            "OpnCBFrmRemoteCtrl": tk.IntVar(value=0),
            "T_SWCTRL_1_OpnCBFrm_ACS": tk.IntVar(value=0),
            "KeyLocDist": tk.IntVar(value=0),
            "ClsCBFrmCtrlPanel": tk.IntVar(value=0),
            "ClsCBFrm_HMI": tk.IntVar(value=0),
            #"Remote": tk.IntVar(value=1),
            "ClsCBFrmRemoteCtrl": tk.IntVar(value=0),
            "T_SWCTRL_1_ClsCBFrm_ACS": tk.IntVar(value=0),
            "CBPosOpn": tk.IntVar(value=0),
            "CBPosCls": tk.IntVar(value=0),
            "DI_HVBCTRL": tk.IntVar(value=0),
            "OperClsCB": tk.IntVar(value=0),
            "DI_SD": tk.IntVar(value=0),
            #"lovn_lo_otkl_avar": tk.IntVar(value=0),
            "ExternalRBRFStart": tk.IntVar(value=0),
            "OpExtOfARC_NN": tk.IntVar(value=0),
            "OpExtOfCBFP_NN": tk.IntVar(value=0),
        }

        self.output_labels = {}

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names()) +
            list(self._get_output_names())  # <-- Добавили выходы
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        # Создание интерфейса
        self.create_widgets()

    def _get_sgf_param_names(self):
        return list(self.sgf_params.keys())
    
    def _get_setting_names(self):
        return list(self.settings.keys())
    
    def _get_input_names(self):
        return list(self.input_vars.keys())

    def create_widgets(self):
        # Frame for SGF parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        
        for key, var in self.sgf_params.items():
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1, 2], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 10:
                row = 0
                col += 2

        # Frame for settings
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0
        for key, var in self.settings.items():

            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            
            # Добавляем tooltip
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 9:
                row = 0
                col += 2

        # Frame for buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")

        # Button Init
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        # Button Start
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        # Button Stop
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        # Button Save
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        # Button Load
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)
        # Button Load JSON
        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        # Button Save JSON
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)

        # Поля для задания имени файла
        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Frame for input values
        input_frame = ttk.LabelFrame(self.root, text="Inputs")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row = 0
        col = 0

        for key, var in self.input_vars.items():
            if isinstance(var, tk.IntVar):
                cb = ttk.Checkbutton(input_frame, text=key, variable=var)
                cb.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(cb, tooltip)
            elif isinstance(var, tk.DoubleVar):
                label = ttk.Label(input_frame, text=key)
                label.grid(row=row, column=col, sticky="w")
                
                # Добавляем tooltip
                tooltip = self.tooltips.get(key)
                if tooltip:
                    ToolTip(label, tooltip)
                
                ttk.Entry(input_frame, textvariable=var).grid(row=row, column=col + 1)

            row += 1
            if row >= 5:
                row = 0
                col += 2

        # Frame for output values
        output_frame = ttk.LabelFrame(self.root, text="Outputs")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")

        row = 0
        col = 0
        for output in self.OUTPUT_PARAMS:
            label = ttk.Label(output_frame, text=output, width=35, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label


            # === ДОБАВЛЯЕМ TOOLTIP ИЗ META.JSON ===
            tooltip = self.tooltips.get(output)
            if tooltip:
                ToolTip(label, tooltip)


            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = SWITCH(
            SGF1_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_EnaDis"].get(),
            SGF2_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkToClsFrmLowIsol"].get(),
            SGF3_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkFrmCBPosFault"].get(),
            SGF4_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_RstFrmCLS"].get(),
            SGF5_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_OCcircuitFailureCtrl"].get(),
            SGF6_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_ConditionForElmgLaunchFault"].get(),
            SGF7_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_KnobCtrl"].get(),
            SGF8_rcbf1_lvcbsup=self.sgf_params["T_LVCBSUP_1_RCBF1_BlkCtrlFrmInsAlm"].get(),
            T1_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),#   self.settings["T_LVCBSUP_1_RCBF1_T_EnBlk"].get(),
            T2_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(),#    self.settings["T_LVCBSUP_1_RCBF1_T_FailureCtrlElmg"].get(),
            T3_rcbf1_lvcbsup=self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get()/1000 if self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get()>=100 else self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),#   self.settings["T_LVCBSUP_1_RCBF1_T_ElmgWorking"].get(),
            SGF1_swctrl=self.sgf_params["T_SWCTRL_1_SWCTRL_EnaDis"].get(),
            SGF1_cbcswi1_swctrl=self.sgf_params["T_SWCTRL_1_CBCSWI1_EnaDis"].get(),
            SGF2_cbcswi1_swctrl=self.sgf_params["T_SWCTRL_1_CBCSWI1_BlkToClsFrmFailureTrip"].get(),

            T1_cbcswi1_swctrl=self.settings["T_SWCTRL_1_CBCSWI1_TchangeCB"].get()/1000 if self.settings["T_SWCTRL_1_CBCSWI1_TchangeCB"].get()>=100 else self.settings["T_SWCTRL_1_CBCSWI1_TchangeCB"].get(),# self.settings["T_SWCTRL_1_CBCSWI1_TchangeCB"].get(),
            T2_cbcswi1_swctrl=self.settings["T_SWCTRL_1_CBCSWI1_Tblk"].get()/1000 if self.settings["T_SWCTRL_1_CBCSWI1_Tblk"].get()>=100 else self.settings["T_SWCTRL_1_CBCSWI1_Tblk"].get(),# self.settings["T_SWCTRL_1_CBCSWI1_Tblk"].get(),

            T3_cbcswi1_swctrl=0.5, #self.settings["T3_cbcswi1_swctrl"].get(),
            T4_cbcswi1_swctr=0.5, #self.settings["T4_cbcswi1_swctrl"].get(),
            SGF1_cbcswi1_hvbctrl=self.sgf_params["T_HVBCTRL_1_CBCSWI1_EnaDis"].get(),
            T1_cbcswi1_hvbctrl=0.5, #self.settings["T1_cbcswi1_hvbctrl"].get(),
            SGF1_tsd=self.sgf_params["T_SwitchDevice_1_SD_EnaDis"].get(),
            SGF1_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_EnaDis"].get(),
            SGF2_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_TPOpnResetCtrl"].get(),
            SGF3_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_CBOSoperationCtrl"].get(),
            SGF4_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_TPClsResetCtrl"].get(),
            SGF5_xcbr1_tsd=self.sgf_params["T_SwitchDevice_1_CB1_CBCSoperationCtrl"].get(),
            #SGF6_xcbr1_tsd=self.sgf_params["SGF6_xcbr1_tsd"].get(),

            T1_xcbr1_tsd= self.settings["T_SwitchDevice_1_CB1_TonFaul"].get()/1000 if self.settings["T_SwitchDevice_1_CB1_TonFaul"].get()>=100 else self.settings["T_SwitchDevice_1_CB1_TonFaul"].get(),# self.settings["T_SwitchDevice_1_CB1_TonFaul"].get(),
            T2_xcbr1_tsd= self.settings["T_SwitchDevice_1_CB1_OpnTPtime"].get()/1000 if self.settings["T_SwitchDevice_1_CB1_OpnTPtime"].get()>=100 else self.settings["T_SwitchDevice_1_CB1_OpnTPtime"].get(),#self.settings["T_SwitchDevice_1_CB1_OpnTPtime"].get(),
            T3_xcbr1_tsd= self.settings["T_SwitchDevice_1_CB1_ClsTPtime"].get()/1000 if self.settings["T_SwitchDevice_1_CB1_ClsTPtime"].get()>=100 else self.settings["T_SwitchDevice_1_CB1_ClsTPtime"].get(),#self.settings["T_SwitchDevice_1_CB1_ClsTPtime"].get(),
            T4_xcbr1_tsd= self.settings["T_SwitchDevice_1_CB1_TextenCls"].get()/1000 if self.settings["T_SwitchDevice_1_CB1_TextenCls"].get()>=100 else self.settings["T_SwitchDevice_1_CB1_TextenCls"].get(),#self.settings["T_SwitchDevice_1_CB1_TextenCls"].get(),

            SGF1_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_lvalh=self.sgf_params["T_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_lvalh=self.sgf_params["T_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_lvalh=self.sgf_params["T_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_lvalh=self.sgf_params["T_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_lvalh=self.sgf_params["T_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF14_lvalh=self.sgf_params["T_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_EnaDis"].get(),
            SGF2_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_BlkToOpnSpeedUp"].get(),
            SGF3_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CurrentPickUp"].get(),
            SGF4_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_CBOSTypeCtrl"].get(),
            SGF5_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_ActUpSwitch"].get(),
            SGF6_rbrf1_tpbrf=self.sgf_params["T_TPBRF_1_GENRBRF1_TypeOfCtrlCurrent"].get(),
            T1_rbrf1_tpbrf= self.settings["T_TPBRF_1_GENRBRF1_Top"].get()/1000 if self.settings["T_TPBRF_1_GENRBRF1_Top"].get()>=100 else self.settings["T_TPBRF_1_GENRBRF1_Top"].get(),#  self.settings["T_TPBRF_1_GENRBRF1_Top"].get(),
            Iset_rbrf1_tpbrf=self.settings["T_TPBRF_1_GENRBRF1_Iop"].get(),
            SGF1_hvcbptrc1_hvtcboff=self.sgf_params["T_HVTCBOFF_1_HVCBPTRC1_EnaDis"].get(),
            T1_hvcbptrc1_hvtcboff= self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get()/1000 if self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get()>=100 else self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(),# self.settings["T_HVTCBOFF_1_HVCBPTRC1_Tpulse"].get(),
        )
        print("part_SWITCH initialized")

    def start_polling(self):
        if self.part is None:
            print("part_SWITCH not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)

            # Обновление выходных значений
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                #if isinstance(label, tuple):
                #print(f"{output}: {value}")
                label.config(text=f"{output}: {round(value, 2)}")

                if int(value) != 0:
                    label.config(background="red", foreground="white")
                else:
                    label.config(background="green", foreground="white")

            time.sleep(0.3)  # Время шага опроса
            self.status_label.config(text="Шаг", background="white", foreground="white")
            time.sleep(0.05)  # Время шага опроса
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="#F0F0F0")

    def save_to_excel(self):
        # Формируем имя файла
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        # Создаем DataFrame для каждой группы данных
        sgf_df = pd.DataFrame({
            key: [var.get()] for key, var in self.sgf_params.items()
        })
        settings_df = pd.DataFrame({
            key: [var.get()] for key, var in self.settings.items()
        })
        inputs_df = pd.DataFrame({
            key: [var.get()] for key, var in self.input_vars.items()
        })
        outputs_df = pd.DataFrame({
            key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()
        })

        # Сохраняем данные в Excel
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        # Применяем форматирование к файлу Excel
        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            """Функция для форматирования листа."""
            for col_num, column in enumerate(sheet.columns, start=1):
                column_letter = openpyxl.utils.get_column_letter(col_num)
                sheet.column_dimensions[column_letter].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:  # Пропускаем заголовки
                        continue
                    try:
                        value = float(cell.value)
                        if value != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        format_sheet(wb["SGF_Parameters"], sgf_df)
        format_sheet(wb["Settings"], settings_df)
        format_sheet(wb["Inputs"], inputs_df)
        format_sheet(wb["Outputs"], outputs_df)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            sgf_df = pd.read_excel(xls, sheet_name="SGF_Parameters")
            for key, var in self.sgf_params.items():
                if key in sgf_df.columns:
                    var.set(sgf_df.at[0, key])

            settings_df = pd.read_excel(xls, sheet_name="Settings")
            for key, var in self.settings.items():
                if key in settings_df.columns:
                    var.set(settings_df.at[0, key])

            inputs_df = pd.read_excel(xls, sheet_name="Inputs")
            for key, var in self.input_vars.items():
                if key in inputs_df.columns:
                    var.set(inputs_df.at[0, key])

            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")

    # === МЕТОДЫ ДЛЯ РАБОТЫ С JSON ФАЙЛАМИ УСТАВОК ===
    
    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла с суффиксом _SG1."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return
        try:
            handler = SettingsHandler.from_json_file(file_path)
            if not self.meta_handler:
                print("⚠️ Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                    elif type_str == "130":  # Integer
                        try:
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()

                            int_val = int(float(value_clean.replace(',', '.')))
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            value_clean = str(value_str).strip()
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            # --- Обновление T-параметров (settings) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                        else:
                            self.settings[key].set(float_val)
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                except Exception as e:
                    print(f"❌ Ошибка при обработке {json_key}: {e}")
            
            print("✅ Параметры обновлены из JSON (с суффиксом _SG1)")
        except FileNotFoundError:
            print(f"❌ Ошибка: Файл не найден: {file_path}")
        except Exception as e:
            print(f"❌ Ошибка загрузки JSON: {e}")
    
    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()


    def _get_output_names(self):
        return self.OUTPUT_PARAMS       


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfSwitchGUI(root)
    root.mainloop()