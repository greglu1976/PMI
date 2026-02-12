########################
### Версия для 2026 года
########################
# По ней выполнено ПМИ М300-Т 10.02.2026
#######################


import tkinter as tk
from tkinter import ttk, messagebox
from tkinter.filedialog import askopenfilename
import threading
import time
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
import openpyxl
import json


from lib2.PARTS.TECH_T import part_TECH_T

from MainConfigHandler import MainConfigHandler
from SettingsHandler import SettingsHandler

from ToolTip import ToolTip


# --- Основной GUI ---
class PartOfTECH_T_GUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Тестирование Газовых и Технологических Защит М300-Т. вер1 от 2026")
        self.part = None
        self.polling_thread = None
        self.is_polling = False
        self.function_name = tk.StringVar(value="Функция")
        self.mode_name = tk.StringVar(value="Режим")

        # === ЗАГРУЗКА МЕТАДАННЫХ ===
        try:
            self.meta_handler = MainConfigHandler.from_json_file("meta.json")
        except Exception as e:
            print(f"⚠️ Не удалось загрузить meta.json: {e}")
            self.meta_handler = None

        # === ГЕНЕРАЦИЯ ПОДСКАЗОК ИЗ JSON ===
        self.tooltips = {}
        all_param_keys = (
            list(self._get_sgf_param_names()) +
            list(self._get_setting_names()) +
            list(self._get_input_names())
        )
        for key in all_param_keys:
            if self.meta_handler:
                desc = self.meta_handler.get_description_by_base_name(key)
                if desc:
                    self.tooltips[key] = desc

        # Инициализация параметров
        self.sgf_params = {name: tk.IntVar(value=0) for name in self._get_sgf_param_names()}
        self.settings = {name: tk.DoubleVar(value=1.0) for name in self._get_setting_names()}
        self.input_vars = {name: tk.IntVar(value=0) for name in self._get_input_names()}

        self.output_labels = {}
        self.create_widgets()

    def _get_sgf_param_names(self):
        return [
            "APTTECHLGC_1_OILPTRC1_EnaDis",
            "APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl",
            "APTTECHLGC_1_WINPTRC1_EnaDis",
            "APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl",
            "APTTECHLGC_1_VLVPTRC1_EnaDis",
            "APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl",
            "ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis",
            "ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis",
            "ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis",
            "TALMGASLGC_1_PTRC1_EnaDis",
            "TALMGASLGC_1_PTRC1_LowIsolSignCtrl",
            "TTRGASLGC_1_PTRC1_EnaDis",
            "TTRGASLGC_1_PTRC1_LowIsolTripCtrl",
            "TLTCGASLGC_1_PTRC1_EnaDis",
            "TLTCGASLGC_1_PTRC1_LowIsolTripCtrl",
            "TOFFLVLGC_1_PTRC1_EnaDis",
            "TOFFLVLGC_1_RBRE1_EnaDis",
            "TOFFLVLGC_1_RBRE1_PVOC2_Ctrl",
            "TOFFLVLGC_1_RBRE1_PVOC3_Ctrl",
            "TOFFLVLGC_1_LVCBRBLC1_EnaDis",
            "TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl",
            "TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl",
            "T_LVALH_1_CALH1_GASSign_Ctl",
            "T_LVALH_1_CALH1_LowIsolGAS_Ctl",
            "T_LVALH_1_CALH1_GASBlock_Ctl",
            "T_LVALH_1_CALH1_TECHSign_Ctl",
            "T_LVALH_1_CALH1_LowIsolTECH_Ctl",
            "T_LVALH_1_CALH1_TECHBlock_Ctl",
            "T_LVALH_1_CALH1_ALMSign_Ctl",
            "T_LVALH_1_CALH1_OCSign_Ctl",
            "T_LVALH_1_CALH1_OCnnSign_Ctl",
            "T_LVALH_1_CALH1_OpExt_Ctl",
            "T_LVALH_1_CALH1_CtlCir_Ctl",
            "T_LVALH_1_CALH1_TestBlock_Ctl",
            "T_LVALH_1_CALH1_SwOperExcTim_Ctl",
            "T_LVALH_1_CALH1_ExtSignGen_Ctl",
            "T_SignAssembly_1_Ctl_SA1",
            "T_SignAssembly_1_Ctl_SA2",
            "T_SignAssembly_1_Ctl_SA3",
            "T_SignAssembly_1_Ctl_SA4",
            "T_SignAssembly_1_Ctl_SA5",
            "T_SignAssembly_1_Ctl_SG1",
            "T_SignAssembly_1_Ctl_SG2",
            "T_SignAssembly_1_Ctl_GAS_OCControl",
            "T_SignAssembly_1_Ctl_TECH_OCControl",
            "T_SignAssembly_1_Ctl_OCcir_CB",
            "T_SignAssembly_1_Ctl_ARCnn_OCControl",
            "T_SignAssembly_1_Ctl_CBFPnn_OCControl",
            "T_SignAssembly_1_Ctl_IEDvt_OCControl",
        ]

    def _get_setting_names(self):
        return [
            "APTTECHLGC_1_LLN0_TopOnBlk",
            "TALMGASLGC_1_PTRC1_TopOnBlk",
            "TTRGASLGC_1_LLN0_TopOnBlk",
            "TLTCGASLGC_1_LLN0_TopOnBlk",
        ]

    def _get_input_names(self):
        return [
            "VYVOD", "OV_tz", "OV_dtm", "OV_dto", "OV_rd", "NaSign_dtm", "NaSign_dto", "NaSign_rd",
            "srabKontOtkl_m", "srabKontSign_m", "srabKontOtkl_o", "srabKontSign_o", "srabKontOtkl_rd",
            "srabKI_m", "srabKI_o", "srabKI_rd", "Sbros", "OV_ts", "OV_pk", "OV_ok", "OV_lev",
            "NaSign_pk", "NaSign_ok", "NaSign_lev", "srabKontOtkl_pk", "srabKontOtkl_ok", "srabKontOtkl_lev",
            "OV_ptrc1_talmgaslgc", "NaOtkl_ptrc1_talmgaslgc", "srabKont_ptrc1_talmgaslgc", "srabKI_ptrc1_talmgaslgc",
            "OV_ptrc1_ttrgaslgc", "NaSign_ptrc1_ttrgaslgc", "srabKont_ptrc1_ttrgaslgc", "srabKI_ptrc1_ttrgaslgc",
            "OV_ptrc1_tltcgaslgc", "NaSign_ptrc1_tltcgaslgc", "srabKont_ptrc1_tltcgaslgc", "srabKI_ptrc1_tltcgaslgc",
            "OV_tofflvlg", "OVlo_tofflvlg", "OVzapv_tofflvlg", "OVzavr_tofflvlg", "oil_t_hi_level",
            "oil_ltc_hi_level", "oil_ltc_lo_level", "oil_ltc_lo_temp",
        ]

    def create_widgets(self):
        # SGF Parameters
        sgf_frame = ttk.LabelFrame(self.root, text="SGF Parameters")
        sgf_frame.grid(row=0, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key in self._get_sgf_param_names():
            var = self.sgf_params[key]
            label = ttk.Label(sgf_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            ttk.Combobox(sgf_frame, textvariable=var, values=[0, 1], state="readonly").grid(row=row, column=col + 1)
            row += 1
            if row >= 14:
                row = 0
                col += 2

        # Settings (T-parameters)
        settings_frame = ttk.LabelFrame(self.root, text="Settings")
        settings_frame.grid(row=1, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key in self._get_setting_names():
            var = self.settings[key]
            label = ttk.Label(settings_frame, text=key)
            label.grid(row=row, column=col, sticky="w")
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(label, tooltip)
            ttk.Entry(settings_frame, textvariable=var).grid(row=row, column=col + 1)
            row += 1
            if row >= 3:
                row = 0
                col += 2

        # Buttons
        buttons_frame = ttk.LabelFrame(self.root, text="Buttons")
        buttons_frame.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        ttk.Button(buttons_frame, text="Init", command=self.init_part).grid(row=0, column=0, pady=10)
        ttk.Button(buttons_frame, text="Start", command=self.start_polling).grid(row=0, column=1, pady=10)
        ttk.Button(buttons_frame, text="Stop", command=self.stop_polling).grid(row=0, column=2, pady=10)
        ttk.Button(buttons_frame, text="Save", command=self.save_to_excel).grid(row=0, column=3, pady=10)
        ttk.Button(buttons_frame, text="Load", command=self.load_from_excel).grid(row=0, column=4, pady=10)

        ttk.Button(buttons_frame, text="Load JSON", command=self.load_settings_from_json).grid(row=0, column=5, padx=2, pady=5)
        ttk.Button(buttons_frame, text="Save JSON", command=self.save_settings_to_json).grid(row=0, column=6, padx=2, pady=5)

        ttk.Label(buttons_frame, text="Функция:").grid(row=0, column=7, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.function_name, width=15).grid(row=0, column=8, padx=5, pady=5)
        ttk.Label(buttons_frame, text="Режим:").grid(row=0, column=9, padx=5, pady=5)
        ttk.Entry(buttons_frame, textvariable=self.mode_name, width=15).grid(row=0, column=10, padx=5, pady=5)

        # Добавляем новый элемент (например, Label) с возможностью изменения цвета
        self.status_label = ttk.Label(buttons_frame, text="Шаг", background="green", foreground="white")
        self.status_label.grid(row=0, column=11, padx=5, pady=5)

        # Input Parameters
        input_frame = ttk.LabelFrame(self.root, text="Входные параметры")
        input_frame.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        row, col = 0, 0
        for key in self._get_input_names():
            var = self.input_vars[key]
            cb = ttk.Checkbutton(input_frame, text=key, variable=var)
            cb.grid(row=row, column=col, sticky="w")
            tooltip = self.tooltips.get(key)
            if tooltip:
                ToolTip(cb, tooltip)
            row += 1
            if row >= 6:
                row = 0
                col += 2

        # Output Parameters
        output_frame = ttk.LabelFrame(self.root, text="Выходные параметры")
        output_frame.grid(row=0, column=1, rowspan=4, padx=10, pady=10, sticky="nsew")
        outputs = [
            "vvod_oilptrc1_apttechlgc", "oper_vyvod_oilptrc1_apttechlgc", "srab_oilptrc1_apttechlgc", "srabsign_oilptrc1_apttechlgc",
            "zablok_oilptrc1_apttechlgc", "ET_oilptrc1_apttechlgc", "vvod_winptrc1_apttechlgc", "oper_vyvod_winptrc1_apttechlgc",
            "srab_winptrc1_apttechlgc", "srabsign_winptrc1_apttechlgc", "zablok_winptrc1_apttechlgc", "ET_winptrc1_apttechlgc",
            "vvod_vlvptrc1_apttechlgc", "oper_vyvod_vlvptrc1_apttechlgc", "srab_vlvptrc1_apttechlgc", "srabsign_vlvptrc1_apttechlgc",
            "zablok_vlvptrc1_apttechlgc", "ET_vlvptrc1_apttechlgc", "vvod_prvlvptrc1_almtechlgc", "oper_vyvod_prvlvptrc1_almtechlgc",
            "srab_prvlvptrc1_almtechlgc", "srabsign_prvlvptrc1_almtechlgc", "vvod_shvlvptrc1_almtechlgc", "oper_vyvod_shvlvptrc1_almtechlgc",
            "srab_shvlvptrc1_almtechlgc", "srabsign_shvlvptrc1_almtechlgc", "vvod_levptrc1_almtechlgc", "oper_vyvod_levptrc1_almtechlgc",
            "srab_levptrc1_almtechlgc", "srabsign_levptrc1_almtechlgc", "vvod_ptrc1_talmgaslgc", "oper_vyvod_ptrc1_talmgaslgc",
            "srab_ptrc1_talmgaslgc", "srabsign_ptrc1_talmgaslgc", "zablok_ptrc1_talmgaslgc", "ET_ptrc1_talmgaslgc",
            "vvod_ptrc1_ttrgaslgc", "oper_vyvod_ptrc1_ttrgaslgc", "srab_ptrc1_ttrgaslgc", "srabsign_ptrc1_ttrgaslgc",
            "zablok_ptrc1_ttrgaslgc", "ET_ptrc1_ttrgaslgc", "vvod_ptrc1_tltcgaslgc", "oper_vyvod_ptrc1_tltcgaslgc",
            "srab_ptrc1_tltcgaslgc", "srabsign_ptrc1_tltcgaslgc", "zablok_ptrc1_tltcgaslgc", "ET_ptrc1_tltcgaslgc",
            "vvod_ptrc1_tofflvlgc", "oper_vyvod_ptrc1_tofflvlgc", "pusk_ptrc1_tofflvlgc", "srab_ptrc1_tofflvlgc",
            "vvod_rblc1_tofflvlgc", "oper_vyvod_rblc1_tofflvlgc", "zapret_rblc1_tofflvlgc",
            "vvod_rbre1_tofflvlgc", "oper_vyvod_rbre1_tofflvlgc", "zapret_rbre1_tofflvlgc",
            "SS_gz_sign", "SS_gz_zablok", "SS_gz_nizk_isol", "SS_tz_sign", "SS_tz_nizk_isol", "SS_tz_zablok",
            "SS_ts_sign", "SS_vnesh_otkl", "SS_vyh_zepi_razobr", "SS_bi_vyved", "SS_ot_sign",
            "SS_neispr_ot_gz", "SS_neispr_ot_tz", "SS_neispr_ot_v", "SS_ot_nn_sign",
            "SS_prev_vrem_per_ka", "SS_obsh_vnesh_sign", "pusk_lvalv"
        ]
        row, col = 0, 0
        for output in outputs:
            label = ttk.Label(output_frame, text=output, width=33, anchor="w")
            label.grid(row=row, column=col, sticky="w")
            self.output_labels[output] = label
            # Подсказки для выходов можно добавить позже, если они есть в JSON
            row += 1
            if row >= 32:
                row = 0
                col += 2

    def init_part(self):
        self.part = part_TECH_T(
            SGF1_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_EnaDis"].get(),
            SGF2_oilptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_OILPTRC1_LowIsolTripCtrl"].get(),
            SGF1_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_EnaDis"].get(),
            SGF2_winptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_WINPTRC1_LowIsolTripCtrl"].get(),
            SGF1_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_EnaDis"].get(),
            SGF2_vlvptrc1_apttechlgc=self.sgf_params["APTTECHLGC_1_VLVPTRC1_LowIsolTripCtrl"].get(),
            T1_apttechlgc=self.settings["APTTECHLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_prvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_PRVLVPTRC1_EnaDis"].get(),
            SGF1_shvlvptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_SHVLVPTRC1_EnaDis"].get(),
            SGF1_levptrc1_almtechlgc=self.sgf_params["ALMTECHLGC_UIRZ_1_LEVPTRC1_EnaDis"].get(),
            SGF1_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_talmgaslgc=self.sgf_params["TALMGASLGC_1_PTRC1_LowIsolSignCtrl"].get(),
            T1_ptrc1_talmgaslgc=self.settings["TALMGASLGC_1_PTRC1_TopOnBlk"].get(),
            SGF1_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_ttrgaslgc=self.sgf_params["TTRGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_ttrgaslgc=self.settings["TTRGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_EnaDis"].get(),
            SGF2_ptrc1_tltcgaslgc=self.sgf_params["TLTCGASLGC_1_PTRC1_LowIsolTripCtrl"].get(),
            T1_ptrc1_tltcgaslgc=self.settings["TLTCGASLGC_1_LLN0_TopOnBlk"].get(),
            SGF1_ptrc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_PTRC1_EnaDis"].get(),
            SGF1_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_EnaDis"].get(),
            SGF2_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC2_Ctrl"].get(),
            SGF3_rbre1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_RBRE1_PVOC3_Ctrl"].get(),
            SGF1_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_EnaDis"].get(),
            SGF2_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC2_Ctrl"].get(),
            SGF3_rblc1_tofflvlgc=self.sgf_params["TOFFLVLGC_1_LVCBRBLC1_PVOC3_Ctrl"].get(),
            SGF1_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASSign_Ctl"].get(),
            SGF2_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolGAS_Ctl"].get(),
            SGF3_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_GASBlock_Ctl"].get(),
            SGF4_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHSign_Ctl"].get(),
            SGF5_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_LowIsolTECH_Ctl"].get(),
            SGF6_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TECHBlock_Ctl"].get(),
            SGF7_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_ALMSign_Ctl"].get(),
            SGF8_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCSign_Ctl"].get(),
            SGF9_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OCnnSign_Ctl"].get(),
            SGF10_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_OpExt_Ctl"].get(),
            SGF11_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_CtlCir_Ctl"].get(),
            SGF12_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_TestBlock_Ctl"].get(),
            SGF13_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_SwOperExcTim_Ctl"].get(),
            SGF14_t_lvalh=self.sgf_params["T_LVALH_1_CALH1_ExtSignGen_Ctl"].get(),
            SGF1_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA1"].get(),
            SGF2_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA2"].get(),
            SGF3_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA3"].get(),
            SGF4_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA4"].get(),
            SGF5_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SA5"].get(),
            SGF6_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SG1"].get(),
            SGF7_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_SG2"].get(),
            SGF8_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_GAS_OCControl"].get(),
            SGF9_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_TECH_OCControl"].get(),
            SGF10_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_OCcir_CB"].get(),
            SGF11_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_ARCnn_OCControl"].get(),
            SGF12_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_CBFPnn_OCControl"].get(),
            SGF13_t_signassembly=self.sgf_params["T_SignAssembly_1_Ctl_IEDvt_OCControl"].get(),
        )
        print("part_TECH_T initialized")

    def start_polling(self):
        if not self.part:
            print("part_TECH_T not initialized")
            return
        self.is_polling = True
        self.polling_thread = threading.Thread(target=self.poll_inputs, daemon=True)
        self.polling_thread.start()

    def stop_polling(self):
        self.is_polling = False
        if self.polling_thread and self.polling_thread.is_alive():
            self.polling_thread.join(timeout=1.0)
        print("Polling stopped")

    def poll_inputs(self):
        while self.is_polling:
            inputs = {key: var.get() for key, var in self.input_vars.items()}
            result = self.part.Step(**inputs)
            for output, value in zip(self.output_labels.keys(), result):
                label = self.output_labels[output]
                label.config(text=f"{output}: {round(value, 2)}")
                bg_color = "red" if int(value) != 0 else "green"
                label.config(background=bg_color, foreground="white")
            time.sleep(0.3)
            self.status_label.config(text="Шаг", background="white", foreground="black")
            time.sleep(0.05)
            self.status_label.config(text="Шаг", background="#F0F0F0", foreground="black")

    def save_to_excel(self):
        function = self.function_name.get().strip()
        mode = self.mode_name.get().strip()
        if not function or not mode:
            print("Поля 'Функция' и 'Режим' должны быть заполнены")
            return
        output_file = f"{function}_{mode}.xlsx"

        sgf_df = pd.DataFrame({key: [var.get()] for key, var in self.sgf_params.items()})
        settings_df = pd.DataFrame({key: [var.get()] for key, var in self.settings.items()})
        inputs_df = pd.DataFrame({key: [var.get()] for key, var in self.input_vars.items()})
        outputs_df = pd.DataFrame({key: [label.cget("text").split(": ")[-1]] for key, label in self.output_labels.items()})

        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            sgf_df.to_excel(writer, sheet_name="SGF_Parameters", index=False)
            settings_df.to_excel(writer, sheet_name="Settings", index=False)
            inputs_df.to_excel(writer, sheet_name="Inputs", index=False)
            outputs_df.to_excel(writer, sheet_name="Outputs", index=False)

        wb = openpyxl.load_workbook(output_file)
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

        def format_sheet(sheet, df):
            for col_num, column in enumerate(sheet.columns, start=1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20
                for row_num, cell in enumerate(column, start=1):
                    if row_num == 1:
                        continue
                    try:
                        if float(cell.value) != 0:
                            cell.fill = red_fill
                    except (ValueError, TypeError):
                        pass

        for sheet_name in ["SGF_Parameters", "Settings", "Inputs", "Outputs"]:
            format_sheet(wb[sheet_name], None)

        wb.save(output_file)
        print(f"Data saved to {output_file} with formatting")

    def load_from_excel(self):
        file_path = askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name, var_dict in [("SGF_Parameters", self.sgf_params),
                                          ("Settings", self.settings),
                                          ("Inputs", self.input_vars)]:
                if sheet_name in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    for key, var in var_dict.items():
                        if key in df.columns:
                            var.set(df.at[0, key])
            print("Data loaded successfully")
        except Exception as e:
            print(f"Error loading data: {e}")




    # МЕТОДЫ ДЛЯ РАБОТЫ С JSON ФАЙЛАМИ УСТАВОК

    def load_settings_from_json(self):
        """Загружает SGF и T-параметры из JSON-файла, добавляя '_SG1' к именам."""
        file_path = askopenfilename(filetypes=[("JSON files", "*.json")])
        if not file_path:
            return

        try:
            handler = SettingsHandler.from_json_file(file_path)
            
            if not self.meta_handler:
                messagebox.showwarning("Предупреждение", "Метаданные не загружены. Используется стандартная обработка.")
            
            # --- Обновление SGF-параметров (с _SG1) ---
            for key in self.sgf_params:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                    
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Преобразование значения в зависимости от типа
                    if type_str == "3":  # Булевое значение
                        value_lower = str(value_str).lower().strip()
                        bool_map = {
                            "true": 1, "1": 1, "on": 1, "вкл": 1, "да": 1, "yes": 1, "enabled": 1,
                            "false": 0, "0": 0, "off": 0, "выкл": 0, "нет": 0, "no": 0, "disabled": 0
                        }
                        
                        if value_lower in bool_map:
                            self.sgf_params[key].set(bool_map[value_lower])
                        else:
                            # Пробуем числовое преобразование
                            try:
                                num_val = float(value_str)
                                self.sgf_params[key].set(1 if num_val != 0 else 0)
                            except ValueError:
                                print(f"⚠️ Неизвестное булевое значение для {json_key}: '{value_str}'")
                                self.sgf_params[key].set(0)
                                
                    elif type_str == "130":  # Integer
                        try:
                            # Удаляем возможные единицы измерения
                            value_clean = str(value_str).strip()
                            for suffix in ['%', '°', '°C', 'мс', 'с', 'м']:
                                if value_clean.endswith(suffix):
                                    value_clean = value_clean[:-len(suffix)].strip()
                            
                            # Преобразуем в int
                            int_val = int(float(value_clean.replace(',', '.')))  # Обрабатываем 1.0, 1,5
                            self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Ошибка преобразования int для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                            
                    else:  # По умолчанию или неизвестный тип - пробуем как int
                        try:
                            # Пробуем разные форматы
                            value_clean = str(value_str).strip()
                            
                            # Сначала пробуем как булевое
                            value_lower = value_clean.lower()
                            bool_map = {
                                "true": 1, "1": 1, "on": 1, "вкл": 1,
                                "false": 0, "0": 0, "off": 0, "выкл": 0
                            }
                            
                            if value_lower in bool_map:
                                self.sgf_params[key].set(bool_map[value_lower])
                            else:
                                # Пробуем как число
                                int_val = int(float(value_clean.replace(',', '.')))
                                self.sgf_params[key].set(int_val)
                        except (ValueError, TypeError) as e:
                            print(f"⚠️ Не удалось преобразовать значение для {json_key}: '{value_str}' - {e}")
                            self.sgf_params[key].set(0)
                            
                except Exception as e:
                    print(f"Ошибка при обработке {json_key}: {e}")

            # --- Обновление T-параметров (с _SG1) ---
            for key in self.settings:
                json_key = key + "_SG1"
                value_str = handler.get_value_by_parameter(json_key)
                if value_str is None:
                    continue
                    
                try:
                    # Определяем тип из метаданных
                    type_str = None
                    if self.meta_handler:
                        param_info = self.meta_handler.get_param_info(json_key)
                        if param_info:
                            type_str = param_info.get("type")
                            print(f"{json_key}: type={type_str}, value='{value_str}'")
                    
                    # Для settings обычно используются float значения
                    value_clean = str(value_str).strip()
                    
                    # Убираем единицы измерения
                    units_to_remove = ['%', '°', '°c', '°с', 'мс', 'с', 'м', 'мм', 'кг', 'кпа', 'па']
                    for unit in units_to_remove:
                        if value_clean.lower().endswith(unit):
                            value_clean = value_clean[:-len(unit)].strip()
                    
                    # Заменяем запятую на точку
                    value_clean = value_clean.replace(',', '.')
                    
                    # Пробуем преобразовать в float
                    try:
                        float_val = float(value_clean)
                        
                        # Проверяем разумные пределы для settings
                        if abs(float_val) > 1000000:
                            print(f"⚠️ Подозрительно большое значение для {json_key}: {float_val}")
                            # Можно установить значение по умолчанию
                            # self.settings[key].set(1.0)
                        else:
                            self.settings[key].set(float_val)
                            
                    except ValueError as e:
                        print(f"⚠️ Невозможно преобразовать в число: {json_key} = '{value_str}' - {e}")
                        
                except Exception as e:
                    print(f"Ошибка при обработке {json_key}: {e}")

            messagebox.showinfo("Успех", "Уставки успешно загружены из JSON-файла.")
            print("Параметры обновлены из JSON (с суффиксом _SG1)")

        except FileNotFoundError:
            messagebox.showerror("Ошибка", f"Файл не найден: {file_path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить уставки:\n{str(e)}")
            print(f"Ошибка загрузки JSON: {e}")


    def save_settings_to_json(self):
        """Сохраняет SGF и T-параметры в JSON-файл с суффиксом _SG1."""
        file_path = askopenfilename(
            title="Сохранить уставки как...",
            filetypes=[("JSON files", "*.json")],
            defaultextension=".json"
        )
        if not file_path:
            return
        try:
            # Пытаемся загрузить существующий файл, иначе создаём пустой
            try:
                handler = SettingsHandler.from_json_file(file_path)
            except (FileNotFoundError, json.JSONDecodeError, KeyError, ValueError):
                handler = SettingsHandler([])  # пустой обработчик
            
            # --- 1. Сохраняем SGF-параметры ---
            for base_key in self.sgf_params:
                json_key = base_key + "_SG1"
                raw_value = self.sgf_params[base_key].get()

                # Определяем тип параметра
                param_type = "3"  # значение по умолчанию — бинарный
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Преобразуем значение в число
                try:
                    numeric_value = float(raw_value)
                    if not numeric_value.is_integer():
                        numeric_value = int(round(numeric_value))
                    else:
                        numeric_value = int(numeric_value)
                except (ValueError, TypeError):
                    numeric_value = 0

                # Форматируем в зависимости от типа
                if param_type == "3":
                    formatted = "1" if numeric_value != 0 else "0"
                elif param_type == "130":
                    formatted = str(numeric_value)
                else:
                    formatted = str(numeric_value)

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 SGF {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # --- 2. Сохраняем T-параметры (settings) ---
            for base_key in self.settings:
                json_key = base_key + "_SG1"  # ← ВАЖНО: тоже добавляем _SG1!
                raw_value = self.settings[base_key].get()

                # Определяем тип параметра из метаданных
                param_type = None
                if self.meta_handler:
                    param_info = self.meta_handler.get_param_info(json_key)
                    if param_info and "type" in param_info:
                        param_type = str(param_info["type"])

                # Для T-параметров обычно используется float
                try:
                    float_val = float(raw_value)
                    # Сохраняем с разумной точностью (убираем лишние нули)
                    formatted = f"{float_val:.6g}"
                except (ValueError, TypeError):
                    formatted = "0.0"

                handler.add_or_update_parameter(json_key, formatted)
                print(f"💾 T   {json_key} = {formatted} (type={param_type}, raw={raw_value})")
            
            # Сохраняем
            handler.save_to_json_file(file_path)
            messagebox.showinfo("Успех", f"Уставки сохранены в:\n{file_path}")
            print(f"✅ Уставки сохранены в {file_path}")
        except Exception as e:
            error_msg = f"Ошибка при сохранении уставок:\n{str(e)}"
            messagebox.showerror("Ошибка", error_msg)
            print(f"❌ {error_msg}")
            import traceback
            traceback.print_exc()


if __name__ == "__main__":
    root = tk.Tk()
    app = PartOfTECH_T_GUI(root)
    root.mainloop()