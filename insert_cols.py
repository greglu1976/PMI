import os
import openpyxl

def process_excel_files():
    """
    1. Удаляет столбцы по списку масок во всех файлах.
    2. Специально для листа 'Inputs':
       - Удаляет столбец 'Remote'
       - Вставляет столбец 'LocKey' (значение 0) между 'OpnCBFrm_HMI' и 'OpnCBFrmRemoteCtrl'
    """
    # Список столбцов для удаления (по частичному совпадению)
    cols_to_delete = []
        #"T4_cbcswi1_swctrl", "T3_cbcswi1_swctrl", 
        #"SGF6_xcbr1_tsd", "T1_cbcswi1_hvbctrl", "distanz"
    #]
    
    # Настройки для специального листа Inputs
    target_sheet_name = "Inputs"
    col_to_delete_specific = "Remote"
    new_col_name = "LocKey"
    new_col_value = 0
    anchor_col_left = "OpnCBFrm_HMI"      # После этого столбца
    anchor_col_right = "OpnCBFrmRemoteCtrl" # Перед этим столбцом

    current_dir = os.getcwd()
    xlsx_files = [f for f in os.listdir(current_dir) if f.endswith('.xlsx') and not f.startswith('~$')]
    
    print(f"Найдено {len(xlsx_files)} xlsx файлов")
    print("-" * 50)
    
    for file_name in xlsx_files:
        file_path = os.path.join(current_dir, file_name)
        print(f"\nОбработка: {file_name}")
        
        try:
            wb = openpyxl.load_workbook(file_path)
            changes_made = False
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                headers = []
                
                # Сбор заголовков
                for col_idx in range(1, sheet.max_column + 1):
                    header = sheet.cell(row=1, column=col_idx).value
                    headers.append((col_idx, str(header) if header else ""))
                
                # --- СПЕЦИАЛЬНАЯ ЛОГИКА ДЛЯ ЛИСТА Inputs ---
                if sheet_name == target_sheet_name:
                    idx_remote = None
                    idx_left = None
                    idx_right = None
                    
                    # Поиск индексов нужных столбцов
                    for col_idx, header in headers:
                        if header == col_to_delete_specific:
                            idx_remote = col_idx
                        if header == anchor_col_left:
                            idx_left = col_idx
                        if header == anchor_col_right:
                            idx_right = col_idx

                    # 1. Удаление столбца 'Remote'
                    if idx_remote:
                        sheet.delete_cols(idx_remote)
                        print(f"  Лист '{sheet_name}': удален столбец '{col_to_delete_specific}'")
                        changes_made = True
                        
                        # Корректировка индексов, если Remote был до правого якоря
                        if idx_right and idx_remote < idx_right:
                            idx_right -= 1
                        if idx_left and idx_remote < idx_left:
                            idx_left -= 1

                    # 2. Вставка столбца 'LocKey'
                    # Вставляем после idx_left (то есть в позицию idx_left + 1)
                    if idx_left:
                        insert_index = idx_left + 1
                        sheet.insert_cols(insert_index)
                        cell = sheet.cell(row=1, column=insert_index, value=new_col_name)
                        cell.font = openpyxl.styles.Font(bold=True)
                        
                        # Заполнение значением 0 (начиная со 2-й строки, чтобы не затереть заголовок)
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(row=row, column=insert_index, value=new_col_value)
                            
                        print(f"  Лист '{sheet_name}': вставлен столбец '{new_col_name}' на позицию {insert_index}")
                        changes_made = True
                    elif not idx_remote:
                        print(f"  Лист '{sheet_name}': столбцы для модификации не найдены ({anchor_col_left} или {col_to_delete_specific})")

                # --- ОБЩАЯ ЛОГИКА УДАЛЕНИЯ ПО МАСКАМ ---
                # Собираем заново заголовки, если лист был изменен (для Inputs), 
                # либо используем старые (для остальных листов)
                if sheet_name == target_sheet_name and changes_made:
                     # Обновляем список заголовков после манипуляций с Inputs, 
                     # чтобы не удалить только что вставленный LocKey по ошибке, 
                     # если он попадет под маску (маловероятно, но надежнее)
                    headers = []
                    for col_idx in range(1, sheet.max_column + 1):
                        header = sheet.cell(row=1, column=col_idx).value
                        headers.append((col_idx, str(header) if header else ""))

                cols_to_remove = []
                for col_idx, header in reversed(headers):
                    if header:
                        for pattern in cols_to_delete:
                            if pattern in header:
                                # Проверка: не удаляем только что вставленный LocKey
                                if sheet_name == target_sheet_name and header == new_col_name:
                                    continue
                                cols_to_remove.append(col_idx)
                                print(f"  Лист '{sheet_name}': удален столбец {col_idx} ('{header}') по маске '{pattern}'")
                                break
                
                # Удаление столбцов по маске
                if cols_to_remove:
                    for col_idx in sorted(cols_to_remove, reverse=True):
                        sheet.delete_cols(col_idx)
                    changes_made = True
            
            # Сохранение
            if changes_made:
                wb.save(file_path)
                print(f"  ✅ Изменения сохранены")
            else:
                print(f"  ℹ️  Изменений не требуется")
            
            wb.close()
            
        except Exception as e:
            print(f"  ❌ Ошибка: {e}")
    
    print("\n" + "-" * 50)
    print("Готово!")
    input("Нажмите Enter, чтобы выйти...")

if __name__ == "__main__":
    process_excel_files()